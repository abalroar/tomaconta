from io import BytesIO
from zipfile import ZipFile

import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import pytest
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn
from pptx.util import Inches

from utils.sgs_credit_pptx_export import (
    exportar_deck_secoes_pptx,
    exportar_figuras_pptx,
    figura_para_painel,
)


SOURCE = "Fonte: Banco Central do Brasil · ConsultaUnificada"


def _meta(**overrides):
    return {
        "chart_title": "Taxas por produto",
        "source": SOURCE,
        "value_format": "0.00%",
        "value_scale": 0.01,
        "formato_data": "diaria",
        **overrides,
    }


def _charts(deck):
    return [shape.chart for slide in deck.slides for shape in slide.shapes if shape.has_chart]


def test_daily_export_keeps_dates_gaps_values_and_editable_workbook():
    datas = pd.to_datetime(["2026-06-01", "2026-06-02", "2026-06-03", "2026-06-04"])
    fig = go.Figure(go.Scatter(x=datas, y=[2.31, None, 2.45, None], name="Banco A", line_color="#117733"))
    fig.update_layout(meta=_meta())
    blob, meta = exportar_figuras_pptx([fig], titulo_deck="Taxas", perfil="taxas")
    deck = Presentation(BytesIO(blob))
    chart = _charts(deck)[0]

    assert meta["detalhe"][0]["categorias"] == 4
    assert [from_excel(float(item)) for item in chart.plots[0].categories] == datas.to_pydatetime().tolist()
    assert chart.category_axis._element.tag == qn("c:dateAx")
    assert chart.category_axis._element.find(qn("c:baseTimeUnit")).get("val") == "days"
    assert list(chart.series[0].values) == pytest.approx([0.0231, None, 0.0245, None])
    assert chart._chartSpace.chart.find(qn("c:dispBlanksAs")).get("val") == "gap"
    assert chart.series[0].format.line.color.rgb.__str__() == "117733"
    texts = [shape.text for shape in deck.slides[0].shapes if shape.has_text_frame]
    assert SOURCE in texts
    assert not any("BCData/SGS" in text for text in texts)
    with ZipFile(BytesIO(blob)) as archive:
        books = [name for name in archive.namelist() if name.startswith("ppt/embeddings/") and name.endswith(".xlsx")]
        assert len(books) == 1
        workbook = load_workbook(BytesIO(archive.read(books[0])), data_only=True)
        rows = list(workbook.active.values)
        assert [row[0] for row in rows[1:]] == datas.to_pydatetime().tolist()
        assert rows[2][1] is None
        assert rows[3][1] == 0.0245
        assert workbook.active["B2"].number_format == "0.00%"


def test_monthly_export_preserves_all_months_in_workbook():
    dates = pd.date_range("2025-01-31", periods=18, freq="ME")
    fig = go.Figure(go.Scatter(x=dates, y=[2.5] * 18, name="Banco A"))
    fig.update_layout(meta=_meta(formato_data="mensal"))
    blob, _ = exportar_figuras_pptx([fig], titulo_deck="Mensal", perfil="taxas")
    chart = _charts(Presentation(BytesIO(blob)))[0]
    categories = list(chart.plots[0].categories)
    assert len(categories) == 18
    assert "\u00a0" not in categories
    assert chart.category_axis._element.find(qn("c:majorUnit")).get("val") == "3"
    assert chart.category_axis._element.find(qn("c:majorTimeUnit")).get("val") == "months"


def test_horizontal_ranking_keeps_order_colors_and_rate_scale():
    frame = pd.DataFrame({"banco": ["Banco de Nome Longo", "Banco B", "Banco C"], "taxa": [1.4, 2.1, 3.5]})
    colors = {"Banco de Nome Longo": "#117733", "Banco B": "#CC6677", "Banco C": "#332288"}
    fig = px.bar(frame, x="taxa", y="banco", color="banco", orientation="h", color_discrete_map=colors)
    fig.update_layout(meta=_meta(subtitulo="Ranking · % a.m."))
    blob, meta = exportar_figuras_pptx([fig], titulo_deck="Ranking", perfil="taxas")
    chart = _charts(Presentation(BytesIO(blob)))[0]
    assert chart.chart_type == XL_CHART_TYPE.BAR_CLUSTERED
    categorias = list(chart.plots[0].categories)
    assert categorias == list(fig.layout.yaxis.categoryarray)
    valores_por_banco = frame.set_index("banco").taxa.to_dict()
    assert list(chart.series[0].values) == pytest.approx([valores_por_banco[banco] / 100 for banco in categorias])
    assert [str(point.format.fill.fore_color.rgb) for point in chart.series[0].points] == [colors[banco][1:] for banco in categorias]
    assert chart.value_axis.minimum_scale == 0
    assert chart.value_axis.tick_labels.number_format == "0.00%"
    assert meta["detalhe"][0]["rotulos"] == 3
    for point in chart.series[0].points:
        label = point.data_label._dLbl
        assert label.find(qn("c:showVal")).get("val") == "1"
        assert label.find(qn("c:showSerName")).get("val") == "0"
        assert label.find(qn("c:dLblPos")).get("val") == "outEnd"


def test_facets_become_separate_native_panels_without_secondary_axes():
    frame = pd.DataFrame([
        {"data": pd.Timestamp("2026-01-31") + pd.offsets.MonthEnd(mes), "taxa": banco + mes / 10, "banco": f"Banco {banco}"}
        for banco in range(1, 6) for mes in range(3)
    ])
    fig = px.line(frame, x="data", y="taxa", color="banco", facet_col="banco", facet_col_wrap=3)
    fig.update_layout(meta=_meta(formato_data="mensal", separar_paineis=True))
    blob, meta = exportar_figuras_pptx([fig], titulo_deck="Painéis", perfil="taxas")
    deck = Presentation(BytesIO(blob))
    assert meta["slides"] == 2
    assert meta["paineis"] == 5
    assert len(_charts(deck)) == 5
    assert all(not panel["eixo_secundario"] for panel in meta["detalhe"])
    assert all(len(chart.series) == 1 for chart in _charts(deck))
    assert all(len(chart._chartSpace.chart.plotArea.findall(qn("c:valAx"))) == 1 for chart in _charts(deck))
    assert len({(chart.value_axis.minimum_scale, chart.value_axis.maximum_scale) for chart in _charts(deck)}) == 1


def test_taxas_layout_is_opt_in_and_each_regular_figure_has_full_slide():
    fig = go.Figure(go.Scatter(x=["2026-01-31", "2026-02-28"], y=[2, 3], name="Banco A"))
    fig.update_layout(meta=_meta(formato_data="mensal"))
    baseline, _ = exportar_figuras_pptx([fig], titulo_deck="Padrão")
    taxas, meta = exportar_figuras_pptx([fig, fig], titulo_deck="Taxas", perfil="taxas")
    shape_baseline = next(shape for shape in Presentation(BytesIO(baseline)).slides[0].shapes if shape.has_chart)
    deck_taxas = Presentation(BytesIO(taxas))
    assert meta["slides"] == 2
    assert shape_baseline.width < Inches(7)
    assert all(shape.width > Inches(12) for slide in deck_taxas.slides for shape in slide.shapes if shape.has_chart)
    assert figura_para_painel(fig).estilo_taxas is False


def test_months_missing_from_first_institution_stay_chronological():
    fig = go.Figure([
        go.Scatter(x=["2026-01-31", "2026-03-31"], y=[2, 3], name="Banco A"),
        go.Scatter(x=["2026-02-28", "2026-03-31"], y=[4, 5], name="Banco B"),
    ])
    fig.update_layout(meta=_meta(formato_data="mensal"))
    painel = figura_para_painel(fig, perfil="taxas")
    assert painel.ordem_categorias == ["2026-01", "2026-02", "2026-03"]
    blob, _ = exportar_figuras_pptx([fig], titulo_deck="Lacunas", perfil="taxas")
    chart = _charts(Presentation(BytesIO(blob)))[0]
    assert list(chart.series[0].values) == [0.02, None, 0.03]
    assert list(chart.series[1].values) == [None, 0.04, 0.05]
    assert chart.has_legend
    assert chart.legend.include_in_layout is False


def test_text_only_sections_survive_and_use_readable_body():
    blob, meta = exportar_deck_secoes_pptx(
        [("Glossário", ("Saldo: valor informado na data de referência.", "Fonte: BCB"), []),
         ("Cobertura", ("", "Fonte consultada indisponível."), [])],
        titulo_deck="Crédito BC",
    )
    deck = Presentation(BytesIO(blob))
    assert meta["secoes"] == 2
    assert meta["slides"] == 3
    assert meta["paineis"] == 0
    text_shape = next(shape for shape in deck.slides[1].shapes if shape.has_text_frame and shape.text.startswith("Saldo:"))
    assert text_shape.text_frame.paragraphs[0].runs[0].font.size.pt == 15
    assert any(shape.has_text_frame and "Fonte consultada indisponível." in shape.text for shape in deck.slides[2].shapes)
