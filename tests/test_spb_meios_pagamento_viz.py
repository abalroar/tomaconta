from __future__ import annotations

import sys
import zipfile
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

from utils.spb_meios_pagamento_viz import (
    ITAU_BBA_PALETTE,
    build_line_figure,
    build_spb_native_pptx,
    build_spb_pptx,
    build_share_figure,
    compute_share,
    format_ano_mes_label,
    format_trimestre_label,
    melt_nucleo_spb,
    latest_summary,
    style_spb_figure,
)


def test_period_labels_match_dashboard_spec():
    assert format_ano_mes_label("202605") == "05-2026"
    assert format_trimestre_label("20251") == "mar-25"
    assert format_trimestre_label("20254") == "dez-25"


def test_melt_nucleo_monthly_adds_sortable_period_and_labels():
    raw = pd.DataFrame(
        {
            "ano_mes": ["202501", "202502"],
            "quantidadePix": [10, 20],
            "valorPix": [100, 200],
            "quantidadeTED": [5, 4],
            "valorTED": [50, 40],
        }
    ).rename(
        columns={
            "quantidadePix": "quantidade_pix",
            "valorPix": "valor_pix",
            "quantidadeTED": "quantidade_ted",
            "valorTED": "valor_ted",
        }
    )

    long_df = melt_nucleo_spb(raw, "ano_mes", "mensal")

    assert set(long_df["instrumento"]) == {"Pix", "TED"}
    assert set(long_df["periodo_label"]) == {"01-2025", "02-2025"}
    assert long_df["periodo_ordem"].min() == pd.Timestamp("2025-01-01")


def test_compute_share_sums_to_100_by_period():
    raw = pd.DataFrame(
        {
            "trimestre": ["20251", "20252"],
            "quantidade_pix": [80, 75],
            "quantidade_ted": [20, 25],
            "valor_pix": [40, 50],
            "valor_ted": [60, 50],
        }
    )
    long_df = melt_nucleo_spb(raw, "trimestre", "trimestral")

    share = compute_share(long_df, "Quantidade (mil)", ["Pix", "TED"])

    totals = share.groupby("periodo_label")["participacao"].sum().round(6).tolist()
    assert totals == [100.0, 100.0]


def test_figures_render_final_point_labels():
    raw = pd.DataFrame(
        {
            "ano_mes": ["202501", "202502"],
            "quantidade_pix": [10, 20],
            "quantidade_ted": [5, 4],
            "valor_pix": [100, 200],
            "valor_ted": [50, 40],
        }
    )
    long_df = melt_nucleo_spb(raw, "ano_mes", "mensal")

    line_fig = build_line_figure(
        long_df,
        tipo="Quantidade (mil)",
        instruments=["Pix", "TED"],
        title="Quantidade",
        yaxis_title="mil transações",
    )
    share_fig = build_share_figure(
        long_df,
        tipo="Quantidade (mil)",
        instruments=["Pix", "TED"],
        title="Participação",
    )

    assert line_fig.data[0].text[-1] == "20"
    assert share_fig.data[0].text[-1].endswith("%")


def test_native_pptx_contains_powerpoint_chart_parts():
    raw = pd.DataFrame(
        {
            "ano_mes": ["202501", "202502"],
            "quantidade_pix": [10, 20],
            "quantidade_ted": [5, 4],
            "valor_pix": [100, 200],
            "valor_ted": [50, 40],
        }
    )
    long_df = melt_nucleo_spb(raw, "ano_mes", "mensal")

    pptx_bytes = build_spb_native_pptx(
        title="Dados de Meios de Pagamento - BCB",
        charts=[
            {
                "title": "Núcleo mensal - quantidade",
                "long_df": long_df,
                "tipo": "Quantidade (mil)",
                "instruments": ["Pix", "TED"],
                "yaxis_title": "mil transações",
            },
            {
                "title": "Participação mensal - quantidade",
                "long_df": long_df,
                "tipo": "Quantidade (mil)",
                "instruments": ["Pix", "TED"],
                "yaxis_title": "% do total selecionado",
                "share": True,
            },
        ],
        summaries={
            "Quantidade": compute_share(long_df, "Quantidade (mil)", ["Pix", "TED"]).rename(
                columns={"instrumento": "Instrumento", "valor": "Valor", "participacao": "Participação", "periodo_label": "Período"}
            )[["Período", "Instrumento", "Valor", "Participação"]],
        },
        source_note="Fonte: BCB",
    )

    with zipfile.ZipFile(BytesIO(pptx_bytes)) as pptx_zip:
        chart_parts = [name for name in pptx_zip.namelist() if name.startswith("ppt/charts/chart") and name.endswith(".xml")]
        slide_parts = [name for name in pptx_zip.namelist() if name.startswith("ppt/slides/slide") and name.endswith(".xml")]
        slide_xml = b"\n".join(pptx_zip.read(name) for name in slide_parts)

    assert len(chart_parts) == 2
    assert len(slide_parts) == 4
    assert b"Imagem do gr" not in slide_xml
    assert "Plano de melhorias aplicado".encode("utf-8") not in slide_xml
    assert "Manter mensal e trimestral separados".encode("utf-8") not in slide_xml


def test_legacy_pptx_export_fails_instead_of_placeholder_when_image_missing():
    with pytest.raises(RuntimeError, match="build_spb_native_pptx"):
        build_spb_pptx(
            title="Dados de Meios de Pagamento - BCB",
            figures={"Núcleo mensal - quantidade": None},
            summaries={},
            source_note="Fonte: BCB",
        )


def test_native_pagination_preserves_all_series_values_and_share_denominator():
    from pptx import Presentation

    instruments = list(ITAU_BBA_PALETTE)[:13]
    periods = pd.date_range("2023-01-01", periods=30, freq="MS")
    long_df = pd.DataFrame([
        {
            "periodo": period, "periodo_ordem": period,
            "periodo_label": period.strftime("%m-%Y"),
            "instrumento": instrument, "tipo": tipo, "valor": 100.0,
        }
        for period in periods for instrument in instruments
        for tipo in ("Quantidade (mil)", "Valor (R$ milhão)")
    ])
    original = long_df.copy(deep=True)
    output = build_spb_native_pptx(
        title="SPB",
        charts=[{
            "title": "Participação", "long_df": long_df,
            "tipo": "Quantidade (mil)", "instruments": instruments,
            "share": True, "yaxis_title": "% do total selecionado",
        }],
        summaries={
            "Quantidade": latest_summary(long_df, "Quantidade (mil)", instruments),
            "Valor": latest_summary(long_df, "Valor (R$ milhão)", instruments),
        },
        source_note="Fonte: BCB",
    )
    presentation = Presentation(BytesIO(output))
    charts = [shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
    tables = [shape.table for slide in presentation.slides for shape in slide.shapes if shape.has_table]
    assert len(presentation.slides) == 5  # cover, two charts, two summary pages
    assert len(charts) == 2
    assert [series.name for chart in charts for series in chart.series] == instruments
    for chart in charts:
        assert len(chart.plots[0].categories) == len(periods)
        for series in chart.series:
            assert tuple(series.values) == pytest.approx([1 / 13] * len(periods))
    assert len(tables) == 4
    summary_names = [row.cells[0].text for table in tables for row in list(table.rows)[1:]]
    assert sorted(summary_names) == sorted(instruments * 2)
    pd.testing.assert_frame_equal(long_df, original)


def test_styling_preserves_observations_and_stable_series_colors():
    import plotly.express as px

    frame = pd.DataFrame({"periodo": ["mar-25", "jun-25"] * 2, "valor": [0, 3.2, 1.8, 2.4], "serie": ["Crédito"] * 2 + ["Débito"] * 2})
    chart = px.line(frame, x="periodo", y="valor", color="serie")
    observations = [(tuple(series.x), tuple(series.y)) for series in chart.data]
    style_spb_figure(chart, title="MDR", yaxis_title="%", series_order=["Crédito", "Débito"])
    assert [(tuple(series.x), tuple(series.y)) for series in chart.data] == observations
    single = px.line(frame[frame.serie == "Débito"], x="periodo", y="valor", color="serie")
    style_spb_figure(single, title="MDR", yaxis_title="%", series_order=["Crédito", "Débito"])
    assert single.data[0].line.color == chart.data[1].line.color


def test_trailing_and_total_missing_values_preserve_gaps_and_zero():
    from pptx import Presentation

    periods = pd.date_range("2025-01-01", periods=3, freq="MS")
    long_df = pd.DataFrame([
        {
            "periodo_ordem": period, "periodo_label": period.strftime("%m-%Y"),
            "instrumento": instrument, "tipo": "Quantidade (mil)", "valor": value,
        }
        for instrument, values in [("Pix", [10.0, 0.0, float("nan")]), ("TED", [float("nan")] * 3)]
        for period, value in zip(periods, values)
    ])
    original = long_df.copy(deep=True)
    figure = build_line_figure(long_df, tipo="Quantidade (mil)", instruments=["Pix", "TED"], title="Quantidade", yaxis_title="mil transações")
    assert figure.data[0].text[1] == "0"
    assert pd.isna(figure.data[0].y[2])
    assert all(pd.isna(value) for value in figure.data[1].y)
    label_traces = [trace for trace in figure.data if trace.meta and trace.meta.get("spb_label_only")]
    assert len(label_traces) == 1
    assert label_traces[0].x[0] == "02-2025"
    assert label_traces[0].text[0].strip() == "0"
    assert label_traces[0].legendgroup == figure.data[0].legendgroup == "Pix"
    all_missing = build_line_figure(long_df[long_df.instrumento == "TED"], tipo="Quantidade (mil)", instruments=["TED"], title="Sem dados", yaxis_title="mil transações")
    assert len(all_missing.data) == 1

    output = build_spb_native_pptx(
        title="SPB", charts=[{"title": "Quantidade", "long_df": long_df, "tipo": "Quantidade (mil)", "instruments": ["Pix", "TED"]}],
        summaries={}, source_note="Fonte: BCB",
    )
    presentation = Presentation(BytesIO(output))
    chart = next(shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart)
    assert [series.name for series in chart.series] == ["Pix"]
    assert tuple(chart.series[0].values) == (10.0, 0.0, None)
    pd.testing.assert_frame_equal(long_df, original)
