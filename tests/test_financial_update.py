"""Revisão de contas, decisão de linha e atualização por etapas."""
import json
from pathlib import Path

import pandas as pd
import pytest

from tests.test_prudential_analysis import simple_sources  # noqa: F401
from utils.ifdata_cache.financial_analysis_review import account_changes, new_decisions, pending_items, row_jumps
from utils.prudential_analysis import dre_definitions, dre_row_for


def frame(rows):
    return pd.DataFrame([{"DATA_BASE": p, "DOCUMENTO": d, "COD_CONGL": c, "CONTA": a, "NOME_CONTA": n, "SALDO": s} for p, d, c, a, n, s in rows])


BASE = [("202605", "4060", "C1", "8199200005", "Provisão", -5e9), ("202605", "4060", "C1", "1110000006", "Caixa", 1e9),
        ("202605", "4060", "C1", "3000000007", "Compensação", 9e12)]


@pytest.mark.parametrize("code,row", [("8199200005", "provision"), ("8197000009", "other_exp"), ("7191500002", "sale"),
                                      ("8172700005", "people"), ("7940000001", "tax_income"), ("7500000006", None)])
def test_row_follows_the_longest_prefix(code, row):
    assert dre_row_for(code) == row


def test_recoded_provision_stops_for_review_and_decision_restores_the_row():
    now = [("202606", "4060", "C1", "8197000009", "Nova", -5e9), ("202606", "4060", "C1", "1110000006", "Caixa", 1e9)]
    items = account_changes(frame(BASE + now), "202606")
    by_kind = {item["kind"]: item for item in items}
    assert by_kind["nova"]["account"] == "8197000009" and by_kind["nova"]["row"] == "other_exp" and by_kind["nova"]["blocking"]
    assert by_kind["sumiu"]["row"] == "provision" and by_kind["sumiu"]["blocking"]
    assert pending_items(items, {by_kind["nova"]["id"]: {"row": "provision"}}) == [by_kind["sumiu"]["id"]]
    rules = {ident: rule for ident, _, _, rule in dre_definitions({"8197000009": "provision"})}
    assert ("8197000009", 1) in rules["provision"] and ("8197000009", -1) in rules["other_exp"]


def test_compensation_small_accounts_and_code_format_do_not_stop():
    now = [("202606", "4060", "C1", "3090000000", "Comp", 9e12), ("202606", "4060", "C1", "8197000009", "Pequena", -1e6),
           ("202606", "4060", "C1", "81992005", "Provisão oito dígitos", -5e9), ("202606", "4060", "C1", "1110000006", "Caixa", 1e9)]
    items = account_changes(frame(BASE + now), "202606")
    assert not pending_items(items, {})
    assert [item["account"] for item in items if item["kind"] == "nova"] == ["8197000009"]


def test_semester_close_is_expected_only_for_result_accounts_after_the_close():
    july = account_changes(frame(BASE[:2] + [("202607", "4060", "C1", "1110000006", "Caixa", 1e9)]), "202607")
    august = account_changes(frame(BASE[:2] + [("202608", "4060", "C1", "1110000006", "Caixa", 1e9)]), "202608")
    assert not july[0]["blocking"] and august[0]["blocking"]


def test_rename_is_visible_without_stopping():
    now = [("202606", "4060", "C1", "8199200005", "Outro nome", -5e9), ("202606", "4060", "C1", "1110000006", "Caixa", 1e9)]
    item = next(item for item in account_changes(frame(BASE + now), "202606") if item["kind"] == "renomeada")
    assert item["previous_name"] == "Provisão" and not item["blocking"]


@pytest.mark.parametrize("current,previous,flag", [(-100, -90, False), (100, -90, True), (-400, -100, True), (-20, -100, True), (0, 0, False)])
def test_row_jump(current, previous, flag):
    assert bool(row_jumps({"provision": [current, previous]}, "202606")) is flag


def test_decision_records_the_account_and_rejects_unknown_rows():
    recorded = new_decisions({"nova:4060:8197000009": {"row": "provision"}, "salto:provision:202606": {"row": None}}, "202606", "run")
    assert [item["account"] for item in recorded] == ["8197000009", None]
    with pytest.raises(ValueError):
        new_decisions({"nova:4060:8197000009": {"row": "x"}}, "202606", "run")
    with pytest.raises(ValueError):
        dre_definitions({"8197000009": "x"})


def test_full_flow_holds_review_then_applies_and_keeps_the_decision(tmp_path, monkeypatch, simple_sources):
    from tools import update_financial_analysis as cli
    from utils.ifdata_cache import BloprudencialCache
    monkeypatch.setattr("utils.ifdata_cache.release_ops.materialize_for_publication", lambda *args, **kwargs: [])
    root = tmp_path / "repo"; root.mkdir()
    path = tmp_path / "source.csv"
    simple_sources.rename(columns={"DATA_BASE": "#DATA_BASE"}).to_csv(path, sep=";", decimal=",", index=False)
    assert cli.update(root, csv=path, offline=True)[0]["status"] == "updated"
    # Mês seguinte com uma conta de resultado inédita, compensada para os totais fecharem.
    following = simple_sources.assign(DATA_BASE="202607")
    extra = following[following.DOCUMENTO.eq("4060")].iloc[:1].assign(CONTA="8197000009", NOME_CONTA="Nova", SALDO=-2e8)
    offset = following[following.DOCUMENTO.eq("4060")].iloc[:1].assign(CONTA="7197000000", NOME_CONTA="Par", SALDO=2e8)
    following = pd.concat([following, extra, offset])
    for account, change in (("8000000002", -2e8), ("8100000005", -2e8), ("7000000003", 2e8), ("7100000006", 2e8)):
        following.loc[following.DOCUMENTO.eq("4060") & following.CONTA.eq(account), "SALDO"] += change
    following.rename(columns={"DATA_BASE": "#DATA_BASE"}).to_csv(path, sep=";", decimal=",", index=False)
    active = BloprudencialCache(root).arquivo_dados
    before = active.read_bytes()
    checked, _ = cli.update(root, csv=path, offline=True, check_only=True)
    assert checked["status"] == "validated", checked
    assert checked["pending_review"] and active.read_bytes() == before
    held, _ = cli.update(root, csv=path, offline=True)
    assert held["status"] == "needs_review" and active.read_bytes() == before
    answers = {item: {"row": "provision"} if item.endswith("8197000009") else {"row": None} for item in held["pending_review"]}
    applied, _ = cli.update(root, csv=path, offline=True, decisions=answers)
    assert applied["status"] == "updated", applied
    cache = BloprudencialCache(root)
    assert cache.row_overrides() == {"8197000009": "provision"}


def test_decisions_survive_a_later_plain_save(tmp_path):
    from utils.ifdata_cache import BloprudencialCache
    cache = BloprudencialCache(tmp_path)
    data = pd.DataFrame({"DATA_BASE": ["202606"], "Período": ["202606"], "SALDO": [1.0]})
    decision = [{"account": "8197000009", "row": "provision"}]
    assert cache.salvar_local(data, fonte="teste", metadata_extra={"cosif_decisions": decision}).sucesso
    assert cache.salvar_local(data, fonte="teste").sucesso
    assert cache.cosif_decisions() == decision


def test_panel_marks_a_killed_process_as_interrupted_and_reads_the_latest_report(tmp_path):
    from tabs.financial_update_panel import _reports_dir, job_state, latest_report
    directory = _reports_dir(tmp_path)
    (directory / "job.json").write_text(json.dumps({"pid": 99999999, "stage": "x", "running": True}))
    assert job_state(tmp_path)["interrupted"]
    (directory / "20260101T000000-a.json").write_text('{"status": "failed"}')
    (directory / "20260201T000000-b.json").write_text('{"status": "updated"}')
    (directory / "decisions.json").write_text("{}")
    assert latest_report(tmp_path)[0]["status"] == "updated"


def test_installed_source_enters_runtime_with_its_own_signature(tmp_path, monkeypatch):
    from types import SimpleNamespace
    import hashlib
    from utils.ifdata_cache import financial_analysis_update as subject
    source = tmp_path / "dados.parquet"; source.write_bytes(b"bytes")
    metadata = tmp_path / "metadata.json"; metadata.write_text('{"sha256": "antiga", "integridade": {}, "periodos": ["202606"]}')
    saved = []
    cache = SimpleNamespace(existe=lambda: False, existe_leitura=lambda: True, arquivo_dados=source, arquivo_metadata=metadata,
                            salvar_arquivo_local=lambda path, data: saved.append(data) or SimpleNamespace(sucesso=True, mensagem=""))
    monkeypatch.setattr("utils.ifdata_cache.release_ops._publication_scope", lambda names: (["a"], []))
    subject.activate_installed_sources(SimpleNamespace(get_cache=lambda name: cache), ["a"])
    assert saved == [{"sha256": hashlib.sha256(b"bytes").hexdigest(), "periodos": ["202606"]}]


def _published(frame, decisions=None):
    from hashlib import sha256
    from io import BytesIO
    from unittest.mock import Mock
    buffer = BytesIO()
    frame.to_parquet(buffer, index=False)
    raw = buffer.getvalue()
    def fake_get(url, **kwargs):
        response = Mock(status_code=200, content=raw)
        response.raise_for_status = lambda: None
        response.json = lambda: {"cosif_decisions": decisions or []}
        return response
    info = {"cache": "bloprudencial", "exists": True, "sha256": sha256(raw).hexdigest(), "record_count": len(frame), "max_period_ref": str(frame.DATA_BASE.max())}
    return info, fake_get, sha256(raw).hexdigest()


def _cosif(*periods):
    return pd.DataFrame([{"DATA_BASE": period, "Período": period, "DOCUMENTO": "4060", "COD_CONGL": "C1", "CONTA": "1000000009", "SALDO": 1.0} for period in periods])


def test_newer_published_month_replaces_the_installed_base_with_its_decisions(tmp_path, monkeypatch):
    from hashlib import sha256
    from utils.ifdata_cache import BloprudencialCache
    from utils.individual_release_cache import ensure_published_bloprudencial
    cache = BloprudencialCache(tmp_path)
    assert cache.salvar_local(_cosif("202606"), fonte="teste").sucesso
    decision = [{"id": "nova:4060:8197000009", "account": "8197000009", "row": "provision"}]
    info, fake_get, digest = _published(_cosif("202606", "202607"), decision)
    monkeypatch.setattr("utils.individual_release_cache.requests.get", fake_get)
    assert ensure_published_bloprudencial(cache, info, "https://exemplo/release") is True
    assert sha256(cache.read_data_file.read_bytes()).hexdigest() == digest
    assert cache.row_overrides() == {"8197000009": "provision"}
    # Uma nova sessão enxerga a mesma competência e não baixa de novo.
    monkeypatch.setattr("utils.individual_release_cache.requests.get", lambda *a, **k: pytest.fail("sem download"))
    fresh = BloprudencialCache(tmp_path)
    assert ensure_published_bloprudencial(fresh, info, "https://exemplo/release") is False
    assert pd.read_parquet(fresh.arquivo_dados).DATA_BASE.max() == "202607"


def test_older_or_equal_release_never_replaces_the_installed_base(tmp_path, monkeypatch):
    from utils.ifdata_cache import BloprudencialCache
    from utils.individual_release_cache import ensure_published_bloprudencial
    cache = BloprudencialCache(tmp_path)
    assert cache.salvar_local(_cosif("202606"), fonte="teste").sucesso
    before = cache.read_data_file.read_bytes()
    monkeypatch.setattr("utils.individual_release_cache.requests.get", lambda *a, **k: pytest.fail("sem download"))
    for period in ("202606", "202605"):
        assert ensure_published_bloprudencial(cache, {"max_period_ref": period, "sha256": "0" * 64}, "https://exemplo/release") is False
    assert cache.read_data_file.read_bytes() == before


def _individual_rows(period, code="60701190", name="BANCO A"):
    amounts = [("1000000009", 1000.0), ("1100000002", 1000.0), ("4000000006", 800.0), ("4100000009", 800.0), ("6000000004", 150.0),
               ("7000000003", 70.0), ("7100000006", 70.0), ("7110000003", 70.0), ("8000000002", -20.0), ("8100000005", -20.0), ("8110000002", -20.0)]
    return pd.DataFrame([{"DATA_BASE": period, "Período": period, "CNPJ": code, "CONTA": account, "NOME_CONTA": account, "SALDO": value * 1e6,
                          "NOME_INSTITUICAO": name, "TAXONOMIA": "X", "GRUPO_FONTE": "Bancos"} for account, value in amounts])


def test_individual_analysis_uses_the_same_rules_and_labels_its_scope(tmp_path):
    from io import BytesIO
    from openpyxl import load_workbook
    from pptx import Presentation
    from tabs.prudential_analysis import load_analysis
    from utils.ifdata_cache import CacheManager
    from utils.prudential_pptx import render_pptx
    from utils.prudential_workbook import render_xlsx
    manager = CacheManager(tmp_path)
    cache = manager.get_cache("balancete_individual")
    for period in ("202512", "202603", "202606"):
        cache.add_period(_individual_rows(period))
    assert cache.available_periods() == ["202512", "202603", "202606"]
    analysis = load_analysis(manager, "60701190", ["202512", "202603", "202606", "202609"], individual=True)
    assert analysis.periods == ["202512", "202603", "202606"] and analysis.scope_label == "Instituição individual"
    assert not any(check["status"] == "Divergência" and check["check"].startswith(("Balanço:", "DRE: resultado")) for check in analysis.checks)
    dre, rows = analysis.sheets["DRE"], analysis.rows["DRE"]
    assert dre.cells[f"G{rows['net']}"].value == pytest.approx(50.0)
    assert "4010" in dre.cells["C5"].value or "4010" in str(dre.cells["C6"].value)
    balance, balance_rows = analysis.sheets["Balanço"], analysis.rows["Balanço"]
    assert balance.cells[f"G{balance_rows['roe']}"].value == pytest.approx(50 / 200 * 2)
    workbook = load_workbook(BytesIO(render_xlsx(analysis)))
    assert workbook["Carteira 4.966"].sheet_state == "hidden"
    deck = Presentation(BytesIO(render_pptx(analysis)))
    assert len(deck.slides) == 1
    text = " ".join(shape.text_frame.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    assert "Instituição individual" in text and "COSIF 4010" in text and "4066" not in text


def test_individual_history_replaces_only_its_period_and_reads_by_institution(tmp_path):
    from utils.ifdata_cache.balancete_individual import BalanceteIndividualCache
    cache = BalanceteIndividualCache(tmp_path)
    cache.add_period(pd.concat([_individual_rows("202603"), _individual_rows("202603", code="11111111", name="BANCO B")]))
    cache.add_period(_individual_rows("202606"))
    changed = _individual_rows("202606"); changed["SALDO"] *= 2
    cache.add_period(changed)
    assert cache.available_periods() == ["202603", "202606"]
    loaded = cache.load_institution("60701190", ["202606"])
    assert loaded.SALDO.max() == 2000e6 and set(loaded.CNPJ) == {"60701190"}
    assert set(cache.catalog().CNPJ) == {"60701190", "11111111"}


def test_individual_review_flags_new_accounts_and_broken_totals(tmp_path):
    from utils.ifdata_cache.balancete_individual import BalanceteIndividualCache
    from utils.ifdata_cache.financial_analysis_update import individual_review
    cache = BalanceteIndividualCache(tmp_path)
    cache.add_period(_individual_rows("202603"))
    now = _individual_rows("202606")
    assert individual_review(_with(cache, now), "202606", {}) == []
    extra = now.iloc[:1].assign(CONTA="8197000009", NOME_CONTA="Nova", SALDO=-2e8)
    items = individual_review(_with(cache, pd.concat([now, extra])), "202606", {})
    assert {item["id"] for item in items} >= {"nova:4010:8197000009", "divergencia:4010:202606"}
    assert all(item["blocking"] for item in items)


def _with(cache, frame):
    cache.add_period(frame)
    return cache


def test_published_individual_history_is_adopted_only_when_newer(tmp_path, monkeypatch):
    from hashlib import sha256
    from unittest.mock import MagicMock
    from utils.ifdata_cache.balancete_individual import BalanceteIndividualCache, file_metadata
    from utils.individual_release_cache import ensure_published_individual_statements
    source = BalanceteIndividualCache(tmp_path / "origem")
    source.add_period(_individual_rows("202606")); source.add_period(_individual_rows("202609"))
    raw, metadata = source.arquivo_dados.read_bytes(), source.arquivo_metadata.read_bytes()
    def fake_get(url, **kwargs):
        response = MagicMock()
        response.__enter__.return_value = response
        response.iter_content = lambda size: [metadata if "metadata" in url else raw]
        return response
    monkeypatch.setattr("utils.individual_release_cache.requests.get", fake_get)
    cache = BalanceteIndividualCache(tmp_path / "site")
    cache.add_period(_individual_rows("202606"))
    info = {"max_period_ref": "202609", "sha256": sha256(raw).hexdigest()}
    assert ensure_published_individual_statements(cache, info, "https://exemplo") is True
    assert cache.available_periods() == ["202606", "202609"]
    assert ensure_published_individual_statements(cache, info, "https://exemplo") is False
    with pytest.raises(ValueError):
        ensure_published_individual_statements(BalanceteIndividualCache(tmp_path / "outro"), {**info, "sha256": "0" * 64}, "https://exemplo")


def test_statement_catalog_cache_is_keyed_by_the_view():
    import ast
    tree = ast.parse(Path("tabs/prudential_analysis.py").read_text(encoding="utf-8"))
    catalog = next(node for node in ast.walk(tree) if isinstance(node, ast.FunctionDef) and node.name == "catalog")
    assert [arg.arg for arg in catalog.args.args] == ["token", "individual_view"]
    # Variáveis livres ficam fora da chave do cache: a visão não pode vir do escopo externo.
    assert "individual" not in {node.id for node in ast.walk(catalog) if isinstance(node, ast.Name)}


def test_screen_note_keeps_every_currency_sign_literal(monkeypatch):
    from types import SimpleNamespace
    import tabs.prudential_analysis as subject
    monkeypatch.setattr(subject, "reference_note", lambda analysis: "* Diferença (R$ milhões).")
    note = subject.statement_note(SimpleNamespace())
    assert note.count(r"\$") == 2 and note.count("$") == 2


def test_new_generation_wins_over_the_installed_publication_after_file_promotion(tmp_path):
    from utils.ifdata_cache.balancete_individual import BalanceteIndividualCache, file_metadata
    origin = BalanceteIndividualCache(tmp_path / "origem")
    origin.add_period(_individual_rows("202606"))
    bundled = tmp_path / "site/data/bundled/balancete_individual"
    bundled.mkdir(parents=True)
    (bundled / "dados.parquet").write_bytes(origin.arquivo_dados.read_bytes())
    (bundled / "metadata.json").write_text(json.dumps({**file_metadata(origin.arquivo_dados), "publication_id": "instalada"}))
    cache = BalanceteIndividualCache(tmp_path / "site")
    assert cache.available_periods() == ["202606"]
    cache.add_period(_individual_rows("202609"))
    assert cache.available_periods() == ["202606", "202609"]
    assert cache.arquivo_dados == cache.arquivo_dados_runtime
