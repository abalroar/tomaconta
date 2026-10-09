from io import BytesIO
from zipfile import ZipFile

import pandas as pd
import pytest
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation

from utils.evolucao_pptx_export import (CHART_HEIGHT, LABEL_HEIGHT, PLOT_HEIGHT,
                                      PLOT_LEFT, PLOT_TOP, PLOT_WIDTH, export_evolucao_powerpoint)
from utils.evolucao_visual import METRIC_LABELS, SERIES


NS = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def sample(n=6):
    periods = [f"dez-{20 + i}" for i in range(n)]
    graph = pd.DataFrame({column: [(i + 1) * (s + 1) * 1e9 for i in range(n)]
                          for s, (column, *_) in enumerate(SERIES)})
    table = pd.DataFrame({"Métrica": list(METRIC_LABELS),
                          **{p: ["13,3%", "5,2x", "12,3%", "14,0%", "15,4%"] for p in periods}})
    return dict(df_graph=graph, df_show=table, periods=periods, institution="ITAU - PRUDENCIAL",
                queried_at="09/10/2026 18:30 -03")


def test_native_combination_table_metadata_and_fonts():
    blob = export_evolucao_powerpoint(**sample())
    prs = Presentation(BytesIO(blob))
    assert len(prs.slides) == 1
    chart_shape = next(s for s in prs.slides[0].shapes if s.has_chart)
    table_shape = next(s for s in prs.slides[0].shapes if s.has_table)
    assert chart_shape.left == table_shape.left
    assert chart_shape.width == table_shape.width
    assert chart_shape.top + chart_shape.height < table_shape.top
    chart = chart_shape.chart
    assert len(chart.plots) == 2
    assert [s.name for s in chart.series] == [s[1] for s in SERIES]
    assert chart.has_legend
    table = table_shape.table
    assert len(table.rows) == 6 and len(table.columns) == 7
    assert [table.cell(r, 0).text for r in range(1, 6)] == list(METRIC_LABELS.values())
    assert table.cell(0, 1).text == "Dez/20"
    assert str(table.cell(0, 0).fill.fore_color.rgb) == "EC7000"
    assert table.cell(1, 1).text == "13,3%"
    assert table.cell(1, 6).text_frame.paragraphs[0].font.bold
    for slide in prs.slides:
        texts = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
        assert "1 instituição" in texts and "Dez/20 a Dez/25" in texts
        assert "BCB IFData" in texts and "09/10/2026 18:30 -03" in texts
        assert "Tipo: Evolução histórica" in texts
        assert str(slide.background.fill.fore_color.rgb) == "FFFFFF"
        assert all(s.left >= 0 and s.top >= 0 and s.left + s.width <= prs.slide_width
                   and s.top + s.height <= prs.slide_height for s in slide.shapes)
        assert "ITAU - PRUDENCIAL" in slide.notes_slide.notes_text_frame.text
    with ZipFile(BytesIO(blob)) as archive:
        assert not any(n.startswith("ppt/media/") for n in archive.namelist())
        assert any(n.startswith("ppt/embeddings/") for n in archive.namelist())
        xml = etree.fromstring(archive.read("ppt/charts/chart1.xml"))
        assert len(xml.xpath(".//c:barChart/c:ser", namespaces=NS)) == 2
        assert len(xml.xpath(".//c:lineChart/c:ser", namespaces=NS)) == 2
        assert len(xml.xpath(".//c:valAx", namespaces=NS)) == 2
        assert float(xml.xpath(".//c:plotArea/c:layout/c:manualLayout/c:w/@val", namespaces=NS)[0]) >= .86
        assert xml.xpath(".//c:plotArea/c:layout/c:manualLayout/c:layoutTarget/@val", namespaces=NS) == ["inner"]
        assert xml.xpath(".//c:legend/c:legendEntry/c:txPr//a:srgbClr/@val", namespaces=NS) == [s[2].lstrip("#").upper() for s in SERIES]
        assert all(0 <= int(value) < 2 ** 31 for value in xml.xpath(".//c:axId/@val | .//c:crossAx/@val", namespaces=NS))
        assert len(xml.xpath(".//c:dispBlanksAs[@val='gap']", namespaces=NS)) == 1
        assert all(node.get("typeface") == "Calibri" for node in xml.xpath(".//a:latin", namespaces=NS))


def test_last_valid_labels_are_dynamic_bold_colored_and_separated_with_gaps():
    args = sample()
    args["df_graph"].loc[5, "Core Funding*"] = None
    args["df_graph"].loc[2, "Carteira de Crédito*"] = None
    before_graph, before_table = args["df_graph"].copy(), args["df_show"].copy()
    blob = export_evolucao_powerpoint(**args)
    pd.testing.assert_frame_equal(args["df_graph"], before_graph)
    pd.testing.assert_frame_equal(args["df_show"], before_table)
    with ZipFile(BytesIO(blob)) as archive:
        xml = etree.fromstring(archive.read("ppt/charts/chart1.xml"))
        positions = []
        for i, series in enumerate(xml.xpath(".//c:barChart/c:ser | .//c:lineChart/c:ser", namespaces=NS)):
            labels = series.xpath("c:dLbls/c:dLbl", namespaces=NS)
            assert len(labels) == 1
            label = labels[0]
            assert label.xpath("c:idx/@val", namespaces=NS) == ["4" if i == 3 else "5"]
            assert label.xpath("c:showVal/@val", namespaces=NS) == ["1"]
            assert label.xpath("c:showSerName/@val", namespaces=NS) == ["0"]
            assert len(label.xpath("c:spPr", namespaces=NS)) == 1
            assert label.xpath(".//a:defRPr[@b='1']", namespaces=NS)
            assert SERIES[i][2].lstrip("#").upper() in label.xpath(".//a:srgbClr/@val", namespaces=NS)
            positions.append(float(label.xpath("c:layout/c:manualLayout/c:y/@val", namespaces=NS)[0]))
            x = float(label.xpath("c:layout/c:manualLayout/c:x/@val", namespaces=NS)[0])
            assert PLOT_LEFT < x < PLOT_LEFT + PLOT_WIDTH - .1
        assert min(b - a for a, b in zip(sorted(positions), sorted(positions)[1:])) >= .139
        assert min(positions) >= PLOT_TOP
        assert max(positions) + LABEL_HEIGHT <= PLOT_TOP + PLOT_HEIGHT + .0001
        assert labels[0].xpath("c:showCatName/@val", namespaces=NS) == ["1"]
        workbook_name = next(n for n in archive.namelist() if n.startswith("ppt/embeddings/"))
        sheet = load_workbook(BytesIO(archive.read(workbook_name)), data_only=True).active
        assert sheet.cell(7, 5).value is None
        assert sheet.cell(4, 4).value is None
        assert sheet.cell(6, 5).value == 20
        assert sheet.cell(2, 2).value == 1


def test_long_windows_keep_every_period_and_source_limitations_visible():
    args = sample(11)
    args["df_show"].loc[2, "dez-28"] = "-†"
    args["status_rows"] = [{"Período": "dez-28", "Indicador": "Core Funding*",
                            "Status analítico": "fallback_capitacoes_view", "Fonte analítica": "Captações"}]
    prs = Presentation(BytesIO(export_evolucao_powerpoint(**args)))
    assert len(prs.slides) == 2
    categories = []
    for slide in prs.slides:
        chart = next(s.chart for s in slide.shapes if s.has_chart)
        categories.extend(c.label for c in chart.plots[0].categories)
    assert categories == [f"Dez/{20 + i}" for i in range(11)]
    assert not any("fallback para Captações" in s.text for s in prs.slides[0].shapes if s.has_text_frame)
    assert any("fallback para Captações: Dez/28" in s.text for s in prs.slides[1].shapes if s.has_text_frame)
    table = next(s.table for s in prs.slides[1].shapes if s.has_table)
    assert table.cell(3, 1).text == "N/D†"
    assert "Captações" in prs.slides[1].notes_slide.notes_text_frame.text


@pytest.mark.parametrize("value", [0, -1e9, 1e9])
def test_coincident_labels_leave_room_for_dates_and_axes(value):
    args = sample(8)
    args["df_graph"].iloc[:, :] = value
    args["df_graph"].iloc[-1, :] = None
    chart = next(s.chart for s in Presentation(BytesIO(export_evolucao_powerpoint(**args))).slides[0].shapes if s.has_chart)
    labels = chart._chartSpace.xpath(".//c:dLbl")
    positions = sorted(float(label.xpath("./c:layout/c:manualLayout/c:y/@val")[0]) for label in labels)
    assert all(label.xpath("./c:showCatName/@val") == ["1"] for label in labels)
    # Two lines (value + valid date) at 12 pt, with at least 0.05 in of air.
    assert min(b - a for a, b in zip(positions, positions[1:])) * CHART_HEIGHT >= 24 / 72 + .05
    assert positions[0] >= PLOT_TOP
    assert positions[-1] + LABEL_HEIGHT <= PLOT_TOP + PLOT_HEIGHT + .0001


def test_zero_negative_and_all_missing_series_preserve_numeric_meaning():
    args = sample(1)
    args["df_graph"].loc[0] = [-1e9, 0, None, None]
    prs = Presentation(BytesIO(export_evolucao_powerpoint(**args)))
    chart = next(s.chart for s in prs.slides[0].shapes if s.has_chart)
    assert [tuple(s.values) for s in chart.series] == [(-1,), (0,), (None,), (None,)]
    assert float(chart._chartSpace.xpath(".//c:valAx[c:axPos/@val='l']/c:scaling/c:min/@val")[0]) < 0
    assert chart.series[0]._element.xpath(".//c:dLbl")
    assert chart.series[1]._element.xpath(".//c:dLbl")
    assert not chart.series[2]._element.xpath(".//c:dLbl")


def test_mismatched_windows_are_rejected():
    args = sample()
    args["periods"] = args["periods"][:-1]
    with pytest.raises(ValueError, match="mesmas competências"):
        export_evolucao_powerpoint(**args)
