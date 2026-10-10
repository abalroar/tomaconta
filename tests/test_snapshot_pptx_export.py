from io import BytesIO
import json
from zipfile import ZipFile

import pandas as pd
import pytest
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation

from tabs.carteira_4966 import EXPECTED_LOSS_COLUMNS, build_carteira_4966_model
from utils.snapshot_pptx_export import export_snapshot_powerpoint


def snapshot():
    labels = [
        "Ativo Total", "Carteira de Crédito", "Patrimônio Líquido", "Índice de Basileia",
        "Lucro Líquido Trimestral", "Lucro Líquido Acum. YTD", "ROE trim. anualizado", "ROE Ac. Anualizado",
        "Crédito / Captações", "Desp. Anualizada Captação / Volume Captação", "Perda Esperada / Estágio 3",
        "Perda Esperada / Carteira", "CET1",
    ]
    return {
        "bank": "ITAU - PRUDENCIAL", "base": "Consolidada / Prudencial", "period_label": "Mar/26",
        "qoq_label": "Dez/25", "yoy_label": "Mar/25", "source_notes": ["Dados trimestrais. Perda Esperada / Estágio 3: cobertura aproximada."],
        "cards": [{"key": str(index), "label": label, "value": "14,77%" if index == 3 else f"{index},25",
                   "qoq": {"display": "↓ −41 bps" if index == 3 else "↑ +7 bps", "tone": "attention", "reason": "Fonte preservada"},
                   "yoy": {"display": "↑ +0,4 p.p.", "tone": "favorable"},
                   "group": "hero" if index < 4 else "profit" if index < 8 else "support",
                   "source": "BCB IFData Rel. 1", "status": "available",
                   "history": {"periods": ["Jun/25", "Set/25", "Dez/25", "Mar/26"], "values": [1e9, None, 1.1e9, 1.2e9], "kind": "bars" if 4 <= index < 8 else "line"}}
                  for index, label in enumerate(labels)],
    }


def carteira():
    source = pd.DataFrame({
        "Período": ["3/2025", "4/2025", "1/2026"], "Total Geral": [100e6, 100e6, 110e6],
        "C1": [10e6, 10e6, 11e6], "C2": [20e6, 20e6, 21e6], "C3": [15e6, 15e6, 17e6],
        "C4": [20e6, 20e6, 21e6], "C5": [30e6, 30e6, 33e6], "Total não Individualizado": [3e6, 3e6, 4e6],
        "Carteira não Informada ou não se Aplica": [2e6, 2e6, 3e6], "Total Exterior": [0, 0, 0],
        "Inadimplência": [2e6, 2.18e6, 110e6 * .0225],
    })
    loss = pd.DataFrame({"Período": source["Período"]})
    loss[EXPECTED_LOSS_COLUMNS[0]] = [-4e6, -2.18e6 * 1.931, -110e6 * .0225 * 1.935]
    for column in EXPECTED_LOSS_COLUMNS[1:]:
        loss[column] = 0.0
    return build_carteira_4966_model(source, loss, source["Período"].tolist())


def peers():
    from utils.peers_table_model import DEFAULT_METRICS, build_query
    periods = ["3/2025", "4/2025", "1/2026"]
    df = pd.DataFrame({"Instituição": ["ITAU - PRUDENCIAL"] * 3, "Período": periods})
    return build_query(df, ["ITAU - PRUDENCIAL"], periods, DEFAULT_METRICS, {}, {},
                       base="Consolidada / Prudencial", cache_token="qa", scale="R$ bilhões", mode="quarter", queried_at="10/10/2026")


def visible_text(slide):
    pieces = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            pieces.append(shape.text)
        elif shape.has_table:
            pieces.extend(cell.text for row in shape.table.rows for cell in row.cells if not cell.is_spanned)
    return "\n".join(pieces)


def assert_schema_order(parent, sequence):
    """Verifica a sequência OOXML, que um ZIP bem formado não valida."""
    names = [etree.QName(child).localname for child in parent]
    assert all(name in sequence for name in names), names
    positions = [sequence.index(name) for name in names]
    assert positions == sorted(positions), names
    assert len(names) == len(set(names)), names


def test_office_chart_and_table_properties_follow_ooxml_child_order():
    # A ordem inválida de spPr levou o PowerPoint a remover os slides 1 e 2.
    # Bordas após solidFill também violam a sequência de tcPr.
    with ZipFile(BytesIO(export_snapshot_powerpoint(snapshot(), peers(), carteira()))) as archive:
        for path in archive.namelist():
            if path.startswith("ppt/charts/chart") and path.endswith(".xml"):
                root = etree.fromstring(archive.read(path))
                assert_schema_order(root, [
                    "date1904", "lang", "roundedCorners", "style", "clrMapOvr", "pivotSource",
                    "protection", "chart", "spPr", "txPr", "externalData", "printSettings", "userShapes", "extLst",
                ])
                ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
                chart = root.find("c:chart", ns)
                assert_schema_order(chart, [
                    "title", "autoTitleDeleted", "pivotFmts", "view3D", "floor", "sideWall", "backWall",
                    "plotArea", "legend", "plotVisOnly", "dispBlanksAs", "showDLblsOverMax", "extLst",
                ])
                for axis in root.xpath(".//c:catAx | .//c:valAx", namespaces=ns):
                    assert [etree.QName(child).localname for child in axis][:4] == ["axId", "scaling", "delete", "axPos"]
                plot_area = chart.find("c:plotArea", ns)
                assert_schema_order(plot_area, ["layout", "lineChart", "barChart", "catAx", "valAx", "dTable", "spPr", "extLst"])
            if path.startswith("ppt/slides/slide") and path.endswith(".xml"):
                root = etree.fromstring(archive.read(path))
                for props in root.xpath(".//a:tcPr", namespaces={"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}):
                    assert_schema_order(props, [
                        "lnL", "lnR", "lnT", "lnB", "lnTlToBr", "lnBlToTr", "cell3D",
                        "noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill", "headers", "extLst",
                    ])


def test_first_two_slides_preserve_snapshot_and_have_native_editable_history_with_gaps():
    payload = snapshot()
    result = export_snapshot_powerpoint(payload, peers(), carteira())
    prs = Presentation(BytesIO(result))
    texts = [visible_text(slide) for slide in prs.slides]
    assert len(prs.slides) >= 6
    assert sum(shape.has_table for slide in list(prs.slides)[:2] for shape in slide.shapes) == 2
    for card in payload["cards"]:
        target = 1 if card["group"] == "support" else 0
        assert card["label"] in texts[target]
        assert card["value"] in texts[target]
    assert "14,77%" in texts[0] and "↓ −41 bps" in texts[0]
    assert "QoQ\nDez/25" in texts[0] and "YoY\nMar/25" in texts[0]
    charts = [shape.chart for slide in list(prs.slides)[:2] for shape in slide.shapes if shape.has_chart]
    assert len(charts) == 13
    assert charts[0].series[0].values == (1e9, None, 1.1e9, 1.2e9)
    workbook = load_workbook(BytesIO(charts[0].part.chart_workbook.xlsx_part.blob), data_only=True)
    assert list(workbook.active.values)[2] == ("Set/25", None)
    with ZipFile(BytesIO(result)) as archive:
        assert not any(path.startswith("ppt/media/") for path in archive.namelist())
        assert len([path for path in archive.namelist() if path.startswith("ppt/embeddings/")]) == 13
        chart_xml = etree.fromstring(archive.read("ppt/charts/chart1.xml"))
        assert chart_xml.xpath("string(.//c:dispBlanksAs/@val)", namespaces={"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}) == "gap"
        namespace = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        for path in archive.namelist():
            if not path.startswith("ppt/charts/chart") or not path.endswith(".xml"):
                continue
            chart_xml = etree.fromstring(archive.read(path))
            identifiers = chart_xml.xpath(".//c:axId/@val | .//c:crossAx/@val", namespaces=namespace)
            # IDs UInt32 acima de Int32 causam reparo no PowerPoint Mac.
            assert all(0 < int(identifier) <= 0x7FFFFFFF for identifier in identifiers)
            axes = chart_xml.xpath(".//c:catAx | .//c:valAx", namespaces=namespace)
            axis_ids = [axis.xpath("string(c:axId/@val)", namespaces=namespace) for axis in axes]
            assert len(axis_ids) == len(set(axis_ids)) == 2
            for axis in axes:
                own_id = axis.xpath("string(c:axId/@val)", namespaces=namespace)
                cross_id = axis.xpath("string(c:crossAx/@val)", namespaces=namespace)
                assert cross_id in axis_ids and cross_id != own_id
            plot_ids = chart_xml.xpath(".//c:lineChart/c:axId/@val | .//c:barChart/c:axId/@val", namespaces=namespace)
            assert set(plot_ids) == set(axis_ids)
    notes = json.loads(prs.slides[0].notes_slide.notes_text_frame.text)
    assert notes["cards"][0]["qoq"]["reason"] == "Fonte preservada"
    assert "Tabela de Peers" in texts[2]
    assert "Carteira 4.966" in texts[-1]


def test_4966_uses_shared_difference_units_and_credit_colors_in_native_table():
    prs = Presentation(BytesIO(export_snapshot_powerpoint(snapshot(), peers(), carteira())))
    table = next(shape.table for shape in prs.slides[-1].shapes if shape.has_table)
    risk = next(row for row in table.rows if "conceito de arrasto" in row.cells[0].text)
    # Final period has two subcolumns, saldo and carteira %. Rate moves 2.18 -> 2.25.
    assert "2,25%" in risk.cells[len(risk.cells) - 1].text
    assert "↑ +7 bps" in risk.cells[len(risk.cells) - 1].text
    assert str(risk.cells[len(risk.cells) - 1].text_frame.paragraphs[1].font.color.rgb) == "B32624"
    coverage = next(row for row in table.rows if row.cells[0].text == "PDD / Créditos vencidos acima de 90 dias (%)")
    assert "193,50%" in coverage.cells[len(coverage.cells) - 2].text
    assert "↑ +0,4 p.p." in coverage.cells[len(coverage.cells) - 2].text
    assert str(coverage.cells[len(coverage.cells) - 2].text_frame.paragraphs[1].font.color.rgb) == "16713B"
    assert "bps inteiros" in visible_text(prs.slides[-1])


def test_missing_sections_keep_order_and_explicit_scope_reason_without_fabricating_values():
    prs = Presentation(BytesIO(export_snapshot_powerpoint(snapshot(), peers_unavailable_reason="IFData individual não disponível.", carteira_unavailable_reason="Relatório 16 exige o perímetro prudencial.")))
    assert len(prs.slides) == 4
    assert "Tabela de Peers" in visible_text(prs.slides[2])
    assert "IFData individual não disponível." in visible_text(prs.slides[2])
    assert "Carteira 4.966" in visible_text(prs.slides[3])
    assert "Relatório 16 exige o perímetro prudencial." in visible_text(prs.slides[3])
    assert "N/D" in visible_text(prs.slides[3])


def test_last_three_periods_are_chronological_and_short_coverage_is_disclosed():
    model = carteira()
    model.periods = ("1/2026",)
    prs = Presentation(BytesIO(export_snapshot_powerpoint(snapshot(), peers(), model)))
    final = visible_text(prs.slides[-1])
    assert "Mar/26\nQoQ vs Dez/25" in final
    assert "1 período(s) disponível(is)" in final
    assert "Set/25\nQoQ vs" not in final


def test_history_mismatched_periods_raise_a_clear_input_error():
    payload = snapshot()
    payload["cards"][0]["history"]["periods"].pop()
    with pytest.raises(ValueError, match="um rótulo por valor"):
        export_snapshot_powerpoint(payload)
