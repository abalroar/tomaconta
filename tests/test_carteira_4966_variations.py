from io import BytesIO

import openpyxl
import pandas as pd
import pytest

from tabs.carteira_4966 import (
    EXPECTED_LOSS_COLUMNS, ROW_SPECS, ROW_BY_KEY, MetricCell,
    build_carteira_4966_model, cell_delta, comparison_source_periods,
    build_carteira_4966_excel, build_carteira_4966_raw_excel,
    render_carteira_4966_html, variation_dataframe,
)


def sources():
    portfolio = pd.DataFrame({
        "Período": ["4/2025", "1/2026"], "Total Geral": [100e6, 110e6],
        "C1": [10e6, 11e6], "C2": [20e6, 21e6], "C3": [15e6, 17e6],
        "C4": [20e6, 21e6], "C5": [30e6, 33e6],
        "Total não Individualizado": [3e6, 4e6],
        "Carteira não Informada ou não se Aplica": [2e6, 3e6],
        "Total Exterior": [0, 0], "Inadimplência": [2.18e6, 110e6 * .0225],
    })
    loss = pd.DataFrame({"Período": ["4/2025", "1/2026"]})
    loss[EXPECTED_LOSS_COLUMNS[0]] = [-2.18e6 * 1.931, -110e6 * .0225 * 1.935]
    for column in EXPECTED_LOSS_COLUMNS[1:]:
        loss[column] = 0.0
    return portfolio, loss


def model(periods=("1/2026",)):
    portfolio, loss = sources()
    return build_carteira_4966_model(portfolio, loss, periods)


def test_percent_deltas_subtract_and_coverage_uses_pp_with_credit_colors():
    result = model()
    risk = cell_delta(result, ROW_BY_KEY["delinquency"], "1/2026", secondary=True)
    assert risk.value == pytest.approx(7)
    assert risk.display == "↑ +7 bps"
    assert risk.tone == "attention"
    coverage = cell_delta(result, ROW_BY_KEY["provision_over_delinquency"], "1/2026")
    assert coverage.value == pytest.approx(.4)
    assert coverage.display == "↑ +0,4 p.p."
    assert coverage.tone == "favorable"
    result.cells["provision_over_delinquency"]["1/2026"] = MetricCell(1.9)
    falling = cell_delta(result, ROW_BY_KEY["provision_over_delinquency"], "1/2026")
    assert falling.display == "↓ −3,1 p.p."
    assert falling.tone == "attention"


@pytest.mark.parametrize("spec", ROW_SPECS, ids=lambda row: row.key)
def test_every_row_uses_exact_reference_and_appropriate_operation(spec):
    result = model()
    for secondary in ([False, True] if spec.layout == "paired" else [False]):
        delta = cell_delta(result, spec, "1/2026", secondary=secondary)
        assert delta.reference_period == "4/2025"
        assert delta.reference is not None
        percent = secondary or spec.layout == "percent_span"
        factor = 10000 if percent and spec.key in {"delinquency", "provision_over_portfolio"} else 100
        expected = (delta.current - delta.reference) * factor if percent else (delta.current - delta.reference) / delta.reference * 100
        assert delta.value == pytest.approx(expected)
        if not percent or spec.group == "classification" or spec.key == "provision_over_portfolio":
            assert delta.tone == "neutral"


def test_hidden_reference_keeps_selected_base_and_does_not_bridge_missing_quarter():
    result = model()
    assert result.base_period == "1/2026"
    assert result.reference_cells["c1"]["4/2025"].secondary == pytest.approx(10 / 110)
    assert result.cells["c1"]["1/2026"].secondary == pytest.approx(11 / 110)
    assert cell_delta(result, ROW_BY_KEY["c1"], "1/2026", secondary=True).value == pytest.approx(100 / 110)
    portfolio, loss = sources()
    portfolio.loc[portfolio["Período"] == "4/2025", "Período"] = "3/2025"
    gap = build_carteira_4966_model(portfolio, loss, ["1/2026"])
    delta = cell_delta(gap, ROW_BY_KEY["delinquency"], "1/2026", secondary=True)
    assert delta.value is None
    assert delta.display == "Base N/D"
    assert comparison_source_periods(["1/2026"]) == ("4/2025", "1/2026")


def test_bad_reference_preserves_delta_but_neutralizes_color_and_identifies_issue():
    portfolio, loss = sources()
    loss.loc[0, EXPECTED_LOSS_COLUMNS[0]] *= -1
    result = build_carteira_4966_model(portfolio, loss, ["1/2026"])
    assert result.cell_quality_issues("provision_over_delinquency", "1/2026") == ()
    assert result.cell_quality_issues("provision_over_delinquency", "4/2025")
    delta = cell_delta(result, ROW_BY_KEY["provision_over_delinquency"], "1/2026")
    assert delta.value == pytest.approx(.4)
    assert delta.display.endswith("*")
    assert delta.tone == "neutral"
    assert "Dez/25" in delta.reason
    assert "sinal" in delta.reason


def test_zero_reference_is_valid_for_percent_subtraction_but_blocks_amount_growth():
    result = model()
    result.reference_cells["delinquency"]["4/2025"] = MetricCell(0, 0)
    delta = cell_delta(result, ROW_BY_KEY["delinquency"], "1/2026", secondary=True)
    assert delta.value == pytest.approx(225)
    assert cell_delta(result, ROW_BY_KEY["delinquency"], "1/2026").display == "Base ≤ 0"
    result.cells["delinquency"]["1/2026"] = MetricCell(0, .0000001)
    assert cell_delta(result, ROW_BY_KEY["delinquency"], "1/2026", secondary=True).display == "↑ <1 bp"


def test_2025_transition_is_blocked_even_when_older_input_exists():
    portfolio, loss = sources()
    portfolio["Período"] = ["4/2024", "1/2025"]
    loss["Período"] = portfolio["Período"]
    result = build_carteira_4966_model(portfolio, loss, ["1/2025"])
    assert cell_delta(result, ROW_BY_KEY["delinquency"], "1/2025", secondary=True).display == "Quebra em 2025"


def test_html_and_both_excel_exports_share_units_numbers_and_colors():
    result = model()
    html = render_carteira_4966_html(result)
    assert 'class="tc-4966-delta attention"' in html
    assert 'class="tc-4966-delta favorable"' in html
    assert "↑ +7 bps" in html
    assert "↑ +0,4 p.p." in html
    portfolio, loss = sources()
    for payload in (build_carteira_4966_excel(result), build_carteira_4966_raw_excel(result, portfolio, loss)):
        workbook = openpyxl.load_workbook(BytesIO(payload))
        sheet = workbook["Variações"]
        risk = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == ROW_BY_KEY["delinquency"].label and row[1].value == "%")
        coverage = next(row for row in sheet.iter_rows(min_row=2) if row[0].value == ROW_BY_KEY["provision_over_delinquency"].label)
        assert risk[6].value == pytest.approx(7)
        assert risk[6].number_format == "0"
        assert risk[7].value == "bps"
        assert risk[8].value == "↑ +7 bps"
        assert risk[8].font.color.rgb == "FFB32624"
        assert coverage[6].value == pytest.approx(.4)
        assert coverage[7].value == "p.p."
        assert coverage[8].font.color.rgb == "FF16713B"
        assert risk[4].value == pytest.approx(.0225)
        assert risk[4].number_format == "0.00%"
    visual = openpyxl.load_workbook(BytesIO(build_carteira_4966_excel(result)))["Modelo 4966"]
    risk_row = next(row[0].row for row in visual if row[0].value == ROW_BY_KEY["delinquency"].label)
    assert visual.cell(risk_row + 1, 3).value == "↑ +7 bps"
    assert visual.cell(risk_row + 1, 3).font.color.rgb == "FFB32624"
    assert visual.cell(risk_row, 3).value == pytest.approx(.0225)
    assert variation_dataframe(result).query("Componente == '%'")["Delta numérico"].notna().all()
