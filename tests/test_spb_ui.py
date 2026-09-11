from __future__ import annotations

import textwrap
from io import BytesIO
from pathlib import Path

import pandas as pd
from pptx import Presentation
from streamlit.testing.v1 import AppTest


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def _spb_app() -> AppTest:
    source = (PROJECT_ROOT / "app1.py").read_text()
    start = source.index('elif menu == "Meios de Pagamento (SPB)":')
    end = source.index('elif menu == "Estatísticas Crédito BC":', start)
    body = textwrap.dedent(source[start:end].split("\n", 1)[1])
    setup = f"""
import sys
sys.path.insert(0, {str(PROJECT_ROOT)!r})
import pandas as pd
import plotly.express as px
import streamlit as _streamlit
from types import SimpleNamespace
# AppTest's ButtonGroup cannot serialize single-selection segmented controls in this
# Streamlit version. Radio uses the same single-value widget/session-state contract.
class _SpbTestStreamlit:
    def __getattr__(self, name):
        return getattr(_streamlit, name)
    def segmented_control(self, label, options, default, key):
        return _streamlit.radio(label, options, index=options.index(default), key=key)
st = _SpbTestStreamlit()
from utils.spb_meios_pagamento_viz import *
from utils.spb_meios_pagamento_viz import (
    ITAU_BBA_PALETTE as SPB_ITAU_BBA_PALETTE,
    MONTHLY_DEFAULT_INSTRUMENTS as SPB_MONTHLY_DEFAULT_INSTRUMENTS,
    QUARTERLY_DEFAULT_INSTRUMENTS as SPB_QUARTERLY_DEFAULT_INSTRUMENTS,
    available_instruments as spb_available_instruments,
    build_line_figure as spb_build_line_figure,
    build_share_figure as spb_build_share_figure,
    default_start_index as spb_default_start_index,
    filter_period_range as spb_filter_period_range,
    format_ano_mes_label as spb_format_ano_mes_label,
    format_trimestre_label as spb_format_trimestre_label,
    latest_summary as spb_latest_summary,
    period_label_map as spb_period_label_map,
    period_options as spb_period_options,
)
samples = {{
    "nucleo_mensal": pd.DataFrame({{"ano_mes": ["202501", "202502", "202503"], "quantidade_pix": [10,20,30], "quantidade_ted": [4,5,6], "valor_pix": [100,200,300], "valor_ted": [40,50,60]}}),
    "nucleo_trimestral": pd.DataFrame({{"trimestre": ["20243", "20244", "20251", "20252"], "quantidade_pix": [10,20,30,40], "quantidade_ted": [4,5,6,7], "valor_pix": [100,200,300,400], "valor_ted": [40,50,60,70]}}),
    "intercambio": pd.DataFrame({{"trimestre": ["20251", "20252"], "funcao_cartao": ["Crédito", "Crédito"], "tarifa_intercambio_ponderada": [1.2,1.3]}}),
    "desconto": pd.DataFrame({{"trimestre": ["20251", "20252"], "funcao_cartao": ["Crédito", "Crédito"], "tx_media_desconto": [2.2,2.3]}}),
    "canais_servicos": pd.DataFrame({{"trimestre": ["20251", "20252"], "operacao": ["Celular", "Celular"], "qtd_transacoes": [100,200]}}),
    "canais_transacoes": pd.DataFrame({{"trimestre": ["20251", "20252"], "canal_acesso": ["Celular", "Celular"], "produto": ["Pix", "Pix"], "qtd_transacoes": [100,200]}}),
    "cartoes": pd.DataFrame({{"trimestre": ["20251"] * 510 + ["20252"] * 10, "valor": list(range(520))}}),
}}
def result(key):
    return SimpleNamespace(sucesso=True, dados=samples.get(key, pd.DataFrame()).copy(deep=True))
cache = SimpleNamespace(carregar=lambda: result("nucleo_trimestral"), carregar_dataset=result)
def get_cache_manager():
    return SimpleNamespace(get_cache=lambda key: cache)
def _cache_version_token(key):
    return "fixture-publication"
"""
    return AppTest.from_string(setup + "\n" + body, default_timeout=15).run()


def test_export_preserves_each_section_filter_and_persists_download():
    app = _spb_app()
    assert not app.exception
    app.selectbox(key="spb_nucleo_trimestral_ini").set_value(pd.Timestamp("2024-12-01"))
    app.multiselect(key="spb_nucleo_trimestral_instrumentos").set_value(["TED"])
    app.run()
    app.radio(key="spb_nucleo_periodicidade").set_value("Mensal").run()
    assert not app.exception
    app.selectbox(key="spb_nucleo_mensal_ini").set_value(pd.Timestamp("2025-02-01"))
    app.multiselect(key="spb_nucleo_mensal_instrumentos").set_value(["Pix"])
    app.run()
    app.radio(key="spb_nucleo_periodicidade").set_value("Trimestral").run()
    assert app.multiselect(key="spb_nucleo_trimestral_instrumentos").value == ["TED"]
    app.button(key="spb_prepare_ppt").click().run()
    assert not app.exception
    payload = app.session_state["spb_prepared_ppt"]["data"]
    presentation = Presentation(BytesIO(payload))
    charts = [shape.chart for slide in presentation.slides for shape in slide.shapes if shape.has_chart]
    assert [series.name for series in charts[0].series] == ["Pix"]
    assert [category.label for category in charts[0].plots[0].categories] == ["02-2025", "03-2025"]
    assert [series.name for series in charts[2].series] == ["TED"]
    assert [category.label for category in charts[2].plots[0].categories] == ["dez-24", "mar-25", "jun-25"]
    assert {series.name for series in charts[4].series} == {"Pix", "TED"}
    assert len(charts[4].plots[0].categories) == 4
    app.selectbox(key="spb_outros_cartoes_pagina").set_value(2).run()
    assert not app.exception
    assert app.session_state["spb_prepared_ppt"]["data"] == payload
    assert len(app.get("download_button")) == 1
    other_table = next(frame.value for frame in app.dataframe if "valor" in frame.value.columns)
    assert other_table["valor"].tolist() == list(range(500, 520))
    app.selectbox(key="spb_outros_cartoes_periodo").set_value("20252").run()
    assert not app.exception
    other_table = next(frame.value for frame in app.dataframe if "valor" in frame.value.columns)
    assert other_table["valor"].tolist() == list(range(510, 520))
    app.multiselect(key="spb_participacao_instrumentos").set_value(["Pix"]).run()
    assert not app.exception
    assert len(app.get("download_button")) == 0
    assert any("Prepare o PPTX novamente" in info.value for info in app.info)
