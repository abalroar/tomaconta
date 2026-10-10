from io import BytesIO
import base64
import ast
import json
from pathlib import Path
from dataclasses import replace
from unittest.mock import Mock
from types import SimpleNamespace
from zipfile import ZipFile

import pandas as pd
import pytest
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation

from utils.peers_table_model import BY_KEY, DEFAULT_METRICS, ARRASTO_ROWS, build_query, delta, reference, get_metric, methodology_rows, variation_tone
from utils.peers_table_exports import export_excel, export_powerpoint
from utils import peers_groups
from tabs.peers_table import table_html, calculation_rows, variation_rows, table_row_label, table_footnote
from utils.formatting import formatar_delta_br


def sample(banks=("A", "B", "C"), periods=("4/2025", "1/2026", "2/2026"), metrics=("Ativo Total", "Custo de Crédito (%)", "Lucro Líquido Acumulado"), **overrides):
    rows, values = [], {}
    for b, bank in enumerate(banks):
        for p in (*periods, "2/2025", "4/2024", "1/2025"):
            rows.append({"Instituição": bank, "Período": p, "Ativo Total": 1e9*(b+1), "Lucro Líquido Acumulado YTD": 1e8, "Trace::Custo de Crédito::PDD Crédito Anualizada": -2e7, "Carteira de Crédito Bruta": 1e9})
            values.update({("Ativo Total", bank, p): 1e9*(b+1), ("Custo de Crédito (%)", bank, p): .02, ("Lucro Líquido Acumulado", bank, p): 1e8})
    df = pd.DataFrame(rows).drop_duplicates(["Instituição", "Período"])
    kwargs = dict(base="Consolidada / Prudencial", cache_token="abc", scale="R$ bilhões", mode="year", queried_at="09/10/2026")
    kwargs.update(overrides)
    return build_query(df, list(banks), list(periods), list(metrics), values, {}, **kwargs), df, values


def test_delta_baseline_units_rounding_and_ytd_guard():
    assert reference("1/2026", "quarter") == "4/2025"
    assert reference("2/2026", "year") == "2/2025"
    assert delta(.030843, .030776, BY_KEY["Custo de Crédito (%)"], "year") == ("up", "↑ +1 bps")
    assert delta(.03, .030001, BY_KEY["Custo de Crédito (%)"], "year") == ("down", "↓ <1 bp")
    assert delta(3, -2, BY_KEY["Ativo Total"], "year")[0] is None
    assert delta(3, 2, BY_KEY["Lucro Líquido Acumulado"], "quarter")[1] == "YTD: janelas diferentes"
    assert formatar_delta_br(11.34, "bps", 0) == "+11 bps"
    assert formatar_delta_br(-11.34, "bps", 0) == "−11 bps"
    assert delta(.03005, .03, BY_KEY["Custo de Crédito (%)"], "year") == ("up", "↑ +1 bps")
    assert delta(.03, .03, BY_KEY["Custo de Crédito (%)"], "year") == ("flat", "= 0 bps")


@pytest.mark.parametrize("metric", [m for m in BY_KEY.values() if m.unit == "%"], ids=lambda m: m.key)
def test_every_percentage_metric_uses_subtraction_in_the_declared_unit(metric):
    for mode in ("quarter", "year"):
        if metric.delta_unit == "p.p.":
            assert delta(.0225,.0218,metric,mode) == ("up", "↑ +0,07 p.p.")
            assert delta(.1477,.1518,metric,mode) == ("down", "↓ −0,41 p.p.")
            assert delta(1.8741,1.9354,metric,mode) == ("down", "↓ −6,13 p.p.")
        else:
            assert delta(.0225,.0218,metric,mode) == ("up", "↑ +7 bps")
            assert delta(.1477,.1518,metric,mode) == ("down", "↓ −41 bps")
            assert delta(1.8741,1.9354,metric,mode) == ("down", "↓ −613 bps")


@pytest.mark.parametrize("key", ["PDD / Inadimplência (arrasto)", "Perda Esperada / Estágio 3", "Perda Esperada / Est2+3", "Custo de Crédito / Receita de Crédito (%)", "Perda Esperada / Carteira de Crédito*"])
def test_coverage_and_expense_shares_use_percentage_point_differences(key):
    metric = BY_KEY[key]
    assert metric.delta_unit == "p.p."
    assert delta(1.935, 1.931, metric, "quarter") == ("up", "↑ +0,40 p.p.")
    assert delta(1.935, 1.971, metric, "year") == ("down", "↓ −3,60 p.p.")


def test_color_depends_on_credit_meaning_and_data_quality():
    assert variation_tone(BY_KEY["Ativo Total"], "up") == "neutral"
    assert variation_tone(BY_KEY["Inadimplência"], "down") == "favorable"
    assert variation_tone(BY_KEY["Custo de Crédito (%)"], "up") == "attention"
    assert variation_tone(BY_KEY["Inadimplência / Carteira Total"], "down") == "favorable"
    assert variation_tone(BY_KEY["Índice de Basileia Total (%)"], "down") == "attention"
    assert variation_tone(BY_KEY["Índice de Basileia Total (%)"], "down", "curated_value") == "attention"
    assert variation_tone(BY_KEY["PDD / Inadimplência (arrasto)"], "up") == "favorable"
    assert variation_tone(BY_KEY["PDD / Inadimplência (arrasto)"], "up", "warning") == "neutral"


ECONOMIC_DELTA_CASES = {
    # Saldos de balanço e resultados precisam de contexto; crescer não implica menor risco.
    **{key: (110, 100, "↑ +10,00 %", "neutral") for key in (
        "Ativo Total", "Ativos Líquidos", "Carteira de Crédito*", "Perda Esperada",
        "Depósitos Totais", "Core Funding*", "Patrimônio Líquido (PL)", "Lucro Líquido Acumulado",
    )},
    # Maior estoque de crédito deteriorado merece atenção, mesmo com carteira crescente.
    **{key: (110, 100, "↑ +10,00 %", "attention") for key in (
        "Inadimplência", "Ativos Estágio 2", "Ativos Estágio 3",
    )},
    **{key: (.0225, .0218, "↑ +7 bps", "attention") for key in (
        "Custo de Crédito (%)", "Ativos Problemáticos / Carteira Total",
        "Inadimplência / Carteira Total", "Ativos Estágio 3 / Carteira de Crédito",
        "Inadimplência / Carteira de Crédito",
    )},
    **{key: (.1477, .1518, "↓ −41 bps", "attention") for key in (
        "Índice de Capital Principal (CET1)", "Índice de Basileia Total (%)", "ROE Acumulado YTD (%)",
    )},
    **{key: (1.935, 1.931, "↑ +0,40 p.p.", "favorable") for key in (
        "PDD / Inadimplência (arrasto)", "Perda Esperada / Estágio 3", "Perda Esperada / Est2+3",
    )},
    "Custo de Crédito / Receita de Crédito (%)": (.2533, .2762, "↓ −2,29 p.p.", "favorable"),
    "Perda Esperada / Carteira de Crédito*": (.0225, .0218, "↑ +0,07 p.p.", "neutral"),
    "Ativo Total / PL": (11.34, 10.22, "↑ +1,12 x", "attention"),
    "Carteira de Crédito* / PL": (11.34, 10.22, "↑ +1,12 x", "attention"),
}


def test_economic_review_covers_the_entire_metric_catalog():
    assert set(ECONOMIC_DELTA_CASES) == set(BY_KEY)


@pytest.mark.parametrize("key", ECONOMIC_DELTA_CASES)
def test_economic_case_has_expected_operation_precision_and_credit_interpretation(key):
    current, old, expected, tone = ECONOMIC_DELTA_CASES[key]
    metric = BY_KEY[key]
    direction, display = delta(current, old, metric, "year")
    assert display == expected
    assert variation_tone(metric, direction) == tone
    reversed_direction, _ = delta(old, current, metric, "year")
    assert variation_tone(metric, reversed_direction) == ({"attention": "favorable", "favorable": "attention"}.get(tone, tone))
    assert delta(None, old, metric, "year") == (None, "Base N/D")
    assert delta(current, None, metric, "year") == (None, "Base N/D")


@pytest.mark.parametrize("key", [key for key in ECONOMIC_DELTA_CASES if BY_KEY[key].unit == "%"])
def test_zero_percent_reference_is_valid_for_subtraction_and_does_not_mean_missing(key):
    direction, display = delta(.0225, 0, BY_KEY[key], "quarter")
    assert direction == "up"
    assert display == ("↑ +225 bps" if BY_KEY[key].delta_unit == "bps" else "↑ +2,25 p.p.")


def test_reference_and_numeric_delta_are_auditable_in_ui_and_excel():
    q,_,values = sample(metrics=("Índice de Basileia Total (%)",), mode="quarter")
    df=pd.DataFrame([{"Instituição":"A","Período":"4/2025"},{"Instituição":"A","Período":"1/2026"},{"Instituição":"A","Período":"1/2025"}])
    values={("Índice de Basileia Total (%)","A","4/2025"):.1518,("Índice de Basileia Total (%)","A","1/2026"):.1477,("Índice de Basileia Total (%)","A","1/2025"):.16}
    common=dict(base=q["base"],cache_token=q["cache_token"],scale=q["scale"],queried_at=q["queried_at"])
    q=build_query(df,["A"],["1/2026"],q["metrics"],values,{},mode="quarter",**common)
    assert q["cells"][0]["delta_value"] == pytest.approx(-41)
    assert q["cells"][0]["delta_unit"] == "bps"
    assert "14,77%" in table_html(q)
    assert "QoQ vs Dez/25" in table_html(q)
    memo=variation_rows(q,q["metrics"][0],"A").iloc[0]
    assert memo["Referência"] == "Dez/25"
    assert "14,7700% − 15,1800%" in memo["Cálculo"]
    yoy=build_query(df,["A"],["1/2026"],q["metrics"],values,{},mode="year",**common)
    assert yoy["cells"][0]["delta_value"] == pytest.approx(-123)
    assert "YoY vs Mar/25" in table_html(yoy)
    ws=load_workbook(BytesIO(export_excel(q)))["Dados e status"]
    records=list(ws.values); headers=records[0]
    assert records[1][headers.index("Delta numérico")] == pytest.approx(-41)
    assert records[1][headers.index("Unidade delta")] == "bps"
    assert methodology_rows(q)[0]["Variação"] == "(razão decimal atual − razão decimal de referência) × 10.000; bps arredondados ao inteiro"


@pytest.mark.parametrize("reference_status", ["warning", "critical"])
def test_reference_quality_alert_preserves_subtraction_and_neutralizes_all_rendered_colors(reference_status):
    key, bank = "Índice de Basileia Total (%)", "A"
    frame = pd.DataFrame([{"Instituição": bank, "Período": "4/2025"}, {"Instituição": bank, "Período": "1/2026"}])
    values = {(key, bank, "4/2025"): .1518, (key, bank, "1/2026"): .1477}
    statuses = {(key, bank, "4/2025"): {"Status analítico": reference_status, "Observação": "RWA da referência requer validação"}}
    query = build_query(frame, [bank], ["1/2026"], [key], values, statuses,
                        base="Consolidada / Prudencial", cache_token="test", scale="R$ bilhões",
                        mode="quarter", queried_at="10/10/2026")
    cell = query["cells"][0]
    assert cell["status"] == "available"
    assert cell["reference_status"] == reference_status
    assert cell["delta_value"] == pytest.approx(-41)
    assert cell["variation"] == "↓ −41 bps"
    rendered = table_html(query)
    assert 'class="delta down neutral"' in rendered
    assert "referência com alerta: RWA da referência requer validação" in rendered
    workbook = load_workbook(BytesIO(export_excel(query)))
    assert workbook["Comparativo"].cell(7, 3).font.color.rgb == "FF666666"
    rows = list(workbook["Dados e status"].values)
    assert rows[1][rows[0].index("Status referência")] == reference_status
    assert rows[1][rows[0].index("Delta numérico")] == pytest.approx(-41)
    deck = Presentation(BytesIO(export_powerpoint(query)))
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert str(table.cell(2, 1).text_frame.paragraphs[1].font.color.rgb) == "666666"


def test_missing_reference_remains_nd_despite_an_older_available_value():
    key, bank = "Índice de Basileia Total (%)", "A"
    frame = pd.DataFrame([{"Instituição": bank, "Período": p} for p in ["3/2025", "4/2025", "1/2026"]])
    values = {(key, bank, "3/2025"): .16, (key, bank, "4/2025"): None, (key, bank, "1/2026"): .1477}
    query = build_query(frame, [bank], ["1/2026"], [key], values, {},
                        base="Consolidada / Prudencial", cache_token="test", scale="R$ bilhões",
                        mode="quarter", queried_at="10/10/2026")
    cell = query["cells"][0]
    assert cell["value"] == .1477
    assert cell["reference_value"] is None
    assert cell["reference_status"] == "missing"
    assert cell["delta_value"] is None
    assert cell["variation"] == "Base N/D"
    assert 'class="delta  neutral"' in table_html(query)


def test_query_preserves_sign_missing_components_order_and_effective_revision():
    q, df, values = sample()
    signed = df.copy()
    signed.loc[(signed.Instituição == "A") & (signed.Período == "2/2026"), "Trace::Custo de Crédito::PDD Crédito Anualizada"] = 2e7
    opts = dict(base=q["base"], cache_token=q["cache_token"], scale=q["scale"], mode=q["mode"], queried_at=q["queried_at"])
    newer = build_query(signed, q["banks"], q["periods"], q["metrics"], values, {}, **opts)
    assert next(c for c in newer["cells"] if c["metric"] == "Custo de Crédito (%)" and c["bank"] == "A" and c["period"] == "2/2026")["value"] == -.02
    assert newer["signature"] != q["signature"]
    assert build_query(df, list(reversed(q["banks"])), q["periods"], q["metrics"], values, {}, **opts)["signature"] != q["signature"]
    assert build_query(df, q["banks"], q["periods"], q["metrics"], values, {}, **{**opts,"queried_at":"tomorrow"})["signature"] == q["signature"]
    liquidity = build_query(df, ["A"], ["2/2026"], ["Ativos Líquidos"], {("Ativos Líquidos", "A", "2/2026"): 10}, {}, **opts)
    assert liquidity["cells"][0]["value"] is None
    assert "incompletos" in liquidity["cells"][0]["reason"]


def test_profit_memory_uses_displayed_result_and_table_has_nine_columns():
    q, df, _ = sample()
    memo = calculation_rows(q, df, "Lucro Líquido Acumulado", "A")
    result = memo[memo.Campo == "Resultado na tabela"]
    assert set(result.Valor) == {"0,10"}
    html = table_html(q)
    assert html.count('colspan="3"') == 3
    assert 'data-metric="Custo de Crédito (%)"' in html
    assert '<script>' not in table_html({**q, "banks": ["<script>alert(1)</script>"], "cells": [{**c, "bank": "<script>alert(1)</script>"} for c in q["cells"] if c["bank"] == "A"]})


def test_table_embeds_units_and_formats_percentages_without_mutating_query():
    q, _, _ = sample()
    q["cells"][0]["value"] = 2739.22e9
    cost = next(c for c in q["cells"] if c["metric"] == "Custo de Crédito (%)")
    cost.update(value=.13333, display="13,33%")
    q["cells"][0]["variation"] = "↑ +13,33 %"
    before = json.dumps(q, sort_keys=True)
    html = table_html(q)
    table = etree.fromstring(html)
    assert "Unidade" not in html
    assert all(len(row) == 10 for row in table.findall("tbody/tr") if row.get("class") != "section")
    assert all(row[0].get("colspan") == "10" for row in table.findall("tbody/tr") if row.get("class") == "section")
    assert "Ativo total (R$ bi)" in html
    assert "Custo de crédito (%)" in html
    assert ">2.739,22<" in html
    assert ">13,33%<" in html
    assert "↑ +13,33%" in html
    assert json.dumps(q, sort_keys=True) == before
    assert table_row_label(BY_KEY["Ativo Total"], "R$ milhões") == "Ativo total (R$ mi)"
    titled = replace(BY_KEY["Custo de Crédito (%)"], label="Custo de crédito (%)")
    assert table_row_label(titled, q["scale"]) == titled.label


def test_table_marks_series_break_values_and_gives_specific_footnote():
    q, _, _ = sample(metrics=("Carteira de Crédito*", "Core Funding*"))
    for cell in q["cells"]:
        cell.update(value=1e9, display="1,00")
    html = table_html(q)
    assert "Quebra em 2025" not in ''.join(etree.fromstring(html).itertext())
    assert html.count(">1,00*<") == 6
    assert html.count(">1,00<") == 12
    note = table_footnote(q)
    assert note.startswith("* Quebra de série em 2025: ")
    assert "e1 + f1 + g1 + h1" in note
    assert "instrumentos elegíveis a capital" in note
    funding = {**q, "cells": [c for c in q["cells"] if c["metric"] == "Core Funding*"]}
    assert "carteira ampliada" not in table_footnote(funding)
    for cell in q["cells"]:
        cell.update(value=None, display="N/D")
    assert "N/D*" not in table_html(q)
    assert "N/D" in table_html(q)
    assert table_footnote(q) == ""


def test_excel_is_numeric_percentage_missing_and_metadata_are_preserved():
    q, _, _ = sample()
    q["cells"][0]["value"] = None
    workbook = load_workbook(BytesIO(export_excel(q)))
    assert set(workbook.sheetnames) == {"Comparativo", "Dados e status", "Metodologia", "Consulta"}
    sheet = workbook["Comparativo"]
    assert sheet.cell(6, 3).value is None
    assert sheet.cell(6, 4).value == 1
    cost = next(r for r in range(1, sheet.max_row+1) if sheet.cell(r, 1).value == "Custo de crédito")
    assert sheet.cell(cost, 3).value == .02
    assert sheet.cell(cost, 3).number_format == "0.00%"
    metadata = {r[0].value: r[1].value for r in workbook["Consulta"] if r[0].value != "Campo"}
    assert metadata["banks"] == "['A', 'B', 'C']"


def test_powerpoint_table_is_native_three_by_three_and_paginated():
    q, _, _ = sample(metrics=tuple(BY_KEY))
    # Outros indicadores ficam N/D no fixture; continuam presentes e editáveis.
    data = export_powerpoint(q)
    prs = Presentation(BytesIO(data))
    assert len(prs.slides) == 3
    for slide in prs.slides:
        table = next(s.table for s in slide.shapes if s.has_table)
        assert len(table.columns) == 10
        assert len(table.rows) <= 12
        assert table.cell(0, 1).is_merge_origin
        assert table.cell(1, 3).text == "Jun/26\nYoY vs Jun/25"
        assert table.cell(2, 0).text_frame.paragraphs[1].font.size.pt == 9
        assert table.cell(2, 0).text_frame.paragraphs[1].font.bold is False
        assert all(s.top+s.height <= prs.slide_height for s in slide.shapes)
    with ZipFile(BytesIO(data)) as archive:
        assert not any(n.startswith("ppt/media/") for n in archive.namelist())


def test_native_line_chart_last_valid_label_bold_color_and_embedded_workbook():
    q, _, _ = sample(metrics=("Ativo Total",))
    for cell in q["cells"]:
        if cell["bank"] == "A" and cell["period"] == "2/2026":
            cell["value"] = None
    data = export_powerpoint(q, charts=True, chart_metrics=["Ativo Total"])
    ns = {"c":"http://schemas.openxmlformats.org/drawingml/2006/chart", "a":"http://schemas.openxmlformats.org/drawingml/2006/main"}
    with ZipFile(BytesIO(data)) as archive:
        assert any(n.startswith("ppt/embeddings/") for n in archive.namelist())
        chart = etree.fromstring(archive.read("ppt/charts/chart1.xml"))
        assert chart.xpath("count(.//c:lineChart)", namespaces=ns) == 1
        series = chart.xpath(".//c:lineChart/c:ser", namespaces=ns)
        labels = series[0].xpath(".//c:dLbl[c:idx/@val='1']", namespaces=ns)
        assert labels
        assert labels[0].xpath(".//a:defRPr[@b='1'] | .//a:rPr[@b='1']", namespaces=ns)
        assert labels[0].xpath(".//a:srgbClr[@val='174A7E']", namespaces=ns)
        assert chart.xpath(".//c:catAx", namespaces=ns)
        assert not chart.xpath(".//c:dateAx", namespaces=ns)
        assert chart.xpath(".//c:dispBlanksAs[@val='gap']", namespaces=ns)
    cols = export_powerpoint(q, charts=True, chart_metrics=["Ativo Total"], chart_type="column_stacked")
    with ZipFile(BytesIO(cols)) as archive:
        chart = etree.fromstring(archive.read("ppt/charts/chart1.xml"))
        assert chart.xpath(".//c:barChart/c:grouping[@val='clustered']", namespaces=ns)
        assert chart.xpath(".//c:barChart/c:ser[1]/c:dLbls/c:dLbl/c:dLblPos[@val='outEnd']", namespaces=ns)
        assert chart.xpath(".//c:barChart/c:ser[1]/c:dLbls/c:dLbl//a:srgbClr[@val='174A7E']", namespaces=ns)


def test_groups_identity_roundtrip_unavailable_members_and_overwrite_gate():
    blank = {"schema_version": 2, "groups": []}
    payload = peers_groups.upsert(blank, "Meus bancos", "Individual", ["A", "B"], {"A":"1", "B":"2"})
    assert peers_groups.resolve(payload["groups"][0], {"Novo nome A":"1"}) == (["Novo nome A"], ["B"])
    with pytest.raises(ValueError, match="Já existe"):
        peers_groups.upsert(payload, "Meus bancos", "Individual", ["A"], {"A":"1"})
    with pytest.raises(ValueError, match="Identidade"):
        peers_groups.upsert(blank, "Grupo", "Individual", ["A"], {})


def test_groups_github_conflict_and_readback_are_required():
    payload = peers_groups.upsert({"schema_version":2,"groups":[]}, "Teste", "Individual", ["A"], {"A":"1"})
    client = Mock()
    client.put.return_value.status_code = 200
    client.put.return_value.json.return_value = {"content":{"sha":"new"}}
    client.get.return_value.status_code = 200
    client.get.return_value.json.return_value = {"sha":"new", "content":base64.b64encode(json.dumps(payload).encode()).decode()}
    assert peers_groups.publish(payload, "org/repo", "main", "token", "old", client=client) == "new"
    assert client.put.call_args.kwargs["json"]["sha"] == "old"
    client.get.return_value.json.return_value["sha"] = "other"
    with pytest.raises(RuntimeError, match="confirmação"):
        peers_groups.publish(payload, "org/repo", "main", "token", "old", client=client)
    client.put.return_value.status_code = 409
    with pytest.raises(ValueError, match="mudou"):
        peers_groups.publish(payload, "org/repo", "main", "token", "old", client=client)


def test_invalid_groups_fail_with_actionable_validation():
    for group in ({"name": 4, "base": "Individual", "members": []}, {"name": "A", "base": [], "members": []}, {"name": "A", "base": "Individual", "members": [{"id": [], "label": "A"}]}):
        with pytest.raises(ValueError):
            peers_groups.validate({"schema_version": 2, "groups": [group]})


def test_schema_break_missing_june_and_duplicate_identity_are_explicit():
    q, df, values = sample(periods=("4/2025",), metrics=("Carteira de Crédito*",))
    assert q["cells"][0]["variation"] == "Quebra em 2025"
    q, df, values = sample(periods=("4/2026",), metrics=("Lucro Líquido Acumulado",))
    assert all(c["value"] is None for c in q["cells"])
    opts = dict(base=q["base"], cache_token=q["cache_token"], scale=q["scale"], mode=q["mode"], queried_at=q["queried_at"])
    with pytest.raises(ValueError, match="mais de um registro"):
        build_query(pd.concat([df, df.iloc[:1]]), q["banks"], q["periods"], q["metrics"], values, {}, **opts)


def test_exports_keep_balances_neutral_and_risk_directions_meaningful():
    q, _, _ = sample(metrics=("Ativo Total",))
    q["cells"][0].update(direction="up", variation="↑ +1,00 %")
    q["cells"][1].update(direction="down", variation="↓ −1,00 %")
    sheet = load_workbook(BytesIO(export_excel(q)))["Comparativo"]
    assert sheet.cell(6, 3).value == 1
    assert sheet.cell(7, 3).font.color.rgb == "FF666666"
    assert sheet.cell(7, 4).font.color.rgb == "FF666666"
    risk, _, _ = sample(metrics=("Custo de Crédito (%)",))
    risk["cells"][0].update(direction="up", variation="↑ +11 bps")
    risk["cells"][1].update(direction="down", variation="↓ −11 bps")
    sheet = load_workbook(BytesIO(export_excel(risk)))["Comparativo"]
    assert sheet.cell(7, 3).font.color.rgb == "FFB32624"
    assert sheet.cell(7, 4).font.color.rgb == "FF16713B"
    prs = Presentation(BytesIO(export_powerpoint(risk)))
    table = next(s.table for s in prs.slides[0].shapes if s.has_table)
    assert str(table.cell(2, 1).text_frame.paragraphs[1].font.color.rgb) == "B32624"
    assert str(table.cell(2, 2).text_frame.paragraphs[1].font.color.rgb) == "16713B"


def test_individual_definitions_and_funding_do_not_inherit_consolidated_formula():
    q, df, values = sample(periods=("4/2025",), metrics=("Core Funding*", "Custo de Crédito (%)"), base="Individual")
    df["Captações"] = 1e9
    opts = dict(base=q["base"], cache_token=q["cache_token"], scale=q["scale"], mode=q["mode"], queried_at=q["queried_at"])
    q = build_query(df, q["banks"], q["periods"], q["metrics"], values, {}, **opts)
    assert q["cells"][0]["value"] == 1e9
    assert q["cells"][0]["variation"] == "= 0,00 %"
    assert q["cells"][0]["source"] == "IFData Rel. 1 individual"
    assert q["cells"][-1]["value"] is None
    assert get_metric("Carteira de Crédito*", "Individual").label == "Carteira de crédito"
    assert get_metric("ROE Acumulado YTD (%)", "Individual").formula.endswith("PL atual")
    assert methodology_rows(q)[0]["Indicador"] == "Captações"
    assert "Captações" in table_html(q)


def test_deeplink_opens_once_and_allows_navigation_to_other_tabs():
    tree = ast.parse(Path("app1.py").read_text())
    node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "_aplicar_navegacao_inicial_mobile")
    state = {}
    namespace = {"st": SimpleNamespace(session_state=state), "_menu_query_param_inicial": lambda: "Tabela de Peers"}
    exec(compile(ast.Module(body=[node], type_ignores=[]), "app1.py", "exec"), namespace)
    namespace["_aplicar_navegacao_inicial_mobile"]()
    assert state["menu_atual"] == "Tabela de Peers"
    state["menu_atual"] = "Snapshot"
    namespace["_aplicar_navegacao_inicial_mobile"]()
    assert state["menu_atual"] == "Snapshot"


def test_legacy_menu_and_exclusive_helpers_are_removed_shared_apis_remain():
    tree = ast.parse(Path("app1.py").read_text())
    assignments = {
        target.id: node.value
        for node in tree.body if isinstance(node, ast.Assign)
        for target in node.targets if isinstance(target, ast.Name)
    }
    assert "Peers (Tabela)" not in ast.literal_eval(assignments["MENU_PRINCIPAL"])
    assert "Tabela de Peers" in ast.literal_eval(assignments["MENU_PRINCIPAL"])
    dependencies = ast.literal_eval(assignments["CACHE_DEPENDENCIAS_POR_ABA"])
    assert "Peers (Tabela)" not in dependencies
    assert dependencies["Tabela de Peers"] == ["critical_screens"]
    routes = [
        value.value for node in ast.walk(tree)
        if isinstance(node, ast.Compare) and isinstance(node.left, ast.Name) and node.left.id == "menu"
        for value in node.comparators if isinstance(value, ast.Constant)
    ]
    assert "Peers (Tabela)" not in routes
    assert routes.count("Tabela de Peers") == 1
    functions = {node.name for node in tree.body if isinstance(node, ast.FunctionDef)}
    assert not functions.intersection({
        "_render_peers_table_html", "_gerar_imagem_peers_tabela",
        "_gerar_excel_peers_tabela", "_gerar_excel_peers_dados_puros",
        "_carregar_peer_groups_salvos", "_salvar_peer_groups_salvos",
        "_filtrar_peer_group", "_peer_groups_disponiveis",
        "_build_memoria_calculo_peers_tabela_metrica", "_metric_definition_html",
        "_build_peers_visual_status_artifacts", "_merge_peers_analytical_tooltips",
        "_decorate_peers_visual_value", "_normalizar_base_dre_peers",
        "_ordenar_periodos_peers_saida", "_timer_begin_measurement",
    })
    assert {
        "_build_peers_export_status_rows", "_build_peers_status_lookup",
        "_montar_tabela_peers", "_preparar_metricas_extra_peers_from_slice",
        "_get_peers_filters_context", "_get_peers_individual_filters_context",
        "_apply_peers_individual_display_names", "_garantir_cache_telas_criticas",
        "_gerar_excel_evolucao_dados_puros", "_build_memoria_calculo_curado_metrica",
    }.issubset(functions)


def test_shared_analytical_status_survives_query_and_native_excel_export():
    import app1
    bank, period = "TEST BANK - PRUDENCIAL", "4/2025"
    df = pd.DataFrame([{
        "Instituição": bank, "Período": period, "Ativo Total": None,
        "Depósitos Totais": 106.0, "Core Funding": None,
        "Trace::Depósitos Totais::Status": "fallback_components",
        "Trace::Core Funding::Status": "missing_required_component",
        "Trace::Core Funding::Campo Selecionado": "Captações (e) + Instrumentos de Dívida Elegíveis a Capital (h)",
        "Trace::Core Funding::Captações (e)": 210.0,
        "Trace::Core Funding::Instrumentos de Dívida Elegíveis a Capital (h)": None,
    }])
    metrics = ["Ativo Total", "Depósitos Totais", "Core Funding*"]
    values = {(metric, bank, period): 106.0 if metric == "Depósitos Totais" else None for metric in metrics}
    columns = {"Ativo Total": "Ativo Total", "Depósitos Totais": "Depósitos Totais", "Core Funding*": "Core Funding"}
    status = app1._build_peers_status_lookup(
        df_base=df, bancos=[bank], periodos=[period], valores=values, colunas_usadas=columns,
    )
    query = build_query(
        df, [bank], [period], metrics, values, status,
        base="Consolidada / Prudencial", cache_token="test", scale="R$ bilhões",
        mode="year", queried_at="09/10/2026",
    )
    cells = {cell["metric"]: cell for cell in query["cells"]}
    assert cells["Depósitos Totais"]["status"] == "fallback_components"
    assert cells["Depósitos Totais"]["value"] == 106.0
    assert cells["Ativo Total"]["status"] == "missing"
    assert "sem valor disponível" in cells["Ativo Total"]["reason"].lower()
    assert cells["Core Funding*"]["value"] is None
    assert table_html(query).count('<span class="value">N/D</span>') == 2
    ledger = load_workbook(BytesIO(export_excel(query)))["Dados e status"]
    rows = {row[0].value: row for row in ledger.iter_rows(min_row=2)}
    assert rows["Depósitos Totais"][3].value == 106.0
    assert rows["Depósitos Totais"][7].value == "fallback_components"
    assert rows["Ativo Total"][3].value is None
    assert rows["Core Funding*"][3].value is None
    assert rows["Ativo Total"][9].value == cells["Ativo Total"]["reason"]


def arrasto_query(**changes):
    row = {"Instituição": "A", "Período": "2/2026", "Carteira Total 4.966": 1000e6,
           "Inadimplência 4.966": 20e6, "Inadimplência": 990e6, "Carteira de Crédito Bruta": 1500e6,
           "Perda Esperada": -900e6,
           **{f"Trace::Perda Esperada::Perda Esperada ({c}2)": -v * 1e6 for c, v in zip("efgh", [12, 2, 3, 3])},
           "Trace::Perda Esperada::Hedge de Valor Justo (e3)": -880e6}
    row.update(changes)
    prior = {**row, "Período": "1/2026", "Inadimplência 4.966": 10e6}
    before_2025 = {**row, "Período": "4/2024"}
    df = pd.DataFrame([row, prior, before_2025])
    q = build_query(df, ["A"], ["4/2024", "2/2026"], list(ARRASTO_ROWS), {}, {},
                    base="Consolidada / Prudencial", cache_token="test", scale="R$ milhões",
                    mode="quarter", queried_at="10/10/2026")
    return q, df


def test_arrasto_uses_rel16_total_and_only_four_expected_losses_from_2025():
    q, df = arrasto_query()
    cells = {c["metric"]: c for c in q["cells"] if c["period"] == "2/2026"}
    assert set(ARRASTO_ROWS).issubset(DEFAULT_METRICS)
    assert cells["Inadimplência"]["value"] == 20e6
    assert cells["Inadimplência / Carteira Total"]["value"] == .02
    assert cells["PDD / Inadimplência (arrasto)"]["value"] == 1
    assert cells["Inadimplência / Carteira Total"]["variation"] == "↑ +100 bps"
    assert cells["PDD / Inadimplência (arrasto)"]["variation"] == "↓ −100,00 p.p."
    html = table_html(q)
    assert "↓ −100,00 p.p." in html and "(razão decimal atual − razão decimal de referência) × 100" in html
    variations = variation_rows(q, "PDD / Inadimplência (arrasto)", "A")
    assert "100,0000% − 200,0000% = −100,0000 p.p." in variations.iloc[-1]["Cálculo"]
    assert "× 100" not in variations.iloc[-1]["Cálculo"]
    sheet = load_workbook(BytesIO(export_excel(q)))["Comparativo"]
    coverage_row = next(row for row in range(1, sheet.max_row + 1) if sheet.cell(row, 1).value == "PDD / vencidos >90 dias (arrasto)")
    assert sheet.cell(coverage_row, 4).number_format == "0.00%"
    assert sheet.cell(coverage_row + 1, 4).value == "↓ −100,00 p.p."
    assert all(c["value"] is None and "mar/2025" in c["reason"] for c in q["cells"] if c["period"] == "4/2024")
    memo = calculation_rows(q, df, "PDD / Inadimplência (arrasto)", "A")
    assert not memo.Campo.str.contains("Hedge|Ajuste a Valor Justo|Carteira de Crédito Bruta").any()
    assert memo[(memo.Período == "Jun/26") & (memo.Campo == "PDD (soma das quatro perdas esperadas)")].Valor.tolist() == ["20,00"]


@pytest.mark.parametrize("changes", [{"Trace::Perda Esperada::Perda Esperada (h2)": None},
                                      {"Inadimplência 4.966": 0}, {"Inadimplência 4.966": None},
                                      {"Inadimplência 4.966": -1}])
def test_arrasto_coverage_missing_component_or_invalid_denominator_stays_missing(changes):
    q, _ = arrasto_query(**changes)
    cell = next(c for c in q["cells"] if c["metric"] == "PDD / Inadimplência (arrasto)" and c["period"] == "2/2026")
    assert cell["value"] is None
    assert cell["status"] == "missing"
    assert cell["reason"]


def test_arrasto_quality_warning_reaches_html_excel_and_powerpoint():
    q, _ = arrasto_query(**{"Trace::Perda Esperada::Perda Esperada (e2)": 12e6})
    cell = next(c for c in q["cells"] if c["metric"] == "PDD / Inadimplência (arrasto)" and c["period"] == "2/2026")
    assert cell["status"] == "warning" and "sinal" in cell["reason"]
    assert cell["display"].endswith("†")
    assert "†" in table_html(q)
    wb = load_workbook(BytesIO(export_excel(q)))
    assert any("sinal" in str(c.value) for row in wb["Dados e status"] for c in row)
    prs = Presentation(BytesIO(export_powerpoint(q)))
    assert any("sinal" in s.notes_slide.notes_text_frame.text for s in prs.slides)
    assert any("†" in c.text for s in prs.slides for shape in s.shapes if shape.has_table for row in shape.table.rows for c in row.cells)


def test_period_values_and_variations_are_centered_in_native_exports():
    q, _, _ = sample()
    wb = load_workbook(BytesIO(export_excel(q)))
    assert wb["Comparativo"].cell(6, 3).alignment.horizontal == "center"
    assert wb["Comparativo"].cell(7, 3).alignment.horizontal == "center"
    from pptx.enum.text import PP_ALIGN
    prs = Presentation(BytesIO(export_powerpoint(q)))
    table = next(s.table for s in prs.slides[0].shapes if s.has_table)
    assert all(p.alignment == PP_ALIGN.CENTER for p in table.cell(2, 1).text_frame.paragraphs)
    assert table.cell(2, 0).text_frame.paragraphs[0].alignment == PP_ALIGN.LEFT


def test_old_peers_menu_links_resolve_to_renamed_tab():
    tree = ast.parse(Path("app1.py").read_text())
    node = next(n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name == "_normalizar_rotulo_menu")
    namespace = {}
    exec(compile(ast.Module(body=[node], type_ignores=[]), "app1.py", "exec"), namespace)
    assert namespace["_normalizar_rotulo_menu"]("Tabela de peers") == "Tabela de Peers"
    assert namespace["_normalizar_rotulo_menu"]("Peers (Tabela Nova)") == "Tabela de Peers"
    assert namespace["_normalizar_rotulo_menu"]("Peers (Tabela)") == "Tabela de Peers"


def test_arrasto_individual_base_does_not_reuse_prudential_values():
    _, df = arrasto_query()
    query = build_query(df, ["A"], ["2/2026"], list(ARRASTO_ROWS), {}, {},
                        base="Individual", cache_token="test", scale="R$ milhões",
                        mode="year", queried_at="10/10/2026")
    assert all(c["value"] is None and "individuais" in c["reason"] for c in query["cells"])
