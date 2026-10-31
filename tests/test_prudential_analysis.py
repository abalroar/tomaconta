from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from utils.prudential_analysis import Accounts, build_analysis, stem
from utils.prudential_workbook import render_xlsx
from utils.ifdata_cache.bloprudencial_ingest import read_source, ingest_source


def records(accounts,period="202606",doc="4066",code="C1"):
    return pd.DataFrame([dict(DATA_BASE=period,DOCUMENTO=doc,COD_CONGL=code,CONTA=c,NOME_CONTA=n,SALDO=v,
                              CNPJ="00001234",NOME_INSTITUICAO="Banco Teste",NOME_CONGL="Teste Prudencial") for c,n,v in accounts])


def test_cosif_hierarchy_respects_level_boundaries_and_check_digit():
    data=records([("1880000009","Diversos",100),("1884000005","Garantia",40),("1884500000","Impostos",60)])
    account=Accounts(data)
    assert stem("1884000005")=="18840"
    assert account.select("188")[0]==["1880000009"]
    assert account.select("1884")[0]==["1884000005"]
    assert Accounts(data.iloc[1:]).select("188")[0]==["1884000005","1884500000"]
    assert stem("18840005")=="18840"


def test_cosif_general_total_is_not_a_compensation_child(simple_sources):
    extra=records([("3000000007","Compensação ativa",100),
                   ("3010000004","Coobrigações",100),
                   ("3999999009","TOTAL GERAL DO ATIVO",1100)])
    a=build_analysis(pd.concat([simple_sources,extra]),{},"C1",["202606"])
    check=next(c for c in a.checks if c["check"].startswith("Subcontas 4066: 3000000007"))
    assert check["status"]=="OK"
    assert Accounts(extra.iloc[1:]).select("3")[0]==["3010000004"]


@pytest.fixture
def simple_sources():
    amounts=[("1000000009","Ativo",1000),("1100000002","Disponibilidades",1000),
             ("4000000006","Passivo",800),("4100000009","Depósitos",800),
             ("4150000004","Prazo",800),("6000000004","PL",200)]
    data=records(amounts)
    pre=records([(c,n,150 if c.startswith("6") else v) for c,n,v in amounts],doc="4060")
    result=records([("7000000003","Credor",70),("7100000006","Receitas",70),("7110000003","Crédito",70),
                    ("8000000002","Devedor",-20),("8100000005","Despesas",-20),("8110000002","Captação",-20)],doc="4060")
    return pd.concat([data,pre,result],ignore_index=True)


def value(a,sheet,row_id,col="C"):
    return a.sheets[sheet].cells[f"{col}{a.rows[sheet][row_id]}"].value


def test_4066_is_not_added_to_4060_and_preclosing_equity_is_bridged(simple_sources):
    a=build_analysis(simple_sources,{},"C1",["202606"])
    assert value(a,"Balanço","assets")==.001
    assert value(a,"Balanço","equity")==.0002
    assert value(a,"DRE","net")==pytest.approx(.00005)
    assert value(a,"Balanço","roe")==pytest.approx(.5)
    assert all(c["status"]=="OK" for c in a.checks if c["check"].startswith(("Balanço: Ativo", "Balanço: composição", "Conciliação patrimonial:","DRE: resultado")))
    fallback=build_analysis(simple_sources[simple_sources.DOCUMENTO.eq("4060")],{},"C1",["202606"])
    assert value(fallback,"Balanço","equity")==.0002


def test_missing_zero_and_identity_are_not_interchangeable(simple_sources):
    capital=pd.DataFrame([{"CodInst":"C2","Período":"2/2026","Capital Principal":100,"RWA Total":1000},
                          {"CodInst":"C1","Período":"2/2026","Capital Principal":0,"RWA Total":1000}])
    a=build_analysis(simple_sources,{"capital":capital},"C1",["202606"])
    assert value(a,"Balanço","cet1")==0
    assert value(a,"Balanço","basel")=="N/D"
    assert value(a,"Carteira 4.966","writeoffs")=="N/D"
    assert value(a,"Carteira 4.966","total")=="N/D"
    with pytest.raises(ValueError,match="Duplicidade"):
        build_analysis(pd.concat([simple_sources,simple_sources.iloc[:1]]),{},"C1",["202606"])


def test_december_dre_is_second_semester_without_adding_june(simple_sources):
    dec=simple_sources.copy();dec.DATA_BASE="202612"
    a=build_analysis(pd.concat([simple_sources,dec]),{},"C1",["202606","202612"])
    assert value(a,"DRE","net","E")==value(a,"DRE","net","C")
    assert "2S26" in a.sheets["DRE"].cells["E5"].value


def test_xlsx_contains_live_formulas_snapshot_style_and_hidden_inputs(simple_sources):
    a=build_analysis(simple_sources,{},"C1",["202606"])
    payload=render_xlsx(a)
    wb=load_workbook(BytesIO(payload),data_only=False)
    assert wb.sheetnames[:4]==["Comitê","Balanço","DRE","Carteira 4.966"]
    assert wb["Comitê"]["C7"].value == "='Balanço'!C7"
    assert "INDEX(" in wb["Mapa COSIF"]["I2"].value and "MATCH(" in wb["Mapa COSIF"]["I2"].value
    assert wb["Balanço"]["B2"].font.name=="Calibri"
    assert wb["Balanço"].row_dimensions[4].hidden
    assert wb["Comitê"].freeze_panes == "C7"
    assert wb["Balanço"]["C7"].number_format.startswith("[$-416]")
    assert all(wb[name].sheet_state == "hidden" for name in ("Mapa COSIF", "Dados COSIF", "Dados IFData", "Validações"))
    assert wb["Comitê"].page_setup.orientation == "landscape"
    assert all(wb[name].page_setup.orientation == "portrait" for name in wb.sheetnames[1:4])
    assert all((wb[name].page_setup.scale or 100) >= 60 for name in wb.sheetnames[:4])
    assert not any(wb[name].sheet_properties.pageSetUpPr and wb[name].sheet_properties.pageSetUpPr.fitToPage for name in wb.sheetnames[:4])
    row=a.rows["Balanço"]["roe"]
    assert wb["Balanço"][f"D{row}"].fill.fgColor.rgb=="FF2F3640"
    assert all("Validações" not in c.value for s in wb if s.title!="Validações" for row in s for c in row if c.data_type=="f")


def source_file(tmp_path,rows):
    path=tmp_path/"202606BLOPRUDENCIAL.CSV"
    path.write_text("Balancete\nData de geracao: 2026-10-19\nFonte: IFs\n#DATA_BASE;DOCUMENTO;CNPJ;AGENCIA;NOME_INSTITUICAO;COD_CONGL;NOME_CONGL;TAXONOMIA;CONTA;NOME_CONTA;SALDO\n"+"\n".join(rows)+"\n",encoding="latin1")
    return path


def test_strict_ingestion_preserves_zero_leading_cnpj_and_replaces_period(tmp_path):
    row="202606;4066;00001234;;Banco;C1;Grupo;Prudencial;1100000002;Caixa;0,00"
    path=source_file(tmp_path,[row])
    data,meta=read_source(path)
    assert data.CNPJ.iloc[0]=="00001234" and data.SALDO.iloc[0]==0
    root=tmp_path/"repo";root.mkdir()
    first=ingest_source(path,root)
    path=source_file(tmp_path,[row.replace("1100000002;Caixa;0,00","1000000009;Ativo;100,00")])
    second=ingest_source(path,root)
    saved=pd.read_parquet(root/"data/cache/bloprudencial/dados.parquet")
    assert len(saved)==1 and saved.CONTA.iloc[0]=="1000000009"
    assert second["comparison"]=={"previous_rows":1,"added":1,"removed":1,"changed_balances":0}
    assert (root/first["archive"]).exists() and (root/second["archive"]).exists()
    source_file(tmp_path,[row,row])
    with pytest.raises(ValueError,match="duplicada"):read_source(path)
    source_file(tmp_path,[row.replace("0,00","")])
    with pytest.raises(ValueError,match="Saldo"):read_source(path)


@pytest.mark.parametrize("code",["C0083539","C0080099","C0080075","C0080185","C0080336","C0084693","C0084710"])
def test_bundled_real_institutions_close_without_synthetic_child_double_count(code):
    from tools.build_prudential_workbook import load_local_sources
    root=Path(__file__).resolve().parents[1]
    c,s=load_local_sources(root,["202512","202606"])
    a=build_analysis(c,s,code,["202512","202606"])
    critical=[x for x in a.checks if x["check"].startswith(("Balanço: Ativo", "Balanço: composição", "DRE: resultado"))]
    assert critical and all(x["status"]=="OK" for x in critical),critical
    assert not any(x["status"]=="Divergência" for x in a.checks if "subcontas apresentadas" in x["check"])


def test_native_pptx_first_slide_has_both_statements_in_last_three_periods(simple_sources):
    import json
    from pptx import Presentation
    from utils.prudential_pptx import render_pptx
    periods = ["202312", "202406", "202412", "202506", "202512", "202606"]
    analysis = build_analysis(pd.concat([simple_sources.assign(DATA_BASE=p) for p in periods]), {}, "C1", periods)
    deck = Presentation(BytesIO(render_pptx(analysis)))
    first = deck.slides[0]
    tables = [shape for shape in first.shapes if shape.has_table]
    assert len(tables) == 2 and tables[0].left < tables[1].left
    assert tables[0].top + tables[0].height <= deck.slide_height
    assert json.loads(first.notes_slide.notes_text_frame.text)["periods"] == periods[-3:]
    assert all(len(shape.table.columns) == 7 for shape in tables)
    for slide in deck.slides:
        assert not any(shape.shape_type == 13 for shape in slide.shapes)


def test_committee_long_selection_uses_legible_pages_with_every_period(simple_sources):
    periods = ["202406", "202412", "202506", "202512", "202603", "202606"]
    analysis = build_analysis(pd.concat([simple_sources.assign(DATA_BASE=p) for p in periods]), {}, "C1", periods)
    workbook = load_workbook(BytesIO(render_xlsx(analysis)))
    committee = workbook["Comitê"]
    assert committee.max_column == 12
    assert len(committee.row_breaks.brk) == 2
    dates = [cell.value.strftime("%Y%m") for row in committee for cell in row if getattr(cell, "is_date", False)]
    assert sorted(set(dates)) == periods
    assert sum(cell.data_type == "f" for row in committee for cell in row) > 100


def test_update_gate_rejects_broken_balance_and_duplicate_new_identity(simple_sources):
    from utils.ifdata_cache.financial_analysis_update import validate_candidate
    baseline = validate_candidate(simple_sources, {}, "202606")
    assert not baseline["errors"]
    broken = simple_sources.copy()
    broken.loc[broken.DOCUMENTO.eq("4066") & broken.CONTA.eq("6000000004"), "SALDO"] += 100
    report = validate_candidate(broken, {}, "202606")
    assert any(item["check"] == "Balanço: Ativo = Passivo + PL" for item in report["errors"])
    missing = simple_sources[~simple_sources.CONTA.str.startswith("6")]
    assert any(item["check"] == "Grupos patrimoniais obrigatórios" for item in validate_candidate(missing, {}, "202606")["errors"])
    new = pd.DataFrame([{"CodInst": "C1", "Período": "2/2026", "Lucro Líquido": 50}]*2)
    with pytest.raises(ValueError, match="chave"):
        validate_candidate(simple_sources, {"dre": new}, "202606")
    legacy = pd.DataFrame([{"CodInst": None, "Período": "1/2015", "Lucro Líquido": 50}]*2)
    assert not validate_candidate(simple_sources, {"dre": legacy}, "202606")["errors"]


def test_update_gate_keeps_cross_document_differences_in_report(simple_sources):
    from utils.ifdata_cache.financial_analysis_update import validate_candidate
    other = simple_sources.copy()
    other.loc[other.DOCUMENTO.eq("4060") & other.CONTA.eq("6000000004"), "SALDO"] += 20
    report = validate_candidate(other, {}, "202606")
    assert not report["errors"]
    assert any(item["check"].startswith("Conciliação patrimonial") for item in report["warnings"])


def test_month_discovery_stops_at_access_error_and_accepts_latest(monkeypatch, tmp_path):
    from datetime import date
    import requests
    from utils.ifdata_cache import financial_analysis_update as update
    calls = []
    def download(period, *args, **kwargs):
        calls.append(period)
        if period == "202608":
            return tmp_path / "source.zip"
        response = requests.Response(); response.status_code = 404
        raise requests.HTTPError(response=response)
    monkeypatch.setattr(update, "download_bloprudencial_zip", download)
    monkeypatch.setattr(update, "extract_bloprudencial_csv", lambda *args: tmp_path / "source.csv")
    assert update.discover_source(tmp_path, today=date(2026, 10, 30))[0] == "202608"
    assert calls == ["202609", "202608"]
    def blocked(*args, **kwargs):
        response = requests.Response(); response.status_code = 403
        raise requests.HTTPError(response=response)
    monkeypatch.setattr(update, "download_bloprudencial_zip", blocked)
    with pytest.raises(requests.HTTPError):
        update.discover_source(tmp_path, today=date(2026, 10, 30))


def test_refresh_is_idempotent_and_rejects_invalid_candidate_without_writing(simple_sources, tmp_path, monkeypatch):
    from tools import update_financial_analysis as cli
    from utils.ifdata_cache import BloprudencialCache
    # The fixture contains only COSIF; there are no dependent reports to rebuild.
    monkeypatch.setattr(cli, "materialize_for_publication", lambda *args, **kwargs: [])
    root = tmp_path / "repo"; root.mkdir()
    path = tmp_path / "source.csv"
    simple_sources.rename(columns={"DATA_BASE": "#DATA_BASE"}).to_csv(path, sep=";", decimal=",", index=False)
    first, receipt = cli.update(root, csv=path, offline=True)
    assert first["status"] == "updated", first
    assert first["sources"] and receipt.is_file()
    active = BloprudencialCache(root).arquivo_dados
    original_bytes = active.read_bytes()
    second, _ = cli.update(root, csv=path, offline=True)
    assert second["status"] == "unchanged", second
    assert active.read_bytes() == original_bytes
    broken = simple_sources.copy()
    broken.loc[broken.DOCUMENTO.eq("4066") & broken.CONTA.eq("6000000004"), "SALDO"] += 100
    broken.rename(columns={"DATA_BASE": "#DATA_BASE"}).to_csv(path, sep=";", decimal=",", index=False)
    third, _ = cli.update(root, csv=path, offline=True)
    assert third["status"] == "failed"
    assert third["validation"]["errors"] and third["sources"]
    assert active.read_bytes() == original_bytes


def test_refresh_restores_previous_sources_when_promotion_fails(tmp_path):
    from types import SimpleNamespace
    from utils.ifdata_cache.financial_analysis_update import promote_legacy
    import json
    class Cache:
        def __init__(self, directory, value, fail=False):
            directory.mkdir(parents=True)
            self.arquivo_dados = directory / "dados.parquet"
            self.arquivo_metadata = directory / "metadata.json"
            self.fail = fail
            pd.DataFrame({"x": [value]}).to_parquet(self.arquivo_dados)
            self.arquivo_metadata.write_text(json.dumps({"fonte": "api"}))
        def existe_leitura(self):
            return True
        def salvar_local(self, data, **kwargs):
            if self.fail:
                return SimpleNamespace(sucesso=False, mensagem="simulated I/O failure")
            data.to_parquet(self.arquivo_dados)
            return SimpleNamespace(sucesso=True)
    old = {"a": Cache(tmp_path / "old/a", 1), "b": Cache(tmp_path / "old/b", 2, fail=True)}
    new = {"a": Cache(tmp_path / "new/a", 10), "b": Cache(tmp_path / "new/b", 20)}
    manager = SimpleNamespace(base_dir=tmp_path, get_cache=old.get)
    with pytest.raises(ValueError, match="simulated"):
        promote_legacy(manager, SimpleNamespace(get_cache=new.get), ["a", "b"])
    assert pd.read_parquet(old["a"].arquivo_dados).x.iloc[0] == 1
    assert pd.read_parquet(old["b"].arquivo_dados).x.iloc[0] == 2



def test_presentation_keeps_negative_small_zero_and_missing_distinct():
    from utils.prudential_analysis import Cell
    from utils.prudential_presentation import format_cell, value_tone
    money = "#,##0;[Red]-#,##0;0"
    assert [format_cell(Cell(v, fmt=money)) for v in (0, 9.5e-13, -.3, .3, -1234.4, "N/D")] == ["0", "0", "-<1", "<1", "-1.234", "–"]
    assert [value_tone(Cell(v, fmt=money)) for v in (0, 9.5e-13, -.3, -5, 5, "N/D")] == ["zero", "zero", "negative", "negative", "", "missing"]
    assert format_cell(Cell(-.001, fmt="0%;[Red]-0%;0%")) == "0%"
    assert format_cell(Cell(.003, fmt="0.00%")) == "0,30%"


def test_committee_panels_end_together_and_long_selection_pages_keep_pairs(simple_sources):
    from utils.prudential_presentation import period_chunks
    assert [len(c) for c in period_chunks(list("abcdefgh"))] == [4, 4]
    assert [len(c) for c in period_chunks(list("abcde"))] == [3, 2]
    periods = ["202303", "202306", "202312", "202406", "202412", "202506", "202512", "202606"]
    analysis = build_analysis(pd.concat([simple_sources.assign(DATA_BASE=p) for p in periods]), {}, "C1", periods)
    workbook = load_workbook(BytesIO(render_xlsx(analysis)))
    committee = workbook["Comitê"]
    assert committee["H7"].value == analysis.sheets["DRE"].cells["B7"].value
    balance = workbook["Balanço"]
    assert balance.print_title_cols == "$A:$B" and "$2:$6" in balance.print_title_rows
    assert [b.id for b in balance.col_breaks.brk] == [10]  # Four periods per page; 2 + 2*4 columns precede the break.


def test_slides_follow_accounting_groups_and_never_open_on_a_subaccount(simple_sources):
    from utils.prudential_pptx import paginate_rows, MAX_ROWS
    analysis = build_analysis(simple_sources, {}, "C1", ["202606"])
    for name in ("Balanço", "DRE", "Carteira 4.966"):
        for _, rows in paginate_rows(analysis, name):
            assert len(rows) <= MAX_ROWS
            assert analysis.sheets[name].cells[f"B{rows[0][1]}"].style not in {"indent", "detail"}


def test_default_periods_are_last_four_closed_consecutive_quarters():
    from datetime import date
    from utils.prudential_presentation import default_periods
    monthly = [f"{y}{m:02d}" for y in (2024, 2025, 2026) for m in range(1, 13) if f"{y}{m:02d}" <= "202607"]
    assert default_periods(monthly, date(2026, 10, 31)) == ["202606", "202603", "202512", "202509"]
    # Fechou hoje: conta. Ainda aberto: não conta, mesmo com o mês publicado.
    assert default_periods(monthly, date(2026, 6, 30))[0] == "202606"
    assert default_periods(monthly, date(2026, 6, 29))[0] == "202603"
    # Lacuna da instituição não puxa trimestres fora da janela consecutiva.
    assert default_periods(["202606", "202512", "202509"], date(2026, 10, 31)) == ["202606", "202512", "202509"]
    assert default_periods(["202604", "202605"], date(2026, 10, 31)) == ["202605", "202604"]


def test_workbook_prints_unavailable_as_dash_but_keeps_the_value(simple_sources):
    a = build_analysis(simple_sources, {}, "C1", ["202606"])
    wb = load_workbook(BytesIO(render_xlsx(a)), data_only=True)
    ws = wb["Balanço"]
    assert any(cell.value == "N/D" and cell.number_format.endswith(';"–"') for row in ws.iter_rows(min_row=7) for cell in row)


def test_nim_uses_average_with_prior_semester_end_and_annualizes(simple_sources):
    from utils.prudential_nim import add_nim_rows, credit_balances, prior_periods, prior_semester_end
    assert [prior_semester_end(p) for p in ("202603", "202606", "202609", "202612")] == ["202512", "202512", "202606", "202606"]
    periods = ["202512", "202606"]
    assert prior_periods(periods) == ["202506"]
    analysis = build_analysis(pd.concat([simple_sources.assign(DATA_BASE=p) for p in periods]), {}, "C1", periods)
    gross_before = analysis.sheets["DRE"].cells[f"C{analysis.rows['DRE']['gross']}"].value
    add_nim_rows(analysis, {"202506": 40.0})
    sheet, rows = analysis.sheets["DRE"], analysis.rows["DRE"]
    assert rows["nim_gross"] == rows["net"] + 1 and rows["nim_after_pdd"] == rows["net"] + 2
    assert sheet.cells[f"C{rows['gross']}"].value == gross_before
    credit = credit_balances(analysis)
    gross = sheet.cells[f"E{rows['gross']}"].value
    if credit.get("202606") and isinstance(gross, (int, float)) and credit["202606"] + credit["202512"] > 0:
        expected = gross / ((credit["202606"] + credit["202512"]) / 2) * 2
        assert sheet.cells[f"E{rows['nim_gross']}"].value == pytest.approx(expected)
    wb = load_workbook(BytesIO(render_xlsx(analysis)))
    assert "AVERAGE" in wb["DRE"][f"E{rows['nim_gross']}"].value
