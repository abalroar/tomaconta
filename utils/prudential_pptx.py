"""Editable financial statements using Snapshot's native PowerPoint tables."""
from __future__ import annotations

from io import BytesIO
import json

from .financial_statement_style import BAND, BAND_MUTED, BAND_NEGATIVE, MISSING, NEGATIVE, SECTION
from .presentation_tables import pin_run_fonts, _slide, _table, _cell, _font, _text, _WIDTH, _HEIGHT, _MARGIN, _CONTENT, _MUTED
from .prudential_presentation import (PAGE_GROUPS, committee_panel, reference_note, display_rows, format_cell, institution_label,
                                      period_cells, period_label, period_source, row_tier, value_tone, with_band_gaps)

MAX_ROWS = 21
SLIDE_PERIODS = 3
TABLE_TOP = 1.24
HEADER_HEIGHT = .50
BODY_HEIGHT = 5.20
ROW_FLOOR = BODY_HEIGHT / MAX_ROWS
ROW_CEILING = .38
# Amount and share of a period read as one block, whatever the number of periods.
PERIOD_WIDTH = {1: 3.2, 2: 2.8, 3: 2.4, 4: 2.05}
PERCENT_WIDTH = .62
CELL_PADDING = .015
ZERO_COLOR = "8E949B"
# Balanço e DRE lado a lado, como na página de comitê.
PANEL_GAP = .25
PANEL_TOP = 1.12
PANEL_HEADER = .36
PANEL_BODY = 5.3
PANEL_AMOUNT, PANEL_SHARE = .72, .38
PANEL_ROW_MAX = .165
PANEL_GAP_ROW = .045


def row_height(label, style):
    # Calibri 12pt: only a label longer than the narrowest column takes a second line.
    width = 62 - (3 if style == "indent" else 6 if style == "detail" else 0)
    return ROW_FLOOR if len(label) <= width else ROW_FLOOR * 1.7


def _accounting_groups(analysis, name):
    rows, groups, start = list(display_rows(analysis, name)), [], 0
    for title, last in PAGE_GROUPS[name]:
        end = next((i + 1 for i, row in enumerate(rows) if row[0] == last), len(rows))
        if rows[start:end]:
            groups.append((title, rows[start:end]))
        start = end
    return groups


def _paginate_group(sheet, rows):
    pages, page, height = [], [], 0
    for index, row in enumerate(rows):
        style = sheet.cells[f"B{row[1]}"].style
        h = row_height(row[2], style)
        # A parent account travels with its subaccounts, so no slide opens on them.
        reserve, following_count = 0, 0
        if style not in {"indent", "detail"}:
            for following in rows[index + 1:]:
                next_style = sheet.cells[f"B{following[1]}"].style
                if next_style not in {"indent", "detail"}:
                    break
                reserve += row_height(following[2], next_style)
                following_count += 1
        if reserve + h > BODY_HEIGHT or following_count >= MAX_ROWS:
            reserve, following_count = 0, 0
        if page and (len(page) + 1 + following_count > MAX_ROWS or height + h + reserve > BODY_HEIGHT):
            pages.append(page)
            page, height = [], 0
        page.append(row)
        height += h
    if page:
        pages.append(page)
    return pages


def paginate_rows(analysis, name):
    """Slides of one statement as (accounting group title, rows)."""
    sheet = analysis.sheets[name]
    slides = []
    for title, rows in _accounting_groups(analysis, name):
        pages = _paginate_group(sheet, rows)
        slides += [(title + (f" ({part}/{len(pages)})" if len(pages) > 1 else ""), page)
                   for part, page in enumerate(pages, 1)]
    return slides


def _edge(cell, side, color, width=6350):
    """Restyle one border already created by the shared cell helper."""
    node = cell._tc.get_or_add_tcPr().find(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{side}")
    node.set("w", str(width))
    node.find("*/*").set("val", color.lstrip("#"))


def _value_color(cell, default, *, band=False):
    negative, muted = (BAND_NEGATIVE, BAND_MUTED) if band else (NEGATIVE, ZERO_COLOR)
    return {"negative": negative, "missing": muted, "zero": muted}.get(value_tone(cell), default)


def _statement_table(slide, analysis, name, rows, periods, offset):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches

    sheet = analysis.sheets[name]
    block = PERIOD_WIDTH[len(periods)]
    widths = [_CONTENT - block * len(periods)] + [w for _ in periods for w in (block - PERCENT_WIDTH, PERCENT_WIDTH)]
    styles = [sheet.cells[f"B{r}"].style for _, r, _ in rows]
    base = min(ROW_CEILING, BODY_HEIGHT / max(1, sum(row_height(row[2], s) / ROW_FLOOR for row, s in zip(rows, styles))))
    heights = [HEADER_HEIGHT] + [base * row_height(row[2], s) / ROW_FLOOR for row, s in zip(rows, styles)]
    table = _table(slide, len(rows) + 1, 1 + 2 * len(periods), TABLE_TOP, heights, widths)
    _cell(table.cell(0, 0), "R$ milhões", header=True, size=12)
    for i, period in enumerate(periods):
        cell = table.cell(0, 1 + 2 * i)
        cell.merge(table.cell(0, 2 + 2 * i))
        _cell(cell, period_label(period), header=True, centered=True, size=12)
        paragraph = cell.text_frame.add_paragraph()
        paragraph.text = period_source(sheet, offset + i)
        _font(paragraph, 9, color="FFFFFF", centered=True)
    for target, ((ident, source_row, label), style) in enumerate(zip(rows, styles), 1):
        tier = row_tier(name, ident, style)
        band = tier == "key"
        bold = tier in {"key", "section", "total", "bold"}
        background = BAND if band else SECTION if tier in {"section", "total"} else None
        size = 11 if tier == "detail" else 12
        label_cell = table.cell(target, 0)
        _cell(label_cell, label, size=size, bold=bold, background=background,
              color="FFFFFF" if band else _MUTED if tier == "detail" else "24282E")
        label_cell.margin_left = Inches(.09 + {"indent": .16, "detail": .30}.get(tier, 0))
        label_cell.margin_top = label_cell.margin_bottom = Inches(CELL_PADDING)
        label_cell.text_frame.paragraphs[0].font.italic = tier in {"detail", "ratio"}
        for i, _ in enumerate(periods):
            amount, share = period_cells(sheet, source_row, offset + i)
            for j, cell in enumerate((amount, share)):
                target_cell = table.cell(target, 1 + 2 * i + j)
                default = "FFFFFF" if band else _MUTED if j else "24282E"
                _cell(target_cell, format_cell(cell), size=10 if j else size, bold=bold,
                      color=_value_color(cell, default, band=band), background=background)
                target_cell.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                target_cell.margin_left = Inches(.02)
                target_cell.margin_top = target_cell.margin_bottom = Inches(CELL_PADDING)
                if j == 0 and not band:
                    _edge(target_cell, "lnL", "E8EAED")
    return table


def _panel(slide, analysis, name, rows, periods, offset, x, label_width, height):
    """One compact statement of the committee page; thin rows separate its bands."""
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    sheet = analysis.sheets[name]
    widths = [label_width] + [w for _ in periods for w in (PANEL_AMOUNT, PANEL_SHARE)]
    heights = [PANEL_HEADER] + [height if row is not None else PANEL_GAP_ROW for row in rows]
    table = _table(slide, len(rows) + 1, 1 + 2 * len(periods), PANEL_TOP, heights, widths)
    # The shared helper anchors tables at the margin; each panel takes its own position.
    slide.shapes[-1].left, slide.shapes[-1].width = Inches(x), Inches(sum(widths))
    _cell(table.cell(0, 0), f"{name} · R$ milhões", header=True, size=10, background=BAND)
    for i, period in enumerate(periods):
        cell = table.cell(0, 1 + 2 * i)
        cell.merge(table.cell(0, 2 + 2 * i))
        _cell(cell, period_label(period), header=True, centered=True, size=10, background=BAND)
    for target, row in enumerate(rows, 1):
        cells = [table.cell(target, c) for c in range(1 + 2 * len(periods))]
        if row is None:
            for cell in cells:
                _cell(cell, "", size=1)
                cell.margin_top = cell.margin_bottom = 0
                _edge(cell, "lnB", "FFFFFF", 0)
                # The row's floor comes from the empty paragraph's own font size.
                paragraph = cell.text_frame.paragraphs[0]._p
                end = paragraph.makeelement(qn("a:endParaRPr"), {"sz": "100"})
                paragraph.append(end)
            continue
        ident, source_row, label, style, tier = row
        band = tier == "key"
        bold = tier in {"key", "section", "total", "bold"}
        background = BAND if band else SECTION if tier in {"section", "total"} else None
        _cell(cells[0], label, size=9, bold=bold, background=background,
              color="FFFFFF" if band else _MUTED if tier == "detail" else "24282E")
        cells[0].margin_left = Inches(.05 + {"indent": .1, "detail": .18}.get(tier, 0))
        cells[0].margin_right = Inches(.02)
        cells[0].margin_top = cells[0].margin_bottom = 0
        cells[0].text_frame.paragraphs[0].font.italic = tier in {"detail", "ratio"}
        for i, _ in enumerate(periods):
            amount, share = period_cells(sheet, source_row, offset + i)
            # Balance subaccounts carry amounts only, as in the committee model.
            if name == "Balanço" and (style in {"indent", "detail"} or ident == "net"):
                share = None
            for j, cell in enumerate((amount, share)):
                target_cell = cells[1 + 2 * i + j]
                default = "FFFFFF" if band else _MUTED if j else "24282E"
                _cell(target_cell, format_cell(cell), size=8 if j else 9, bold=bold,
                      color=_value_color(cell, default, band=band), background=background)
                target_cell.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
                target_cell.margin_left = target_cell.margin_right = Inches(.03)
                target_cell.margin_top = target_cell.margin_bottom = 0


def _committee_slide(prs, analysis, institution):
    """Balanço e DRE dos últimos períodos em uma página; o detalhe completo fica no Excel."""
    periods = analysis.periods[-SLIDE_PERIODS:]
    offset = len(analysis.periods) - len(periods)
    span = " / ".join(period_label(p) for p in periods)
    slide = _slide(prs, "Balanço e DRE", f"{institution}  ·  Conglomerado prudencial  ·  {span}", len(prs.slides) + 1)
    panels = {name: committee_panel(analysis, name) for name in ("Balanço", "DRE")}
    # The balance sheet is continuous; the income statement breathes only around its results.
    panels["DRE"] = with_band_gaps(panels["DRE"])
    fixed = max(sum(row is not None for row in rows) for rows in panels.values())
    height = min(PANEL_ROW_MAX, PANEL_BODY / fixed)
    panel_width = (_CONTENT - PANEL_GAP) / 2
    label_width = panel_width - len(periods) * (PANEL_AMOUNT + PANEL_SHARE)
    x = _MARGIN
    for name in ("Balanço", "DRE"):
        _panel(slide, analysis, name, panels[name], periods, offset, x, label_width, height)
        x += panel_width + PANEL_GAP
    # A ressalva, quando existe, entra na mesma área do rodapé, sem encostar nas tabelas.
    _text(slide, _MARGIN, 6.88 if reference_note(analysis) else 7.04, _CONTENT, .45,
          f"Fonte: Banco Central, COSIF 4060/4066 e IF.data. R$ milhões. DRE acumulada no ano. {MISSING}: dado não disponível. {reference_note(analysis)}".rstrip(), 9, color=_MUTED)
    slide.notes_slide.notes_text_frame.text = json.dumps({
        "institution": analysis.code, "periods": periods, "statement": "Balanço e DRE",
        "rows": {name: [row[0] for row in rows if row] for name, rows in panels.items()},
        "methodology": analysis.limitations,
    }, ensure_ascii=False, default=str)


def _portfolio_statement_slides(prs, analysis, institution):
    """Without the Carteira 4.966 menu model, the calculated statement keeps the deck complete."""
    name = "Carteira 4.966"
    periods = analysis.periods[-SLIDE_PERIODS:]
    offset = len(analysis.periods) - len(periods)
    span = " / ".join(period_label(p) for p in periods)
    for group, rows in paginate_rows(analysis, name):
        slide = _slide(prs, f"{name} · {group}", f"{institution}  ·  Conglomerado prudencial  ·  {span}", len(prs.slides) + 1)
        _statement_table(slide, analysis, name, rows, periods, offset)
        _text(slide, _MARGIN, 7.04, _CONTENT, .22,
              f"Fonte: Banco Central, COSIF 4060/4066 e IF.data. R$ milhões. {MISSING}: dado não disponível.", 9, color=_MUTED)
        slide.notes_slide.notes_text_frame.text = json.dumps({
            "institution": analysis.code, "periods": periods, "statement": name,
            "rows": [row[0] for row in rows], "methodology": analysis.limitations,
        }, ensure_ascii=False, default=str)


def committee_slides(prs, analysis, evolution=None):
    """Balanço, DRE e Evolução para compor o material de comitê do Snapshot."""
    from .peers_table_model import short_bank
    _committee_slide(prs, analysis, short_bank(institution_label(analysis.institution)))
    if evolution is not None:
        from .evolucao_pptx_export import add_evolucao_slides
        add_evolucao_slides(prs, **evolution)


def render_pptx(analysis, *, portfolio=None, evolution=None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(_WIDTH), Inches(_HEIGHT)
    institution = institution_label(analysis.institution)
    prs.core_properties.title = f"{institution} · Análise financeira"
    prs.core_properties.author = "Toma Conta"
    _committee_slide(prs, analysis, institution)
    if portfolio is not None and portfolio.periods:
        from .snapshot_pptx_export import _carteira_slides
        _carteira_slides(prs, {"bank": institution, "base": "Conglomerado prudencial"}, portfolio, "")
    else:
        _portfolio_statement_slides(prs, analysis, institution)
    if evolution is not None:
        from .evolucao_pptx_export import add_evolucao_slides
        add_evolucao_slides(prs, **evolution)
    pin_run_fonts(prs)
    output = BytesIO()
    prs.save(output)
    return output.getvalue()
