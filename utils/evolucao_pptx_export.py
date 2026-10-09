"""Native Office export of the exact institution/window displayed in Evolução.

The application uses the same chart/OOXML engine as Peers and Crédito BC.
Values remain editable in the embedded workbook; missing values remain gaps.
"""

from io import BytesIO
import json
import math

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from .evolucao_visual import HEADER_COLOR, METRIC_LABELS, SERIES, _axis_range, period_label
from .scr_pptx_export import (
    ORDEM_CHART_SPACE, ORDEM_DLBL, ORDEM_LEGENDA, _estilizar_eixos,
    _layout_do_rotulo, _mover_para_eixo_secundario, _ordenar_filhos,
    _rotular_apenas_ultimo_ponto, _sem_preenchimento_nem_contorno,
    _texto_do_eixo,
)


PERIODS_PER_SLIDE = 8
NUMBER_FORMAT = '[$-416]#,##0.0'
CHART_TOP, CHART_HEIGHT = .94, 3.48
PLOT_LEFT, PLOT_WIDTH, PLOT_TOP, PLOT_HEIGHT = .065, .865, .22, .64
LABEL_GAP, LABEL_HEIGHT = .14, .12
TABLE_TOP = 4.53


def _text(slide, text, x, y, w, h, *, size=12, bold=False, color="222222"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for i, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name, paragraph.font.size, paragraph.font.bold = "Calibri", Pt(size), bold
        paragraph.font.color.rgb = RGBColor.from_string(color)
        paragraph.space_before = paragraph.space_after = Pt(0)
    return shape


def _manual_layout(parent, *, x, y, w=None, h=None, inner=False):
    layout = parent.find(qn("c:layout"))
    if layout is None:
        layout = OxmlElement("c:layout")
        parent.insert(0, layout)
    else:
        layout.clear()
    manual = OxmlElement("c:manualLayout")
    layout.append(manual)
    if inner:
        target = OxmlElement("c:layoutTarget")
        target.set("val", "inner")
        manual.append(target)
    for tag, value in (("xMode", "edge"), ("yMode", "edge"), ("wMode", "factor"), ("hMode", "factor"),
                       ("x", x), ("y", y), ("w", w), ("h", h)):
        if value is not None:
            node = OxmlElement("c:" + tag)
            node.set("val", str(value))
            manual.append(node)


def _label_positions(points):
    """Separate final labels across both axes in the same chart coordinates."""
    ordered = sorted(points, key=lambda point: point[1])
    positions = []
    for _, target in ordered:
        positions.append(max(target, positions[-1] + LABEL_GAP if positions else PLOT_TOP + .01))
    if positions:
        shift = max(0, positions[-1] - (PLOT_TOP + PLOT_HEIGHT - LABEL_HEIGHT))
        positions = [position - shift for position in positions]
    return {index: position for (index, _), position in zip(ordered, positions)}


def _chart(slide, frame, periods):
    data = CategoryChartData()
    data.categories = [period_label(period) for period in periods]
    values = {}
    for column, name, _, _, _ in SERIES:
        numbers = pd.to_numeric(frame[column], errors="coerce")
        values[column] = [float(v) / 1e9 if pd.notna(v) and math.isfinite(float(v)) else None for v in numbers]
        data.add_series(name, values[column], number_format=NUMBER_FORMAT)
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.38), Inches(CHART_TOP), Inches(12.55), Inches(CHART_HEIGHT), data,
    ).chart
    chart.has_title = False
    # python-pptx's template uses signed axis IDs; OOXML requires unsigned IDs.
    # Normalize references together before the shared engine adds the second pair.
    axis_ids = list(dict.fromkeys(chart._chartSpace.xpath(".//c:axId/@val")))
    replacement_ids = {old: str(100_000_000 + i) for i, old in enumerate(axis_ids)}
    for node in chart._chartSpace.xpath(".//c:axId | .//c:crossAx"):
        node.set("val", replacement_ids[node.get("val")])
    _sem_preenchimento_nem_contorno(chart._chartSpace, ORDEM_CHART_SPACE)
    ranges = {
        axis: _axis_range([v for column, _, _, series_axis, _ in SERIES if series_axis == axis
                          for v in values[column] if v is not None])
        for axis in ("y", "y2")
    }
    _estilizar_eixos(chart, len(periods), NUMBER_FORMAT, escala=ranges["y"])
    chart.plots[0].gap_width = 70
    chart.plots[0].has_data_labels = False
    finals = []
    for index, (column, _, color, axis, _) in enumerate(SERIES):
        series = chart.series[index]
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(color.lstrip("#"))
        series.format.line.color.rgb = RGBColor.from_string(color.lstrip("#"))
        series.format.line.width = Pt(2)
        valid = [i for i, v in enumerate(values[column]) if v is not None]
        if valid:
            last = valid[-1]
            _rotular_apenas_ultimo_ponto(series, last, NUMBER_FORMAT, com_nome=False)
            low, high = ranges[axis]
            target = PLOT_TOP + (1 - (values[column][last] - low) / (high - low)) * PLOT_HEIGHT - .08
            finals.append((index, target, last))
    _mover_para_eixo_secundario(chart, [2, 3], NUMBER_FORMAT, escala=ranges["y2"], cor=RGBColor.from_string("222222"))
    positions = _label_positions([(i, y) for i, y, _ in finals])
    for index, _, last in finals:
        label = chart.series[index].points[last].data_label
        # Keep labels inside the plot near their actual last valid category.
        # A small inset protects the right-axis ticks from long numeric labels.
        point_x = PLOT_LEFT + PLOT_WIDTH * (last + .5) / len(periods)
        label_x = min(PLOT_LEFT + PLOT_WIDTH - .115, max(PLOT_LEFT + .015, point_x - .045))
        _layout_do_rotulo(label, label_x, positions[index])
        label.font.name, label.font.size, label.font.bold = "Calibri", Pt(12), True
        label.font.color.rgb = RGBColor.from_string(SERIES[index][2].lstrip("#"))
        # A series ending before the window end must identify its valid date.
        label._dLbl.find(qn("c:showCatName")).set("val", "1" if last < len(periods) - 1 else "0")
        separator = OxmlElement("c:separator")
        separator.text = "\n"
        label._dLbl.append(separator)
        # A white backing keeps labels legible when a line passes behind them.
        background = label._dLbl.find(qn("c:spPr"))
        if background is None:
            background = OxmlElement("c:spPr")
            label._dLbl.append(background)
        else:
            background.clear()
        fill = OxmlElement("a:solidFill")
        color = OxmlElement("a:srgbClr")
        color.set("val", "FFFFFF")
        fill.append(color)
        background.append(fill)
        outline = OxmlElement("a:ln")
        outline.append(OxmlElement("a:noFill"))
        background.append(outline)
        _ordenar_filhos(label._dLbl, ORDEM_DLBL)
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.TOP
    chart.legend.include_in_layout = False
    chart.legend.font.name, chart.legend.font.size = "Calibri", Pt(12)
    _manual_layout(chart.legend._element, x=.035, y=.005, w=.93, h=.095)
    for index, (_, _, color, _, _) in enumerate(SERIES):
        entry = OxmlElement("c:legendEntry")
        idx = OxmlElement("c:idx")
        idx.set("val", str(index))
        entry.append(idx)
        entry.append(_texto_do_eixo(RGBColor.from_string(color.lstrip("#"))))
        chart.legend._element.append(entry)
    _ordenar_filhos(chart.legend._element, ORDEM_LEGENDA)
    _manual_layout(chart._chartSpace.chart.plotArea, x=PLOT_LEFT, y=PLOT_TOP, w=PLOT_WIDTH, h=PLOT_HEIGHT, inner=True)
    _text(slide, "Lucro e PL (R$ bi)", .38, 1.43, 5.9, .2, size=10, color="555555")
    axis_caption = _text(slide, "Carteira e Core Funding (R$ bi)", 7.03, 1.43, 5.9, .2, size=10, color="222222")
    axis_caption.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    blanks = chart._chartSpace.chart.find(qn("c:dispBlanksAs"))
    if blanks is None:
        blanks = OxmlElement("c:dispBlanksAs")
        chart._chartSpace.chart.append(blanks)
    blanks.set("val", "gap")
    # Explicit fonts apply to the secondary axis as well as the primary axes.
    for props in chart._chartSpace.xpath(".//a:rPr | .//a:defRPr"):
        props.set("lang", "pt-BR")
        props.set("sz", str(max(1100, int(props.get("sz", "1100")))))
        latin = props.find(qn("a:latin"))
        if latin is None:
            latin = OxmlElement("a:latin")
            props.append(latin)
        latin.set("typeface", "Calibri")
    for axis in (chart.category_axis, chart.value_axis):
        axis.tick_labels.font.name, axis.tick_labels.font.size = "Calibri", Pt(11)


def _table(slide, frame, periods):
    table = slide.shapes.add_table(
        len(frame) + 1, len(periods) + 1, Inches(.38), Inches(TABLE_TOP), Inches(12.55), Inches(.34 + .32 * len(frame)),
    ).table
    table.columns[0].width = Inches(3.5)
    for column in range(1, len(periods) + 1):
        table.columns[column].width = Inches(9.05 / len(periods))
    table.rows[0].height = Inches(.34)
    for row in list(table.rows)[1:]:
        row.height = Inches(.32)
    table.cell(0, 0).text = "Indicador"
    for c, period in enumerate(periods, 1):
        table.cell(0, c).text = period_label(period)
    for r, (_, row) in enumerate(frame.iterrows(), 1):
        metric = str(row["Métrica"])
        table.cell(r, 0).text = METRIC_LABELS.get(metric, metric)
        for c, period in enumerate(periods, 1):
            value = str(row[period])
            table.cell(r, c).text = "N/D" + value[1:] if value.startswith("-") and value in ("-", "-†") else value
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(HEADER_COLOR.lstrip("#") if r == 0 else "FFFFFF")
            cell.margin_left = cell.margin_right = Inches(.08)
            cell.margin_top = cell.margin_bottom = Inches(.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.name, paragraph.font.size = "Calibri", Pt(12)
                paragraph.font.bold = r == 0 or (r > 0 and c == len(periods))
                paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if r == 0 else "222222")
                paragraph.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER if r == 0 else PP_ALIGN.RIGHT
                paragraph.space_before = paragraph.space_after = Pt(0)
            properties = cell._tc.get_or_add_tcPr()
            border = OxmlElement("a:lnB")
            border.set("w", "6350")
            fill = OxmlElement("a:solidFill")
            color = OxmlElement("a:srgbClr")
            color.set("val", HEADER_COLOR.lstrip("#") if r == 0 else "E5E5E5")
            fill.append(color)
            border.append(fill)
            # CT_TableCellProperties: borders precede the cell fill.
            properties.insert(0, border)


def export_evolucao_powerpoint(*, df_graph, df_show, periods, institution, queried_at, status_rows=()):
    """One slide with native chart and indicators per eight displayed periods."""
    periods = list(periods)
    if not periods or len(df_graph) != len(periods) or len(set(periods)) != len(periods):
        raise ValueError("A janela do gráfico e da tabela deve ter as mesmas competências, sem duplicatas.")
    if df_show.empty or "Métrica" not in df_show or any(period not in df_show for period in periods):
        raise ValueError("A tabela de indicadores está incompleta.")
    if any(column not in df_graph for column, *_ in SERIES):
        raise ValueError("As quatro séries da Evolução são obrigatórias.")
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    prs.core_properties.title = "Evolução | " + str(institution)
    prs.core_properties.author = "Toma Conta"
    query_type = "Evolução histórica; dezembro e competência final; lucro acumulado no ano"
    footer = (f"Fonte: BCB IFData, relatórios 1, 2, 3 e 5. Consolidada / Prudencial. 1 instituição. "
              f"{period_label(periods[0])} a {period_label(periods[-1])}.\n"
              f"Consulta: {queried_at}. Tipo: {query_type}.")
    notes = json.dumps({"instituicao": institution, "competencias": periods, "consulta": queried_at,
                        "tipo_consulta": query_type, "status_analitico": list(status_rows)}, ensure_ascii=False, default=str)
    for start in range(0, len(periods), PERIODS_PER_SLIDE):
        block = periods[start:start + PERIODS_PER_SLIDE]
        block_status = [row for row in status_rows if period_label(str(row.get("Período", ""))) in {period_label(p) for p in block}]
        caveats = []
        for status, description in (("fallback_capitacoes_view", "Core Funding em fallback para Captações"),
                                    ("fallback_net_components", "Carteira em fallback para componentes líquidos e+f+g+h")):
            dates = list(dict.fromkeys(period_label(row["Período"]) for row in block_status if row.get("Status analítico") == status))
            if dates:
                caveats.append(description + ": " + ", ".join(dates) + ".")
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        _text(slide, "Evolução", .38, .2, 12.55, .35, size=22, bold=True)
        _text(slide, f"{institution} | {period_label(block[0])} a {period_label(block[-1])}", .38, .63, 12.55, .25, size=13)
        _chart(slide, df_graph.iloc[start:start + len(block)], block)
        _table(slide, df_show, block)
        note = "Lucro acumulado no ano. * Mudança do layout IFData em 2025 nas séries Carteira e Core Funding."
        if any("*" in str(df_show[p].tolist()) for p in block):
            note += " * Nos indicadores: fallback analítico."
        if any("†" in str(df_show[p].tolist()) for p in block):
            note += " † Indisponibilidade com causa identificada."
        _text(slide, "\n".join([note, *caveats]), .38, 6.52, 12.55, .46, size=9, color="555555")
        _text(slide, footer, .38, 7.03, 12.55, .32, size=8, color="555555")
        slide.notes_slide.notes_text_frame.text = notes
    output = BytesIO()
    prs.save(output)
    return output.getvalue()
