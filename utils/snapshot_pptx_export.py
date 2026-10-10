"""Resumo do Snapshot em PowerPoint, com tabelas e gráficos Office editáveis.

O exportador recebe os valores e as comparações já fechados pela tela. Peers e
4.966 usam seus próprios modelos de leitura; este módulo só compõe os slides.
Nenhuma tabela, mini-histórico ou indicador é convertido em imagem.
"""
from __future__ import annotations

from io import BytesIO
import json
import math
from collections.abc import Mapping

from .comparison_table_style import HEADER_BACKGROUND, VARIATION_COLORS
from .peers_table_model import (
    get_metric, methodology_rows, period_comparison_label, period_label,
    period_sort, short_bank, variation_tone,
)


_WIDTH, _HEIGHT = 13.333, 7.5
_MARGIN, _CONTENT = .42, 12.493
_INK, _MUTED, _GRID, _SECTION = "24282E", "626970", "E8EAED", "F3F5F6"


def _json(value) -> str:
    return json.dumps(value, ensure_ascii=False, default=str, indent=2)


def _font(paragraph, size=14, *, bold=False, color=_INK, centered=False):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    paragraph.font.name = "Calibri"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(str(color).lstrip("#"))
    paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
    paragraph.space_before = paragraph.space_after = Pt(0)


def _text(slide, x, y, width, height, text, size=14, *, bold=False, color=_INK):
    from pptx.util import Inches
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    for index, line in enumerate(str(text).split("\n")):
        paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
        paragraph.text = line
        _font(paragraph, size, bold=bold, color=color)
    return box


def _slide(prs, title, subtitle, page):
    from pptx.dml.color import RGBColor
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    _text(slide, _MARGIN, .24, _CONTENT - .5, .48, title, 25, bold=True)
    _text(slide, _MARGIN, .84, _CONTENT, .4, subtitle, 12, color=_MUTED)
    _text(slide, 12.53, .33, .36, .3, str(page), 10, color=_MUTED)
    return slide


def _borders(cell, *, header=False):
    """Grades discretas, sem faixas ou separadores verticais entre grupos."""
    from pptx.oxml.xmlchemy import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    for name in ("lnL", "lnR", "lnT", "lnB"):
        for existing in tc_pr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{name}"):
            tc_pr.remove(existing)
    # DrawingML exige as bordas antes do preenchimento da célula.
    for index, name in enumerate(("lnL", "lnR", "lnT", "lnB")):
        edge = OxmlElement(f"a:{name}")
        edge.set("w", "6350" if name == "lnB" and not header else "0")
        fill = OxmlElement("a:solidFill")
        color = OxmlElement("a:srgbClr")
        color.set("val", _GRID if not header else HEADER_BACKGROUND.lstrip("#"))
        fill.append(color)
        edge.append(fill)
        tc_pr.insert(index, edge)


def _cell(cell, text, *, size=14, bold=False, header=False, centered=False, color=_INK, background=None):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches
    cell.text = str(text or "")
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string((background or (HEADER_BACKGROUND if header else "FFFFFF")).lstrip("#"))
    cell.margin_left = cell.margin_right = Inches(.09)
    cell.margin_top = cell.margin_bottom = Inches(.035)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text_frame.word_wrap = True
    _borders(cell, header=header)
    for paragraph in cell.text_frame.paragraphs:
        _font(paragraph, size, bold=bold or header, color="FFFFFF" if header else color, centered=centered)


def _cell_delta(cell, display, tone, *, size=10):
    paragraph = cell.text_frame.add_paragraph()
    paragraph.text = str(display or "")
    _font(paragraph, size, color=VARIATION_COLORS.get(tone, VARIATION_COLORS["neutral"]), centered=True)


def _table(slide, rows, columns, y, heights, widths):
    from pptx.util import Inches
    table = slide.shapes.add_table(rows, columns, Inches(_MARGIN), Inches(y), Inches(_CONTENT), Inches(sum(heights))).table
    for row, height in zip(table.rows, heights):
        row.height = Inches(height)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    return table


def _comparison(card, key):
    value = card.get(key)
    if isinstance(value, Mapping):
        return str(value.get("display") or "—"), str(value.get("tone") or "neutral")
    return str(value or "—"), str(card.get(f"{key}_tone") or "neutral")


def _history(slide, history, x, y, width, height):
    """Minigráfico nativo com planilha incorporada e lacunas preservadas."""
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_MARKER_STYLE
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Inches, Pt

    if not isinstance(history, Mapping):
        return False
    values, categories = list(history.get("values") or ()), list(history.get("periods") or ())
    if len(values) != len(categories):
        raise ValueError("O histórico do Snapshot exige um rótulo por valor.")
    numbers = []
    for value in values:
        try:
            number = float(value) if value is not None else None
        except (TypeError, ValueError):
            number = None
        numbers.append(number if number is not None and math.isfinite(number) else None)
    if not any(value is not None for value in numbers):
        return False
    data = CategoryChartData()
    data.categories = [str(category) for category in categories]
    data.add_series("Evolução", numbers)
    bars = history.get("kind") == "bars"
    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED if bars else XL_CHART_TYPE.LINE,
        Inches(x), Inches(y), Inches(width), Inches(height), data,
    ).chart
    # IDs são locais ao gráfico. O Office Mac rejeita alguns IDs acima de
    # Int32 mesmo quando cabem em UInt32; usa IDs positivos pequenos e mantém
    # todas as referências do gráfico e dos eixos no mesmo mapeamento.
    axis_ids = {
        node.get("val"): str(index + 1)
        for index, node in enumerate(chart._chartSpace.xpath(".//c:catAx/c:axId | .//c:valAx/c:axId"))
    }
    for axis_id in chart._chartSpace.xpath(".//c:axId | .//c:crossAx"):
        axis_id.set("val", axis_ids[axis_id.get("val")])
    chart.has_legend = chart.has_title = False
    chart.plots[0].has_data_labels = False
    for axis in (chart.category_axis, chart.value_axis):
        delete = axis._element.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}delete")
        if delete is None:
            delete = OxmlElement("c:delete")
            axis._element.insert_element_before(delete, "c:axPos")
        delete.set("val", "1")
    for grid in chart._chartSpace.xpath(".//c:majorGridlines | .//c:minorGridlines"):
        grid.getparent().remove(grid)
    series = chart.series[0]
    if bars:
        chart.plots[0].gap_width = 45
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string("8798A8")
        series.format.line.fill.background()
    else:
        series.marker.style = XL_MARKER_STYLE.NONE
        series.format.line.color.rgb = RGBColor.from_string("61758A")
        series.format.line.width = Pt(1.7)
        chart.plots[0].smooth = False
    # A planilha mantém o valor de cada ponto. A regra visual segue o Snapshot:
    # períodos sem valor ficam como lacuna, sem interpolação ou zero artificial.
    blanks = chart._chartSpace.chart.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}dispBlanksAs")
    if blanks is None:
        blanks = OxmlElement("c:dispBlanksAs")
        chart._chartSpace.chart.insert_element_before(blanks, "c:showDLblsOverMax", "c:extLst")
    blanks.set("val", "gap")
    for parent, successors in (
        (chart._chartSpace, ("c:txPr", "c:externalData", "c:printSettings", "c:userShapes", "c:extLst")),
        (chart._chartSpace.chart.plotArea, ("c:extLst",)),
    ):
        shape_pr = OxmlElement("c:spPr")
        shape_pr.append(OxmlElement("a:noFill"))
        line = OxmlElement("a:ln")
        line.append(OxmlElement("a:noFill"))
        shape_pr.append(line)
        # spPr precede txPr/externalData na sequência OOXML.
        parent.insert_element_before(shape_pr, *successors)
    return True


def _snapshot_slides(prs, snapshot):
    cards = list(snapshot.get("cards") or ())
    if not cards:
        raise ValueError("O resumo exige os indicadores do Snapshot.")
    grouped = [
        [card for card in cards if card.get("group") != "support"],
        [card for card in cards if card.get("group") == "support"],
    ]
    if not grouped[1]:
        grouped = [cards[:8], cards[8:]]
    if not grouped[1] or any(len(group) > 8 for group in grouped):
        raise ValueError("Os dois slides de Snapshot comportam até oito indicadores cada.")
    bank, base = str(snapshot.get("bank") or "Instituição"), str(snapshot.get("base") or "")
    current = str(snapshot.get("period_label") or snapshot.get("period") or "N/D")
    qoq_label = str(snapshot.get("qoq_label") or snapshot.get("qoq_period_label") or "Trimestre anterior")
    yoy_label = str(snapshot.get("yoy_label") or snapshot.get("yoy_period_label") or "Mesmo trimestre anterior")
    titles = ("Porte, capital e rentabilidade", "Funding e qualidade da carteira")
    for index, group in enumerate(grouped):
        slide = _slide(prs, "Snapshot · " + short_bank(bank), f"{current}  ·  {base}  ·  {titles[index]}", len(prs.slides) + 1)
        row_height = .565 if index == 0 else .68
        widths = [4.00, 2.05, 2.04, 2.04, _CONTENT - 10.13]
        heights = [.58] + [row_height] * len(group)
        y = 1.47
        table = _table(slide, len(group) + 1, 5, y, heights, widths)
        headers = ["Indicador", f"Valor · {current}", f"QoQ\n{qoq_label}", f"YoY\n{yoy_label}", "Evolução\nAté 8 períodos"]
        for column, header in enumerate(headers):
            _cell(table.cell(0, column), header, size=12, header=True, centered=column > 0)
        for row, card in enumerate(group, 1):
            _cell(table.cell(row, 0), card.get("label") or "Indicador", size=15)
            formatted = str(card.get("value") or "N/D")
            marker = str(card.get("status_marker") or "")
            _cell(table.cell(row, 1), formatted + (marker if marker and not formatted.endswith(marker) else ""), size=18, bold=True, centered=True)
            for column, comparison in ((2, "qoq"), (3, "yoy")):
                display, tone = _comparison(card, comparison)
                _cell(table.cell(row, column), display, size=13, centered=True, color=VARIATION_COLORS.get(tone, VARIATION_COLORS["neutral"]))
            _cell(table.cell(row, 4), "", centered=True)
            chart_y = y + heights[0] + (row - 1) * row_height + .06
            if not _history(slide, card.get("history"), _MARGIN + sum(widths[:4]) + .05, chart_y, widths[4] - .10, row_height - .12):
                _cell(table.cell(row, 4), "—", size=13, centered=True, color=_MUTED)
        footer = "BCB IFData: trimestral; Cadoc 4060: mensal, alinhado ao fechamento trimestral. "
        footer += "Cores: direção usual de crédito; cinza: contextual. N/D: dado ou comparação indisponível."
        if index == 1:
            notes = snapshot.get("source_notes") or snapshot.get("notes") or ()
            if isinstance(notes, str):
                notes = [notes]
            visible = [str(note) for note in notes if note][:3]
            if visible:
                _text(slide, _MARGIN, 5.73, _CONTENT, .68, "\n".join(visible), 10, color=_MUTED)
        _text(slide, _MARGIN, 6.70, _CONTENT, .52, footer, 9.5, color=_MUTED)
        slide.notes_slide.notes_text_frame.text = _json({
            "section": titles[index], "snapshot": {key: value for key, value in snapshot.items() if key != "cards"},
            "cards": group,
            "methodology": "Níveis, deltas, unidades e cores são os mesmos enviados pela tela. Históricos são gráficos Office editáveis com os mesmos pontos e lacunas.",
        })


def _unavailable_slide(prs, snapshot, title, reason):
    slide = _slide(prs, title, f"{snapshot.get('bank', 'Instituição')}  ·  {snapshot.get('base', '')}", len(prs.slides) + 1)
    _text(slide, _MARGIN, 1.64, _CONTENT, .55, "N/D para o recorte do Snapshot", 23, bold=True)
    _text(slide, _MARGIN, 2.47, 10.9, 1.5, reason or "Os dados necessários não estão disponíveis no período e perímetro selecionados.", 17, color=_MUTED)
    _text(slide, _MARGIN, 6.65, _CONTENT, .5, "A ausência permanece explícita. Os valores de outro perímetro não substituem este recorte.", 10, color=_MUTED)
    slide.notes_slide.notes_text_frame.text = _json({"section": title, "status": "unavailable", "reason": reason, "bank": snapshot.get("bank"), "base": snapshot.get("base")})


def _peers_slides(prs, snapshot, query, reason):
    if not query or not query.get("periods") or not query.get("metrics") or not query.get("banks"):
        _unavailable_slide(prs, snapshot, "Tabela de Peers", reason)
        return
    periods = sorted(set(query["periods"]), key=period_sort)[-3:]
    lookup = {(cell["metric"], cell["bank"], cell["period"]): cell for cell in query["cells"]}
    # O contexto do Snapshot normalmente tem uma instituição. A paginação
    # também conserva legibilidade quando o chamador fornece até três pares.
    for bank_start in range(0, len(query["banks"]), 3):
        banks = query["banks"][bank_start:bank_start + 3]
        for start in range(0, len(query["metrics"]), 8):
            keys = query["metrics"][start:start + 8]
            slide = _slide(prs, "Tabela de Peers", f"{query['base']}  ·  {query['scale']}  ·  {', '.join(short_bank(bank) for bank in banks)}", len(prs.slides) + 1)
            rows = []
            section = None
            for key in keys:
                metric = get_metric(key, query["base"])
                if section != metric.section:
                    rows.append(("section", metric.section))
                    section = metric.section
                rows.append(("metric", key))
            column_count = 1 + len(banks) * len(periods)
            widths = [4.1] + [(_CONTENT - 4.1) / (column_count - 1)] * (column_count - 1)
            metric_height = min(.48, (4.95 - .64 - .23 * sum(kind == "section" for kind, _ in rows)) / len(keys))
            heights = [.64] + [.23 if kind == "section" else metric_height for kind, _ in rows]
            table = _table(slide, len(rows) + 1, column_count, 1.45, heights, widths)
            _cell(table.cell(0, 0), "Indicador", header=True, size=13)
            for bank_index, bank in enumerate(banks):
                for period_index, period in enumerate(periods):
                    title = (short_bank(bank) + " · " if len(banks) > 1 else "") + period_label(period)
                    _cell(table.cell(0, 1 + bank_index * len(periods) + period_index), title + "\n" + period_comparison_label(period, query["mode"]), size=11, header=True, centered=True)
            for row_index, (kind, value) in enumerate(rows, 1):
                if kind == "section":
                    table.cell(row_index, 0).merge(table.cell(row_index, column_count - 1))
                    _cell(table.cell(row_index, 0), value, size=11, bold=True, background=_SECTION)
                    continue
                metric = get_metric(value, query["base"])
                unit = query["scale"].replace("R$ bilhões", "R$ bi").replace("R$ milhões", "R$ mm") if metric.unit == "R$" else metric.unit
                row_label = metric.label if metric.label.endswith(f"({unit})") else f"{metric.label} ({unit})"
                _cell(table.cell(row_index, 0), row_label, size=13)
                for bank_index, bank in enumerate(banks):
                    for period_index, period in enumerate(periods):
                        cell = lookup.get((value, bank, period), {})
                        target = table.cell(row_index, 1 + bank_index * len(periods) + period_index)
                        display = cell.get("display") or "N/D"
                        if cell.get("variation") == "Quebra em 2025" and cell.get("value") is not None:
                            display += "*"
                        if cell.get("status") in {"warning", "critical"}:
                            display += "†"
                        _cell(target, display, size=14, bold=True, centered=True)
                        _cell_delta(target, cell.get("variation") or "", variation_tone(metric, cell.get("direction"), cell.get("status", "unavailable")), size=10)
            footer = "BCB IFData (trimestral) / Cadoc 4060 (mensal). Taxas: Δ em bps inteiros; coberturas e custo/receita: Δ em p.p.; saldos: Δ em %."
            footer += "\nVerde: favorável; vermelho: atenção; cinza: contextual. †: alerta de qualidade; *: quebra de série. Detalhes de cálculo e fontes nas notas do slide."
            if len(periods) < 3:
                footer += f" Apenas {len(periods)} período(s) disponível(is) neste recorte."
            _text(slide, _MARGIN, 6.60, _CONTENT, .61, footer, 9.5, color=_MUTED)
            slide.notes_slide.notes_text_frame.text = _json({
                "query": {key: value for key, value in query.items() if key != "cells"},
                "methodology": methodology_rows({**query, "metrics": keys}),
                "cells": [cell for cell in query["cells"] if cell["metric"] in keys and cell["bank"] in banks and cell["period"] in periods],
            })


def _carteira_slides(prs, snapshot, model, reason):
    from tabs.carteira_4966 import (
        ROW_SPECS, MetricCell, VARIATION_NOTE, COLOR_NOTE, cell_delta,
        format_brl_millions, format_percentage, quality_issue_message,
    )
    if model is None or not model.periods:
        _unavailable_slide(prs, snapshot, "Carteira 4.966", reason)
        return
    periods = sorted(set(model.periods), key=period_sort)[-3:]
    groups = [
        ("Classificação da carteira", [spec for spec in ROW_SPECS if spec.group == "classification"]),
        ("Atrasos, provisão e cobertura", [spec for spec in ROW_SPECS if spec.group != "classification"]),
    ]
    for title, specs in groups:
        slide = _slide(prs, "Carteira 4.966", f"{str(snapshot.get('bank') or 'Instituição')}  ·  {snapshot.get('base') or 'Consolidada / Prudencial'}  ·  {title}", len(prs.slides) + 1)
        columns = 1 + len(periods) * 2
        widths = [4.15] + [(_CONTENT - 4.15) / (columns - 1)] * (columns - 1)
        row_height = .49 if len(specs) > 5 else .64
        heights = [.53, .30] + [row_height] * len(specs)
        table = _table(slide, len(specs) + 2, columns, 1.47, heights, widths)
        table.cell(0, 0).merge(table.cell(1, 0))
        _cell(table.cell(0, 0), "Indicador", header=True, size=13)
        for index, period in enumerate(periods):
            first = 1 + index * 2
            table.cell(0, first).merge(table.cell(0, first + 1))
            previous = cell_delta(model, specs[0], period).reference_period
            reference = model.period_labels.get(previous, period_label(previous)) if previous else "N/D"
            _cell(table.cell(0, first), model.period_labels.get(period, period_label(period)) + "\nQoQ vs " + reference, header=True, centered=True, size=11)
            for offset, unit in enumerate(("R$ mm", "%")):
                _cell(table.cell(1, first + offset), unit, header=True, centered=True, size=10)
        for row_index, spec in enumerate(specs, 2):
            _cell(table.cell(row_index, 0), spec.label, size=13, bold=spec.emphasis)
            for period_index, period in enumerate(periods):
                value = model.cells.get(spec.key, {}).get(period, MetricCell(None))
                first = 1 + period_index * 2
                if spec.layout == "paired":
                    targets = [(table.cell(row_index, first), value.primary, False), (table.cell(row_index, first + 1), value.secondary, True)]
                else:
                    table.cell(row_index, first).merge(table.cell(row_index, first + 1))
                    targets = [(table.cell(row_index, first), value.primary, spec.layout == "percent_span")]
                marker = "*" if model.cell_quality_issues(spec.key, period) else ""
                for target, number, percent in targets:
                    display = format_percentage(number, spec.percent_decimals) if percent else format_brl_millions(number)
                    _cell(target, display + marker, size=14, bold=spec.emphasis, centered=True)
                    delta = cell_delta(model, spec, period, secondary=percent and spec.layout == "paired")
                    _cell_delta(target, delta.display, delta.tone, size=10)
        base = model.period_labels.get(model.base_period, period_label(model.base_period)) if model.base_period else "N/D"
        footer = f"BCB IFData Rel. 16 e Rel. 2 · trimestral · {snapshot.get('base') or 'Consolidada / Prudencial'} · 4.966 desde 2025. Classificação: % da carteira-base de {base}."
        footer += "\nVencidos: operação integral com parcela >90 dias (arrasto), % da carteira do trimestre. PDD: |e2+f2+g2+h2|. Coberturas usam a PDD total."
        footer += "\nΔ taxas: bps inteiros; coberturas/base comum: p.p.; saldos: %. *: alerta de qualidade; consulte as notas do slide."
        source_notes = snapshot.get("carteira_notes") or snapshot.get("source_notes") or ()
        if isinstance(source_notes, str):
            source_notes = [source_notes]
        if source_notes:
            footer += "\n" + str(source_notes[0])
        if len(periods) < 3:
            footer += f" {len(periods)} período(s) disponível(is)."
        _text(slide, _MARGIN, 6.42, _CONTENT, .87, footer, 9.2, color=_MUTED)
        slide.notes_slide.notes_text_frame.text = _json({
            "periods": periods, "base_period": model.base_period, "base_value": model.base_value,
            "bank": snapshot.get("bank"), "scope": snapshot.get("base"), "source_notes": source_notes,
            "methodology": [VARIATION_NOTE, COLOR_NOTE],
            "rows": [{"key": spec.key, "label": spec.label, "definition": spec.help_text, "denominator": spec.denominator_label} for spec in specs],
            "cells": [{"key": spec.key, "period": period, "value": vars(model.cells.get(spec.key, {}).get(period, MetricCell(None))),
                       "deltas": [vars(cell_delta(model, spec, period, secondary=secondary)) for secondary in ([False, True] if spec.layout == "paired" else [False])]}
                      for spec in specs for period in periods],
            "quality": [quality_issue_message(issue) for issue in (*model.quality_issues, *model.reference_quality_issues)],
        })


def export_snapshot_powerpoint(snapshot: Mapping, peers_query=None, carteira_model=None, *, peers_unavailable_reason="", carteira_unavailable_reason="") -> bytes:
    """Exporta Snapshot (2 slides), Peers e 4.966, nessa ordem.

    ``snapshot.cards`` contém ``label``, ``value`` (texto já formatado), ``group``
    (hero/profit/support), ``qoq`` e ``yoy`` (texto ou display/tone/reason), além
    de ``history`` opcional (periods/values/kind). Fontes, status e ressalvas do
    payload completo são preservados nas notas editáveis de cada slide.
    """
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(_WIDTH), Inches(_HEIGHT)
    prs.core_properties.title = f"Resumo do Snapshot — {snapshot.get('bank', 'Instituição')}"
    prs.core_properties.subject = f"{snapshot.get('period_label', snapshot.get('period', ''))} | {snapshot.get('base', '')}"
    prs.core_properties.author = "Toma Conta"
    prs.core_properties.keywords = "Snapshot; Tabela de Peers; Carteira 4.966; editável"
    _snapshot_slides(prs, snapshot)
    _peers_slides(prs, snapshot, peers_query, peers_unavailable_reason)
    _carteira_slides(prs, snapshot, carteira_model, carteira_unavailable_reason)
    output = BytesIO()
    prs.save(output)
    return output.getvalue()


# Nome explícito para chamadores que identificam o formato pela extensão.
export_snapshot_pptx = export_snapshot_powerpoint
