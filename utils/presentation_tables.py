"""Native PowerPoint helpers extracted from Snapshot without changing its layout."""
from .financial_statement_style import (
    FONT, ACCENT as HEADER_BACKGROUND, SLIDE_WIDTH as _WIDTH, SLIDE_HEIGHT as _HEIGHT,
    SLIDE_MARGIN as _MARGIN, SLIDE_CONTENT as _CONTENT,
)
from .financial_statement_style import INK, MUTED, GRID, SECTION
_INK, _MUTED, _GRID, _SECTION = (value.lstrip("#") for value in (INK, MUTED, GRID, SECTION))


def _font(paragraph, size=14, *, bold=False, color=_INK, centered=False):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    paragraph.font.name = FONT
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = RGBColor.from_string(str(color).lstrip("#"))
    paragraph.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
    paragraph.space_before = paragraph.space_after = Pt(0)


def _text(slide, x, y, width, height, text, size=14, *, bold=False, color=_INK):
    from pptx.util import Inches
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    box.text_frame.margin_top = box.text_frame.margin_bottom = 0
    for index, line in enumerate(str(text).split("\n")):
        paragraph = box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
        paragraph.text = line
        _font(paragraph, size, bold=bold, color=color)
    return box


def _slide(prs, title, subtitle, page):
    from pptx.dml.color import RGBColor
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    _text(slide, _MARGIN, .24, _CONTENT - .5, .48, title, 25, bold=True)
    _text(slide, _MARGIN, .84, _CONTENT, .4, subtitle, 12, color=_MUTED)
    _text(slide, 12.53, .33, .36, .3, str(page), 10, color=_MUTED)
    return slide


def _borders(cell, *, header=False):
    """Grades discretas, sem faixas ou separadores verticais entre grupos."""
    from pptx.oxml.xmlchemy import OxmlElement
    tc_pr = cell._tc.get_or_add_tcPr()
    for name in ("lnL", "lnR", "lnT", "lnB"):
        for existing in tc_pr.findall(f"{{http://schemas.openxmlformats.org/drawingml/2006/main}}{name}"):
            tc_pr.remove(existing)
    # DrawingML exige as bordas antes do preenchimento da célula.
    for index, name in enumerate(("lnL", "lnR", "lnT", "lnB")):
        edge = OxmlElement(f"a:{name}")
        edge.set("w", "6350" if name == "lnB" and not header else "0")
        fill = OxmlElement("a:solidFill")
        color = OxmlElement("a:srgbClr")
        color.set("val", _GRID if not header else HEADER_BACKGROUND.lstrip("#"))
        fill.append(color)
        edge.append(fill)
        tc_pr.insert(index, edge)


def _cell(cell, text, *, size=14, bold=False, header=False, centered=False, color=_INK, background=None):
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR
    from pptx.util import Inches
    cell.text = str(text or "")
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string((background or (HEADER_BACKGROUND if header else "FFFFFF")).lstrip("#"))
    cell.margin_left = cell.margin_right = Inches(.09)
    cell.margin_top = cell.margin_bottom = Inches(.035)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text_frame.word_wrap = True
    _borders(cell, header=header)
    for paragraph in cell.text_frame.paragraphs:
        _font(paragraph, size, bold=bold or header, color="FFFFFF" if header else color, centered=centered)


def _table(slide, rows, columns, y, heights, widths):
    from pptx.util import Inches
    table = slide.shapes.add_table(rows, columns, Inches(_MARGIN), Inches(y), Inches(_CONTENT), Inches(sum(heights))).table
    for row, height in zip(table.rows, heights):
        row.height = Inches(height)
    for column, width in zip(table.columns, widths):
        column.width = Inches(width)
    return table


