"""Exports da nova tabela usando a mesma consulta fechada da interface."""
from __future__ import annotations

from io import BytesIO
from math import ceil
import re
from types import SimpleNamespace

import pandas as pd

from .comparison_table_style import HEADER_BACKGROUND, SECTION_BACKGROUND

from .peers_table_model import BY_KEY, get_metric, SCALES, BASELINES, COLORS, period_label, period_sort, short_bank, methodology_rows, comparison_label, period_comparison_label, variation_tone, VARIATION_COLORS, VARIATION_NOTE, COLOR_NOTE


def _lookup(query):
    return {(c["metric"], c["bank"], c["period"]): c for c in query["cells"]}


def footer(query, banks=None, metric_keys=None):
    n = len(banks or query["banks"])
    dates = ", ".join(period_label(p) for p in query["periods"])
    sources = [get_metric(key, query["base"]).source for key in (metric_keys or query["metrics"])]
    reports = sorted({int(r) for source in sources for r in re.findall(r"Rel\.\s*(\d+)", source)})
    source = "BCB IFData Rel. " + ", ".join(map(str, reports)) if reports else "BCB"
    if any("4060" in s for s in sources):
        source += " / Cadoc 4060"
    note = (
        "Δ por subtração: taxas em bps inteiros; coberturas e proporções em p.p. (2 casas); múltiplos em x. Saldos: Δ relativo em %.\n"
        "Verde: favorável; vermelho: atenção; cinza: contextual. Volumes de risco também dependem da carteira."
    )
    if any(c["status"] in {"warning", "critical"} or c.get("reference_status") in {"warning", "critical"} for c in query["cells"]):
        note += " † Alerta de qualidade; consulte Dados e status e a memória de cálculo."
    return f"{source}. {query['base']}. {n} instituições. {dates}.\n{query['query_type']}. Consulta: {query['queried_at']}.\n{note}"


def export_excel(query):
    import xlsxwriter
    buffer = BytesIO()
    with xlsxwriter.Workbook(buffer, {"in_memory": True, "strings_to_formulas": False, "strings_to_urls": False}) as workbook:
        base = {"font_name": "Calibri", "font_size": 11, "valign": "vcenter"}
        header = workbook.add_format({**base, "bold": True, "bg_color": HEADER_BACKGROUND, "font_color": "#FFFFFF", "border": 1, "border_color": "#D9D9D9", "align": "center", "text_wrap": True})
        text = workbook.add_format(base)
        section = workbook.add_format({**base, "bold": True, "bg_color": SECTION_BACKGROUND})
        variations = {tone: workbook.add_format({**base, "font_size": 9, "font_color": color, "align": "center"}) for tone, color in VARIATION_COLORS.items()}
        formats = {}
        for key in query["metrics"]:
            metric = get_metric(key, query["base"])
            numeric_format = "0." + "0" * metric.value_decimals
            formats[key] = workbook.add_format({**base, "num_format": numeric_format + "%" if metric.unit == "%" else numeric_format + '"x"' if metric.unit == "x" else "#,##" + numeric_format, "align": "center"})
        sheet = workbook.add_worksheet("Comparativo")
        sheet.freeze_panes(4, 2)
        sheet.set_column(0, 0, 39)
        sheet.set_column(1, 1, 14)
        sheet.set_column(2, 1 + len(query["banks"]) * len(query["periods"]), 14)
        sheet.write_string(0, 0, "Tabela de Peers", header)
        sheet.write_string(1, 0, query["base"] + " / " + BASELINES[query["mode"]], text)
        sheet.merge_range(2, 0, 3, 0, "Indicador", header)
        sheet.merge_range(2, 1, 3, 1, "Unidade", header)
        for b, bank in enumerate(query["banks"]):
            first = 2 + b * len(query["periods"])
            if len(query["periods"]) > 1:
                sheet.merge_range(2, first, 2, first + len(query["periods"]) - 1, bank, header)
            else:
                sheet.write_string(2, first, bank, header)
            for p, period in enumerate(query["periods"]):
                sheet.write_string(3, first + p, period_label(period) + "\n" + period_comparison_label(period, query["mode"]), header)
        sheet.set_row(3, 32 if query["mode"] != "none" else 20)
        cells = _lookup(query)
        row, current = 4, None
        for key in query["metrics"]:
            metric = get_metric(key, query["base"])
            if metric.section != current:
                sheet.merge_range(row, 0, row, 1 + len(query["banks"]) * len(query["periods"]), metric.section, section)
                current = metric.section
                row += 1
            sheet.write_string(row, 0, metric.label, text)
            sheet.write_string(row, 1, query["scale"] if metric.unit == "R$" else metric.unit, text)
            for b, bank in enumerate(query["banks"]):
                for p, period in enumerate(query["periods"]):
                    cell = cells[key, bank, period]
                    value = cell["value"]
                    col = 2 + b * len(query["periods"]) + p
                    if value is None:
                        sheet.write_blank(row, col, None, formats[key])
                    else:
                        sheet.write_number(row, col, value / SCALES[query["scale"]] if metric.unit == "R$" else value, formats[key])
                    comment = "\n".join(filter(None, [cell["variation"], cell["source"], cell["reason"], f"Status: {cell['status']}", f"Status referência: {cell.get('reference_status', 'N/D')}", cell.get("reference_reason")]))
                    sheet.write_comment(row, col, comment)
            row += 1
            if query["mode"] != "none":
                sheet.write_string(row, 0, "Variação", variations["neutral"])
                sheet.write_string(row, 1, metric.delta_unit, variations["neutral"])
                sheet.set_row(row, 14)
                for b, bank in enumerate(query["banks"]):
                    for p, period in enumerate(query["periods"]):
                        cell = cells[key, bank, period]
                        sheet.write_string(row, 2 + b * len(query["periods"]) + p, cell["variation"], variations[variation_tone(metric, cell["direction"], cell["status"], reference_status=cell.get("reference_status"))])
                row += 1
        sheet.write_string(row + 1, 0, footer(query), text)
        sheet.set_landscape()
        sheet.fit_to_pages(1, 0)
        sheet.repeat_rows(2, 3)
        numeric = workbook.add_worksheet("Dados e status")
        fields = ["metric", "bank", "period", "value", "reference", "reference_value", "variation", "status", "source", "reason", "delta_value", "delta_unit", "reference_status", "reference_reason"]
        titles = ["Indicador", "Instituição", "Período", "Valor (R$, decimal ou x)", "Referência", "Valor referência", "Variação", "Status", "Fonte", "Motivo", "Delta numérico", "Unidade delta", "Status referência", "Motivo referência"]
        for col, title in enumerate(titles):
            numeric.write_string(0, col, title, header)
        for row, cell in enumerate(query["cells"], 1):
            for col, field in enumerate(fields):
                value = cell.get(field)
                if isinstance(value, (int, float)):
                    numeric.write_number(row, col, value, text)
                elif value is not None:
                    numeric.write_string(row, col, str(value), text)
        numeric.freeze_panes(1, 3)
        numeric.autofilter(0, 0, len(query["cells"]), len(fields) - 1)
        numeric.set_column(0, 2, 32)
        numeric.set_column(3, 11, 25)
        for name, rows in (("Metodologia", methodology_rows(query)), ("Consulta", [{"Campo": k, "Valor": str(v)} for k, v in query.items() if k != "cells"])):
            page = workbook.add_worksheet(name)
            fields = list(rows[0])
            for col, field in enumerate(fields):
                page.write_string(0, col, field, header)
            for row, item in enumerate(rows, 1):
                for col, field in enumerate(fields):
                    page.write_string(row, col, str(item[field]), text)
            page.set_column(0, len(fields) - 1, 38)
    return buffer.getvalue()


def _text(slide, x, y, w, h, text, size=12, bold=False, color="222222"):
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.text_frame.word_wrap = True
    box.text_frame.margin_left = box.text_frame.margin_right = 0
    for i, line in enumerate(str(text).split("\n")):
        p = box.text_frame.paragraphs[0] if i == 0 else box.text_frame.add_paragraph()
        p.text = line
        p.font.name, p.font.size, p.font.bold = "Calibri", Pt(size), bold
        p.font.color.rgb = RGBColor.from_string(color)
    return box


def export_powerpoint(query, *, charts=False, chart_metrics=(), colors=None, chart_type="line"):
    # Reutiliza o exportador Office já adotado no SCR/SGS; a tabela é nativa.
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.oxml.xmlchemy import OxmlElement
    from .scr_pptx_export import exportar_paineis_pptx, _rotular_apenas_ultimo_ponto
    colors = colors or {bank: COLORS[i % len(COLORS)] for i, bank in enumerate(query["banks"])}
    cells = _lookup(query)
    if charts:
        panels = []
        for key in chart_metrics:
            metric = get_metric(key, query["base"])
            for start in range(0, len(query["banks"]), 3):
                banks = query["banks"][start:start + 3]
                records = []
                for bank in banks:
                    for period in query["periods"]:
                        y, q = period_sort(period)
                        value = cells[key, bank, period]["value"]
                        if value is not None and metric.unit == "R$":
                            value /= SCALES[query["scale"]]
                        records.append({"data_base": f"{y}-{q*3:02d}", "serie": bank, "valor": value, "denominador": pd.NA})
                series = pd.DataFrame(records)
                if series["valor"].notna().any():
                    panels.append(SimpleNamespace(titulo=metric.label, subtitulo=(query["scale"] if metric.unit == "R$" else metric.unit) + "; " + query["base"], fonte="", produto=metric.label, series=series, ordem_series=banks, ordem_categorias=sorted(series.data_base.unique()), cores=colors, tracejadas=[], metrica="peers", carteira_final_rs_mil=0, formato_numero="0.00%" if metric.unit == "%" else "0.00", formato_secundario="0.00", series_secundarias=[], tipo_grafico=chart_type, rotular_todos_pontos=False, estilo_taxas=False, preservar_lacunas=True, formato_data="mensal", cores_categorias={}, inverter_categorias=False))
                    panels[-1].metric_key = key
        if not panels:
            raise ValueError("Não há valores disponíveis nos indicadores selecionados para gráficos.")
        data, _ = exportar_paineis_pptx(panels, rotulo_serie_fn=short_bank, blocos_por_slide=[1] * len(panels))
        prs = Presentation(BytesIO(data))
        for slide, panel in zip(prs.slides, panels):
            _text(slide, .38, 6.95, 12.55, .45, footer(query, panel.ordem_series, [panel.metric_key]), size=9, color="555555")
            slide.notes_slide.notes_text_frame.text = VARIATION_NOTE + "\n" + COLOR_NOTE + "\n" + str(methodology_rows(query)) + "\n" + str({k: v for k, v in query.items() if k != "cells"})
            for shape in slide.shapes:
                if shape.has_chart:
                    chart = shape.chart
                    if chart_type == "column_stacked":
                        # O engine fornece as colunas nativas; peers usa agrupamento lado a lado.
                        for grouping in chart._chartSpace.xpath(".//c:barChart/c:grouping"):
                            grouping.set("val", "clustered")
                        for overlap in chart._chartSpace.xpath(".//c:barChart/c:overlap"):
                            overlap.set("val", "0")
                        chart.plots[0].gap_width = 70
                        from pptx.enum.chart import XL_DATA_LABEL_POSITION
                        for series, bank in zip(chart.series, panel.ordem_series):
                            valid = [i for i, v in enumerate(series.values) if v is not None]
                            if valid:
                                index = valid[-1]
                                _rotular_apenas_ultimo_ponto(series, index, panel.formato_numero, com_nome=False)
                                label = series.points[index].data_label
                                label.position = XL_DATA_LABEL_POSITION.OUTSIDE_END
                                label.font.name, label.font.size, label.font.bold = "Calibri", Pt(11), True
                                label.font.color.rgb = RGBColor.from_string(colors[bank].lstrip("#"))
                                for props in label._dLbl.findall("{http://schemas.openxmlformats.org/drawingml/2006/chart}spPr"):
                                    label._dLbl.remove(props)
                    # Lacunas continuam lacunas nos gráficos nativos.
                    blanks = chart._chartSpace.chart.find("{http://schemas.openxmlformats.org/drawingml/2006/chart}dispBlanksAs")
                    if blanks is None:
                        blanks = OxmlElement("c:dispBlanksAs")
                        chart._chartSpace.chart.append(blanks)
                    blanks.set("val", "gap")
                    for axis in (chart.category_axis, chart.value_axis):
                        axis.tick_labels.font.name = "Calibri"
                        axis.tick_labels.font.size = Pt(11)
                    # Mantém o negrito/cor por ponto do engine; define fonte explicitamente.
                    for props in chart._chartSpace.xpath(".//a:rPr | .//a:defRPr"):
                        props.set("sz", str(max(int(props.get("sz", "1100")), 1100)))
                        latin = props.find("{http://schemas.openxmlformats.org/drawingml/2006/main}latin")
                        if latin is None:
                            latin = OxmlElement("a:latin")
                            props.append(latin)
                        latin.set("typeface", "Calibri")
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        paragraph.font.name = "Calibri"
                        for run in paragraph.runs:
                            run.font.name = "Calibri"
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        page_size = ceil(len(query["metrics"]) / ceil(len(query["metrics"]) / 10))
        for start in range(0, len(query["banks"]), 3):
            banks = query["banks"][start:start + 3]
            for row_start in range(0, len(query["metrics"]), page_size):
                keys = query["metrics"][row_start:row_start + page_size]
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
                _text(slide, .38, .22, 12.55, .45, "Tabela de Peers", 22, True)
                _text(slide, .38, .75, 12.55, .4, query["base"] + "; " + BASELINES[query["mode"]] + "; " + query["scale"], 11)
                columns = 1 + len(banks) * len(query["periods"])
                table = slide.shapes.add_table(len(keys) + 2, columns, Inches(.38), Inches(1.3), Inches(12.55), Inches(.72 + .46 * len(keys))).table
                table.columns[0].width = Inches(3.25)
                for col in range(1, columns):
                    table.columns[col].width = Inches(9.3 / (columns - 1))
                for b, bank in enumerate(banks):
                    first = 1 + b * len(query["periods"])
                    if len(query["periods"]) > 1:
                        table.cell(0, first).merge(table.cell(0, first + len(query["periods"]) - 1))
                    table.cell(0, first).text = bank
                    for p, period in enumerate(query["periods"]):
                        table.cell(1, first + p).text = period_label(period) + "\n" + period_comparison_label(period, query["mode"])
                table.cell(0, 0).text = "Indicador"
                table.cell(0, 0).merge(table.cell(1, 0))
                for r, key in enumerate(keys, 2):
                    metric = get_metric(key, query["base"])
                    table.cell(r, 0).text = metric.label + "\n" + (query["scale"] if metric.unit == "R$" else metric.unit)
                    for b, bank in enumerate(banks):
                        for p, period in enumerate(query["periods"]):
                            cell = cells[key, bank, period]
                            target = table.cell(r, 1 + b * len(query["periods"]) + p)
                            target.text = cell["display"]
                            if cell["variation"]:
                                paragraph = target.text_frame.add_paragraph()
                                paragraph.text = cell["variation"]
                                paragraph.font.size = Pt(9)
                                paragraph.font.color.rgb = RGBColor.from_string(VARIATION_COLORS[variation_tone(metric, cell["direction"], cell["status"], reference_status=cell.get("reference_status"))].lstrip("#"))
                for r, row in enumerate(table.rows):
                    for c, cell in enumerate(row.cells):
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = RGBColor.from_string(HEADER_BACKGROUND.lstrip("#") if r < 2 else "FFFFFF")
                        cell.margin_left = cell.margin_right = Inches(.07)
                        cell.margin_top = cell.margin_bottom = Inches(.03)
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        for i, paragraph in enumerate(cell.text_frame.paragraphs):
                            paragraph.font.name = "Calibri"
                            paragraph.font.size = Pt(11 if i == 0 else 9)
                            paragraph.font.bold = r < 2
                            paragraph.space_before = paragraph.space_after = Pt(0)
                            if c == 0 and i > 0:
                                paragraph.font.color.rgb = RGBColor.from_string("666666")
                            if i == 0:
                                paragraph.font.color.rgb = RGBColor.from_string("FFFFFF" if r < 2 else "222222")
                            paragraph.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                _text(slide, .38, 6.55, 12.55, .75, footer(query, banks, keys), 9, color="555555")
                slide.notes_slide.notes_text_frame.text = VARIATION_NOTE + "\n" + COLOR_NOTE + "\n" + str(methodology_rows({**query, "metrics": keys})) + "\n" + str({k: v for k, v in query.items() if k != "cells"}) + "\n" + str([c for c in query["cells"] if c["metric"] in keys and c["bank"] in banks and c["reason"]])
    buffer = BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def export_png(query):
    from matplotlib.figure import Figure
    from matplotlib.backends.backend_agg import FigureCanvasAgg
    from matplotlib import font_manager
    from pathlib import Path
    font_path = Path("/Applications/Microsoft Excel.app/Contents/Resources/DFonts/Calibri.ttf")
    if font_path.exists():
        font_manager.fontManager.addfont(str(font_path))
    cells = _lookup(query)
    headers = ["Indicador"] + [short_bank(bank) + "\n" + period_label(p) + "\n" + period_comparison_label(p, query["mode"]) for bank in query["banks"] for p in query["periods"]]
    data, directions = [], {}
    for key in query["metrics"]:
        m = get_metric(key, query["base"])
        data.append([m.label + "\n" + (query["scale"] if m.unit == "R$" else m.unit)] + [cells[key, b, p]["display"] for b in query["banks"] for p in query["periods"]])
        if query["mode"] != "none":
            data.append(["Variação"] + [cells[key, b, p]["variation"] for b in query["banks"] for p in query["periods"]])
            for c, (b, p) in enumerate(((b, p) for b in query["banks"] for p in query["periods"]), 1):
                cell = cells[key, b, p]
                directions[len(data), c] = variation_tone(m, cell["direction"], cell["status"], reference_status=cell.get("reference_status"))
    footer_text = footer(query) + "\n" + BASELINES[query["mode"]]
    height = 2.2 + len(data) * .35
    table_bottom = max(.12, .22 * (footer_text.count("\n") + 2) / height)
    fig = Figure(figsize=(max(13, 3 + len(headers) * 1.25), height), facecolor="white")
    FigureCanvasAgg(fig)
    ax = fig.add_axes([.02, table_bottom, .96, .9 - table_bottom])
    ax.set_axis_off()
    table = ax.table(cellText=data, colLabels=headers, cellLoc="center", colWidths=[.25] + [.75 / (len(headers) - 1)] * (len(headers) - 1), bbox=[0, 0, 1, 1])
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    for (r, c), cell in table.get_celld().items():
        cell.set_edgecolor("#D9D9D9")
        cell.set_linewidth(.5)
        cell.set_facecolor(HEADER_BACKGROUND if r == 0 else "white")
        cell.get_text().set_fontfamily("Calibri" if font_path.exists() else "DejaVu Sans")
        if r == 0:
            cell.set_height(cell.get_height() * 1.5)
            cell.get_text().set_weight("bold")
            cell.get_text().set_color("white")
        if c == 0:
            cell.get_text().set_horizontalalignment("left")
        if (r, c) in directions:
            cell.get_text().set_fontsize(9)
            cell.get_text().set_color(VARIATION_COLORS[directions[r,c]])
    family = "Calibri" if font_path.exists() else "DejaVu Sans"
    fig.text(.02, .95, "Tabela de Peers", fontsize=18, fontweight="bold", fontfamily=family)
    fig.text(.02, .025, footer_text, fontsize=9, fontfamily=family)
    output = BytesIO()
    fig.savefig(output, format="png", dpi=180)
    return output.getvalue()
